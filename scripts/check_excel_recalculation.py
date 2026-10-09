"""Recalculate disposable exports with LibreOffice; never modify source templates.

This is a real calculation-engine gate, not Excel/WPS visual acceptance.
Run with PYTHONPATH=. python scripts/check_excel_recalculation.py --output result.json
"""

import argparse
import json
import math
import shutil
import subprocess
import tempfile
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook

from app.core.context import AppContext
from app.core.schemas import DefectInput, InspectionInput, RecordFilter

NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}

# Literal expected results, independent of application statistics and formulas.
# Each tuple is: inspected, sampled, defective, rework batches, project a, project b.
CASES = [
    ("annual", "2026-01-01", "2026-12-31", "全部",
     (1500, 300, 15, 3, 3, 3), (500, 100, 5, 0, 1, 1), "2026-W53"),
    ("january", "2026-01-01", "2026-12-31", "2026-01",
     (300, 60, 3, 1, 1, 1), (500, 100, 5, 1, 1, 1), "2026-W05"),
    ("february", "2026-01-01", "2026-12-31", "2026-02",
     (700, 140, 7, 2, 1, 1), (400, 80, 4, 1, 0, 1), "2026-W09"),
    ("empty-month", "2026-01-01", "2026-12-31", "2026-03",
     (0, 0, 0, 0, 0, 0), (0, 0, 0, 0, 0, 0), "2026-W14"),
    ("december", "2026-01-01", "2026-12-31", "2026-12",
     (500, 100, 5, 0, 1, 1), (500, 100, 5, 0, 1, 1), "2026-W53"),
    ("cross-year", "2020-12-30", "2021-01-01", "全部",
     (200, 30, 3, 2, 1, 1), (200, 30, 3, 2, 1, 1), "2020-W53"),
    ("cross-year-january", "2020-12-30", "2021-01-01", "2021-01",
     (110, 20, 2, 1, 0, 1), (200, 30, 3, 2, 1, 1), "2020-W53"),
    ("leap-day", "2024-02-28", "2024-03-01", "2024-02",
     (210, 30, 3, 1, 1, 1), (210, 30, 3, 1, 1, 1), "2024-W09"),
]


def check_caches(path, wb):
    with ZipFile(path) as archive:
        charts = [p for p in archive.namelist() if p.startswith("xl/charts/chart") and p.endswith(".xml")]
        assert len(charts) == 6
        for part in charts:
            xml = ET.fromstring(archive.read(part))
            assert len(xml.findall(".//c:ser", NS)) == 1
            for kind in ("num", "str"):
                for ref in xml.findall(f".//c:{kind}Ref", NS):
                    sheet, cells = ref.findtext("c:f", namespaces=NS).rsplit("!", 1)
                    expected = [c.value for row in wb[sheet.strip("'")][cells] for c in row]
                    cache = ref.find(f"c:{kind}Cache", NS)
                    assert cache is not None, (part, cells)
                    assert int(cache.find("c:ptCount", NS).get("val")) == len(expected)
                    points = cache.findall("c:pt", NS)
                    values = {int(p.get("idx")): p.findtext("c:v", "", NS) for p in points}
                    assert len(values) == len(points) and all(0 <= i < len(expected) for i in values)
                    for i, value in enumerate(expected):
                        if kind == "num":
                            assert i in values and math.isclose(float(values[i]), value or 0, abs_tol=1e-10), (part, i)
                        else:
                            assert values.get(i, "") == str(value or ""), (part, i)


def check_result(path, case):
    name, _, _, selection, period, week, iso_week = case
    wb = load_workbook(path, data_only=True)
    try:
        ws = wb["数据分析表"]
        assert ws["A2"].value == selection, name
        assert ws["A32"].value == iso_week, name
        for row, expected in ((4, period), (34, week)):
            cells = [f"H{row}", f"I{row}", f"J{row}", f"F{row + 8}", f"B{row}", f"B{row + 1}"]
            actual = tuple(ws[c].value for c in cells)
            assert actual == expected, (name, actual, expected)
            rate = expected[2] / expected[1] if expected[1] else 0
            assert math.isclose(ws[f"K{row}"].value, rate, abs_tol=1e-10), name
        assert all(c.data_type != "e" for sheet in wb for row in sheet for c in row), name
        check_caches(path, wb)
        return {"case": name, "selection": selection, "statistics": "PASS", "six_chart_caches": "PASS"}
    finally:
        wb.close()


def run():
    soffice = shutil.which("soffice")
    if not soffice:
        raise RuntimeError("LibreOffice Calc is required for the recalculation gate")
    version = subprocess.check_output([soffice, "--version"], text=True, timeout=30).strip()
    with tempfile.TemporaryDirectory(prefix="qc-recalculation-") as temp:
        root = Path(temp)
        source, converted = root / "source", root / "converted"
        source.mkdir()
        converted.mkdir()
        ctx = AppContext(root / "data")
        try:
            rows = [
                ("2020-12-31", 90, 10, 1, "返工", [1]),
                ("2021-01-01", 110, 20, 2, "返工", [2]),
                ("2024-02-29", 210, 30, 3, "返工", [1, 2]),
                ("2026-01-05", 100, 20, 1, "返工", [1]),
                ("2026-01-31", 200, 40, 2, "合格", [2]),
                ("2026-02-01", 300, 60, 3, "返工", [1]),
                ("2026-02-28", 400, 80, 4, "返工", [2]),
                ("2026-12-31", 500, 100, 5, "合格", [1, 2]),
            ]
            for day, inspected, sampled, defective, judgment, ids in rows:
                ctx.inspections.save(InspectionInput(
                    inspection_date=day, team="U1", work_order=f"模拟-{day}",
                    inspection_quantity=inspected, sampling_quantity=sampled,
                    defect_quantity=defective, judgment=judgment, inspector="测试员", source="demo",
                    defects=[DefectInput(defect_id=i) for i in ids],
                ))
            for name, low, high, selection, *_ in CASES:
                path = ctx.excel.export(source / f"{name}.xlsx", RecordFilter(
                    start=low, end=high, source="demo",
                ), legacy=True, prefer_com=False)
                wb = load_workbook(path)
                wb["数据分析表"]["A2"] = selection
                wb.save(path)
                wb.close()
        finally:
            ctx.db.dispose()

        def convert(paths, out):
            subprocess.run([
                soffice, f"-env:UserInstallation={(root / 'profile').as_uri()}",
                "--headless", "--convert-to", "xlsx", "--outdir", str(out), *map(str, paths),
            ], check=True, capture_output=True, text=True, timeout=180)

        convert(sorted(source.glob("*.xlsx")), converted)
        results = [check_result(converted / f"{case[0]}.xlsx", case) for case in CASES]
        # Change an already recalculated/saved January export back to all months.
        wb = load_workbook(converted / "january.xlsx")
        wb["数据分析表"]["A2"] = "全部"
        back = source / "back-to-all.xlsx"
        wb.save(back)
        wb.close()
        convert([back], converted)
        results.append(check_result(converted / back.name, ("back-to-all", *CASES[0][1:])))
    return {"engine": version, "cases": results, "excel_wps_visual_acceptance": "未完成实机验收"}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    result = run()
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False))
