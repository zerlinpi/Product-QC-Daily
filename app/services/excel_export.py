import logging
import os
from calendar import monthrange
from copy import copy
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.data_source import AxDataSource, NumData, NumVal, StrData, StrRef, StrVal
from openpyxl.chart.text import RichText
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.text import (
    CharacterProperties,
    Paragraph,
    ParagraphProperties,
    RichTextProperties,
)
from openpyxl.drawing.text import Font as DrawingFont
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.workbook.properties import CalcProperties
from openpyxl.workbook.views import BookView
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.print_settings import PrintArea
from openpyxl.worksheet.views import Pane, Selection

from app.core.labels import source_label
from app.services.excel_common import load_compatible, write_text
from app.services.statistics_service import rates

LEGACY_FORM_ROWS = 500

HEADERS = [
    "填写ID",
    "时间",
    "检验:组别",
    "检验:加工单号",
    "检验:检验数量",
    "检验:抽检数",
    "检验:不良数",
    "检验:不良项目",
    "检验:判定",
    "签名:图片",
    "检验员",
    "备注",
    "数据来源",
    "年份",
    "月份",
    "日期",
]


def display_width(value) -> int:
    text = str(value if value is not None else "")
    return sum(2 if ord(char) > 127 else 1 for char in text)


def header_footer_text(value, limit=72) -> str:
    """Keep user text safe inside Excel header/footer control syntax."""
    text = " ".join(str(value or "").split())[:limit]
    return text.replace("&", "&&")


def wrapped_row_height(*fields, base=22, line_height=15, maximum=90) -> float:
    """Estimate a readable Excel row height for wrapped Chinese/English text."""
    lines = 1
    for value, column_width in fields:
        width = max(int(float(column_width or 10) * 0.9), 8)
        text_lines = str(value if value is not None else "").splitlines() or [""]
        wrapped_lines = sum(
            max(1, (display_width(text) + width - 1) // width) for text in text_lines
        )
        lines = max(lines, wrapped_lines)
    return min(maximum, max(base, lines * line_height + 6))


def apply_standard_report_identity(wb, settings: dict, filters, start, end) -> None:
    company = str(settings.get("company") or "").strip()
    factory = str(settings.get("factory") or "").strip()
    identity = " · ".join(value for value in (company, factory) if value) or "成品日检管理系统"
    date_range = f"{start or '未限定'} 至 {end or '未限定'}"
    source = source_label(filters.source)

    wb.properties.creator = identity
    wb.properties.lastModifiedBy = "Product-QC-Daily"
    wb.properties.title = f"成品日检质量报表 · {date_range}"
    wb.properties.subject = f"{source} · {date_range}"

    for ws in wb:
        ws.oddHeader.left.text = header_footer_text(identity, 64)
        ws.oddHeader.center.text = header_footer_text(f"质量报表 · {ws.title}", 48)
        ws.oddHeader.right.text = header_footer_text(date_range, 64)
        ws.oddFooter.left.text = header_footer_text(source, 32)
        ws.oddFooter.center.text = "第 &P 页 / 共 &N 页"
        ws.oddFooter.right.text = "&F"


def style_table(ws):
    """Format the standard workbook as a print-ready QC report.

    The legacy/template export deliberately bypasses this function so its
    original layout remains byte-for-byte independent from standard styling.
    """
    last_row, last_col = max(ws.max_row, 1), max(ws.max_column, 1)
    last_letter = get_column_letter(last_col)
    border = Border(
        left=Side(style="thin", color="D9E2F3"),
        right=Side(style="thin", color="D9E2F3"),
        top=Side(style="thin", color="D9E2F3"),
        bottom=Side(style="thin", color="D9E2F3"),
    )
    header_fill = PatternFill("solid", fgColor="1F4E78")
    band_fill = PatternFill("solid", fgColor="F5F9FC")
    header_border = Border(
        left=Side(style="thin", color="D9E2F3"),
        right=Side(style="thin", color="D9E2F3"),
        top=Side(style="thin", color="D9E2F3"),
        bottom=Side(style="medium", color="17365D"),
    )
    total_border = Border(
        left=Side(style="thin", color="D9E2F3"),
        right=Side(style="thin", color="D9E2F3"),
        top=Side(style="medium", color="4472C4"),
        bottom=Side(style="thin", color="D9E2F3"),
    )

    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.page_margins.header = 0.2
    ws.page_margins.footer = 0.2
    ws.print_options.horizontalCentered = True
    ws.print_title_rows = "1:1"

    filter_sheets = {"检验记录", "不良明细", "不良项目", "月度统计"}
    if ws.title in filter_sheets:
        ws.auto_filter.ref = f"A1:{last_letter}{last_row}"
    else:
        ws.auto_filter.ref = None
    ws.sheet_view.zoomScale = {
        "检验记录": 85,
        "不良明细": 90,
        "不良项目": 95,
        "统计摘要": 100,
        "月度统计": 85,
    }.get(ws.title, 100)

    for cell in ws[1]:
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.border = header_border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28

    for row in range(2, last_row + 1):
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 18, 22)
        # ws[row] rescans all populated cells to rediscover max_column for
        # every row. The bounds above already cover the complete table.
        for column in range(1, last_col + 1):
            cell = ws.cell(row, column)
            cell.font = Font(name="Microsoft YaHei", size=9)
            cell.border = border
            cell.alignment = Alignment(vertical="center")
            if row % 2 == 0 and ws.title not in {"统计摘要"}:
                cell.fill = band_fill

    width_overrides = {
        "检验记录": {
            "A": 22,
            "B": 24,
            "C": 12,
            "D": 22,
            "E": 13,
            "F": 12,
            "G": 12,
            "H": 20,
            "I": 11,
            "J": 16,
            "K": 14,
            "L": 34,
            "M": 12,
            "N": 9,
            "O": 11,
            "P": 13,
        },
        "不良明细": {"A": 22, "B": 10, "C": 20, "D": 18, "E": 36},
        "不良项目": {"A": 10, "B": 22, "C": 16, "D": 10, "E": 10, "F": 38},
        "统计摘要": {"A": 24, "B": 48},
        "月度统计": {"A": 12, "B": 12, "C": 14, "D": 14, "E": 12, "F": 12, "G": 12, "H": 12},
    }
    for column in range(1, last_col + 1):
        letter = get_column_letter(column)
        if letter in width_overrides.get(ws.title, {}):
            ws.column_dimensions[letter].width = width_overrides[ws.title][letter]
            continue
        measured = max(
            (display_width(ws.cell(row, column).value) for row in range(1, min(last_row, 300) + 1)),
            default=10,
        )
        ws.column_dimensions[letter].width = min(36, max(10, measured + 2))

    if ws.title == "检验记录":
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        for row in range(2, last_row + 1):
            ws.cell(row, 1).alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            ws.cell(row, 2).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row, 4).alignment = Alignment(vertical="center", wrap_text=True)
            for column in (3, 9, 13, 14, 15, 16):
                ws.cell(row, column).alignment = Alignment(horizontal="center", vertical="center")
            for column in (5, 6, 7):
                ws.cell(row, column).alignment = Alignment(horizontal="right", vertical="center")
                ws.cell(row, column).number_format = "#,##0"
            ws.cell(row, 8).alignment = Alignment(vertical="center", wrap_text=True)
            ws.cell(row, 12).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row, 16).number_format = "yyyy-mm-dd"
            ws.row_dimensions[row].height = max(
                ws.row_dimensions[row].height or 22,
                wrapped_row_height(
                    (ws.cell(row, 1).value, ws.column_dimensions["A"].width),
                    (ws.cell(row, 4).value, ws.column_dimensions["D"].width),
                    (ws.cell(row, 8).value, ws.column_dimensions["H"].width),
                    (ws.cell(row, 12).value, ws.column_dimensions["L"].width),
                    maximum=96,
                ),
            )
            defect_quantity = ws.cell(row, 7)
            if isinstance(defect_quantity.value, (int, float)) and defect_quantity.value > 0:
                defect_quantity.fill = PatternFill("solid", fgColor="FCE4D6")
                defect_quantity.font = Font(
                    name="Microsoft YaHei", size=9, bold=True, color="C65911"
                )
            judgment = ws.cell(row, 9)
            if judgment.value == "合格":
                judgment.fill = PatternFill("solid", fgColor="E2F0D9")
                judgment.font = Font(name="Microsoft YaHei", size=9, bold=True, color="375623")
            elif judgment.value == "返工":
                judgment.fill = PatternFill("solid", fgColor="FCE4D6")
                judgment.font = Font(name="Microsoft YaHei", size=9, bold=True, color="C65911")
        ws.print_area = f"A1:{last_letter}{last_row}"
        ws.sheet_properties.tabColor = "5B9BD5"
    elif ws.title == "不良明细":
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        for row in range(2, last_row + 1):
            ws.cell(row, 1).alignment = Alignment(vertical="center", wrap_text=True)
            ws.cell(row, 2).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row, 4).alignment = Alignment(horizontal="right", vertical="center")
            ws.cell(row, 4).number_format = "#,##0"
            ws.cell(row, 5).alignment = Alignment(vertical="top", wrap_text=True)
            ws.row_dimensions[row].height = max(
                ws.row_dimensions[row].height or 22,
                wrapped_row_height(
                    (ws.cell(row, 1).value, ws.column_dimensions["A"].width),
                    (ws.cell(row, 3).value, ws.column_dimensions["C"].width),
                    (ws.cell(row, 5).value, ws.column_dimensions["E"].width),
                    maximum=90,
                ),
            )
        ws.print_area = f"A1:{last_letter}{last_row}"
        ws.sheet_properties.tabColor = "ED7D31"
    elif ws.title == "不良项目":
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        for row in range(2, last_row + 1):
            ws.cell(row, 1).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row, 4).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row, 5).alignment = Alignment(horizontal="right", vertical="center")
            ws.cell(row, 5).number_format = "#,##0"
            ws.cell(row, 6).alignment = Alignment(vertical="top", wrap_text=True)
            ws.row_dimensions[row].height = max(
                ws.row_dimensions[row].height or 22,
                wrapped_row_height(
                    (ws.cell(row, 2).value, ws.column_dimensions["B"].width),
                    (ws.cell(row, 6).value, ws.column_dimensions["F"].width),
                    maximum=78,
                ),
            )
        ws.print_area = f"A1:{last_letter}{last_row}"
        ws.sheet_properties.tabColor = "A5A5A5"
    elif ws.title == "统计摘要":
        ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
        note_labels = {"口径", "数据范围", "日期范围"}
        for row in range(2, last_row + 1):
            label_cell = ws.cell(row, 1)
            value = ws.cell(row, 2)
            label_cell.font = Font(name="Microsoft YaHei", size=9, bold=True, color="44546A")
            label_cell.alignment = Alignment(horizontal="left", vertical="center")
            value.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
            if label_cell.value in note_labels:
                label_cell.fill = PatternFill("solid", fgColor="F2F2F2")
                value.fill = PatternFill("solid", fgColor="F2F2F2")
                value.font = Font(name="Microsoft YaHei", size=9, color="595959")
                value.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                ws.row_dimensions[row].height = max(
                    ws.row_dimensions[row].height or 22,
                    wrapped_row_height(
                        (value.value, ws.column_dimensions["B"].width),
                        base=28,
                        maximum=60,
                    ),
                )
            else:
                ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 22, 24)
                metric = str(label_cell.value or "")
                if metric in {"不良件数", "不良率", "返工批次", "返工率"}:
                    metric_fill = PatternFill("solid", fgColor="FCE4D6")
                    metric_color = "C65911"
                elif metric in {"合格批次", "合格率"}:
                    metric_fill = PatternFill("solid", fgColor="E2F0D9")
                    metric_color = "375623"
                else:
                    metric_fill = PatternFill("solid", fgColor="D9EAF7")
                    metric_color = "1F4E78"
                label_cell.fill = metric_fill
                value.fill = metric_fill
                value.font = Font(name="Microsoft YaHei", size=11, bold=True, color=metric_color)
            if isinstance(value.value, (int, float)) and value.number_format != "0.00%":
                value.number_format = "#,##0"
        ws.print_area = f"A1:B{last_row}"
        ws.sheet_properties.tabColor = "70AD47"
    elif ws.title == "月度统计":
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        total_row = last_row if ws.cell(last_row, 1).value == "合计" else None
        data_last_row = last_row - 1 if total_row else last_row
        ws.auto_filter.ref = f"A1:H{max(data_last_row, 1)}"
        for row in range(2, last_row + 1):
            ws.cell(row, 1).alignment = Alignment(horizontal="center", vertical="center")
            for column in range(2, 9):
                ws.cell(row, column).alignment = Alignment(horizontal="right", vertical="center")
            for column in (2, 3, 4, 5, 7):
                ws.cell(row, column).number_format = "#,##0"
        if total_row:
            ws.row_dimensions[total_row].height = max(ws.row_dimensions[total_row].height or 22, 24)
            for column in range(1, 9):
                cell = ws.cell(total_row, column)
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
                cell.font = Font(name="Microsoft YaHei", size=9, bold=True, color="1F1F1F")
                cell.border = total_border
        if data_last_row >= 2:
            for column, color in {
                "B": "5B9BD5",
                "C": "5B9BD5",
                "D": "5B9BD5",
                "E": "ED7D31",
                "G": "A5A5A5",
            }.items():
                ws.conditional_formatting.add(
                    f"{column}2:{column}{data_last_row}",
                    DataBarRule(
                        start_type="num",
                        start_value=0,
                        end_type="max",
                        color=color,
                        showValue=True,
                    ),
                )
            for column, color in (("F", "ED7D31"), ("H", "A5A5A5")):
                ws.conditional_formatting.add(
                    f"{column}2:{column}{data_last_row}",
                    DataBarRule(
                        start_type="num",
                        start_value=0,
                        end_type="num",
                        end_value=1,
                        color=color,
                        showValue=True,
                    ),
                )
        ws.print_area = f"A1:W{max(last_row, 32)}"
        ws.sheet_properties.tabColor = "4472C4"


def month_keys(start: date | None, end: date | None) -> list[str]:
    if start is None or end is None:
        return []
    current = date(start.year, start.month, 1)
    last = date(end.year, end.month, 1)
    result = []
    while current <= last:
        result.append(current.strftime("%Y-%m"))
        current = (
            date(current.year + 1, 1, 1)
            if current.month == 12
            else date(current.year, current.month + 1, 1)
        )
    return result


def embed_chart_data(chart, categories: list[str], series_values: list[list], titles: list[str], formats: list[str]):
    """Embed chart caches so Excel/WPS can render data without recalculating first."""
    if not (len(chart.series) == len(series_values) == len(titles) == len(formats)):
        raise ValueError("图表系列数据不完整")
    for series, values, title, number_format in zip(
        chart.series, series_values, titles, formats, strict=True
    ):
        category_formula = None
        if series.cat is not None:
            for reference_name in ("strRef", "numRef"):
                reference = getattr(series.cat, reference_name, None)
                if reference is not None and reference.f:
                    category_formula = reference.f
                    break
        if category_formula:
            series.cat = AxDataSource(
                strRef=StrRef(
                    f=category_formula,
                    strCache=StrData(
                        ptCount=len(categories),
                        pt=[StrVal(idx=index, v=str(value)) for index, value in enumerate(categories)],
                    ),
                )
            )
        if series.val is not None and series.val.numRef is not None:
            series.val.numRef.numCache = NumData(
                formatCode=number_format,
                ptCount=len(values),
                pt=[
                    NumVal(idx=index, v=0 if value is None else value)
                    for index, value in enumerate(values)
                ],
            )
        if series.tx is not None and series.tx.strRef is not None:
            series.tx.strRef.strCache = StrData(
                ptCount=1,
                pt=[StrVal(idx=0, v=title)],
            )


def new_chart_snapshot():
    return {
        "inspection_quantity": 0,
        "sampling_quantity": 0,
        "defect_quantity": 0,
        "defect_batches": {code: 0 for code in "abcdefghijklmnopqrstuvwx"},
        "rework_batches": {},
    }


def update_chart_snapshot(snapshot, row):
    snapshot["inspection_quantity"] += row["inspection_quantity"]
    snapshot["sampling_quantity"] += row["sampling_quantity"]
    snapshot["defect_quantity"] += row["defect_quantity"]
    if row["judgment"] == "返工":
        team = row["team"]
        snapshot["rework_batches"][team] = snapshot["rework_batches"].get(team, 0) + 1
    for code in {item["code"] for item in row["defects"]}:
        if code in snapshot["defect_batches"]:
            snapshot["defect_batches"][code] += 1


def embed_legacy_analysis_chart_data(wb, period, week, teams, defect_names, other_teams=False):
    analysis = wb["数据分析表"]
    team_categories = [teams[index] if index < len(teams) else "" for index in range(8)]
    defect_categories = [defect_names.get(chr(97 + index), "") for index in range(24)]
    totals_categories = ["检验数量", "抽检数", "不良数", "不良率"]

    def team_values(snapshot):
        result = [snapshot["rework_batches"].get(team, 0) if team else 0 for team in team_categories]
        if other_teams:
            result[-1] = sum(snapshot["rework_batches"].values()) - sum(result[:-1])
        return result

    def totals(snapshot):
        rate = (
            snapshot["defect_quantity"] / snapshot["sampling_quantity"]
            if snapshot["sampling_quantity"]
            else 0
        )
        return [
            snapshot["inspection_quantity"],
            snapshot["sampling_quantity"],
            snapshot["defect_quantity"],
            rate,
        ]

    values_by_range = {
        "$F$4:$F$11": (
            team_categories,
            team_values(period),
            "#,##0",
        ),
        "$H$4:$K$4": (totals_categories, totals(period), "General"),
        "$B$4:$B$27": (
            defect_categories,
            [period["defect_batches"][chr(97 + index)] for index in range(24)],
            "#,##0",
        ),
        "$H$34:$K$34": (totals_categories, totals(week), "General"),
        "$B$34:$B$57": (
            defect_categories,
            [week["defect_batches"][chr(97 + index)] for index in range(24)],
            "#,##0",
        ),
        "$F$34:$F$41": (
            team_categories,
            team_values(week),
            "#,##0",
        ),
    }
    for chart in analysis._charts:
        if not chart.series:
            continue
        series = chart.series[0]
        formula = (
            series.val.numRef.f.replace("'", "")
            if series.val is not None and series.val.numRef is not None
            else ""
        )
        match = next((key for key in values_by_range if formula.endswith(key)), None)
        if match is None:
            continue
        categories, values, number_format = values_by_range[match]
        embed_chart_data(chart, categories, [values], [""], [number_format])


def add_monthly_analysis(wb, monthly: dict[str, dict], start: date | None, end: date | None):
    ws = wb.create_sheet("月度统计")
    ws.append(
        ["月份", "检验批次", "检验数量", "抽检数量", "不良件数", "不良率", "返工批次", "返工率"]
    )
    months = month_keys(start, end) or sorted(monthly)
    totals = {
        "batches": 0,
        "inspection_quantity": 0,
        "sampling_quantity": 0,
        "defect_quantity": 0,
        "rework_batches": 0,
    }
    for month in months:
        values = monthly.get(month, {key: 0 for key in totals})
        defect_rate = (
            values["defect_quantity"] / values["sampling_quantity"]
            if values["sampling_quantity"]
            else 0
        )
        rework_rate = values["rework_batches"] / values["batches"] if values["batches"] else 0
        ws.append(
            [
                month,
                values["batches"],
                values["inspection_quantity"],
                values["sampling_quantity"],
                values["defect_quantity"],
                defect_rate,
                values["rework_batches"],
                rework_rate,
            ]
        )
        for key in totals:
            totals[key] += values[key]
        ws.cell(ws.max_row, 6).number_format = "0.00%"
        ws.cell(ws.max_row, 8).number_format = "0.00%"

    month_last_row = 1 + len(months)
    total_defect_rate = (
        totals["defect_quantity"] / totals["sampling_quantity"] if totals["sampling_quantity"] else 0
    )
    total_rework_rate = totals["rework_batches"] / totals["batches"] if totals["batches"] else 0
    ws.append(
        [
            "合计",
            totals["batches"],
            totals["inspection_quantity"],
            totals["sampling_quantity"],
            totals["defect_quantity"],
            total_defect_rate,
            totals["rework_batches"],
            total_rework_rate,
        ]
    )
    total_row = ws.max_row
    for column in range(1, 9):
        cell = ws.cell(total_row, column)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.font = Font(name="Microsoft YaHei", size=9, bold=True, color="1F1F1F")
    ws.cell(total_row, 6).number_format = "0.00%"
    ws.cell(total_row, 8).number_format = "0.00%"

    if months:
        categories = Reference(ws, min_col=1, min_row=2, max_row=month_last_row)
        volume = BarChart()
        volume.style = 10
        volume.title = "月度检验批次趋势"
        volume.y_axis.title = "批次"
        volume.x_axis.title = "月份"
        volume.x_axis.tickLblSkip = max(1, (len(months) + 17) // 18)
        volume.x_axis.tickMarkSkip = volume.x_axis.tickLblSkip
        volume.height = 7
        volume.width = 14
        volume.legend = None
        volume.add_data(
            Reference(ws, min_col=2, max_col=2, min_row=1, max_row=month_last_row),
            titles_from_data=True,
        )
        volume.set_categories(categories)
        embed_chart_data(
            volume,
            months,
            [[ws.cell(row, 2).value for row in range(2, month_last_row + 1)]],
            ["检验批次"],
            ["0"],
        )
        ws.add_chart(volume, "J2")

        rates = LineChart()
        rates.style = 13
        rates.title = "月度质量率趋势"
        rates.y_axis.title = "比例"
        rates.y_axis.numFmt = "0.0%"
        rates.y_axis.scaling.min = 0
        rates.x_axis.title = "月份"
        rates.x_axis.tickLblSkip = max(1, (len(months) + 17) // 18)
        rates.x_axis.tickMarkSkip = rates.x_axis.tickLblSkip
        rates.height = 7
        rates.width = 14
        rates.legend.position = "b"
        rates.add_data(
            Reference(ws, min_col=6, max_col=8, min_row=1, max_row=month_last_row),
            titles_from_data=True,
            from_rows=False,
        )
        # Keep only 不良率 and 返工率; column G is a batch count, not a rate.
        del rates.series[1]
        rates.set_categories(categories)
        embed_chart_data(
            rates,
            months,
            [
                [ws.cell(row, 6).value for row in range(2, month_last_row + 1)],
                [ws.cell(row, 8).value for row in range(2, month_last_row + 1)],
            ],
            ["不良率", "返工率"],
            ["0.00%", "0.00%"],
        )
        ws.add_chart(rates, "J18")
    return ws


def extend_legacy_form(ws, row_style, row_height, data_rows=0):
    """Keep the original form visually continuous after the last saved record."""
    last_row = max(LEGACY_FORM_ROWS + 1, data_rows + 1)
    for row in range(2, last_row + 1):
        ws.row_dimensions[row].height = row_height
        if row <= data_rows + 1:
            continue
        for col, style in enumerate(row_style, 1):
            cell = ws.cell(row, col)
            cell._style = copy(style)
            cell.value = None


def extend_legacy_print_area(ws, record_count):
    """Extend only the exported copy; preserve the template's print columns."""
    last = max(record_count + 1, 2)
    areas = list(PrintArea.from_string(ws.print_area).ranges) if ws.print_area else [CellRange("B1:K2")]
    for area in areas:
        area.min_row = 1
        area.max_row = max(area.max_row, last)
    ws.print_area = [str(area) for area in areas]
    if not ws.print_title_rows:
        ws.print_title_rows = "1:1"


def label_legacy_chart_scope(ws, start, end):
    """Change title text only, retaining native title and chart formatting."""
    if start == date(start.year, 1, 1) and end == date(start.year, 12, 31):
        period = "全年"
    elif start == end:
        period = "当日"
    elif (
        (start.year, start.month) == (end.year, end.month)
        and start.day == 1
        and end.day == monthrange(end.year, end.month)[1]
    ):
        period = "本月"
    else:
        period = "区间"
    ws["B2"] = f"{period}出货抽检不良统计表"
    ws["B32"] = "截止周抽检不良统计表"
    # L32 is Monday; its Thursday determines the ISO week-year.
    ws["A32"] = '=YEAR(L32+3)&"-W"&TEXT(WEEKNUM(L32,21),"00")'
    ws["A61"] = (
        "项目统计为出现批次；不良率=不良件数/抽检件数。逐项已知数量见标准报表。"
        "截止周仅统计本次导出明细内、结束日期所在的周一至周日，不代表全年周度汇总。"
    )
    titles = {
        "$F$4:$F$11": f"{period}返工批次",
        "$H$4:$K$4": f"{period}抽检状况统计",
        "$B$4:$B$27": f"{period}不良项目批次",
        "$H$34:$K$34": "截止周抽检状况统计",
        "$B$34:$B$57": "截止周不良项目批次",
        "$F$34:$F$41": "截止周返工批次",
    }
    for chart in ws._charts:
        if not chart.series:
            continue
        series = chart.series[0]
        ref = series.val.numRef if series.val is not None else None
        title = next((value for key, value in titles.items() if ref and ref.f.endswith(key)), None)
        if title is None:
            continue
        rich = chart.title.tx.rich if chart.title and chart.title.tx else None
        runs = [run for paragraph in rich.p for run in paragraph.r] if rich else []
        if runs:
            runs[0].t = title
            for run in runs[1:]:
                run.t = ""
        else:
            chart.title = title


def improve_legacy_sheet_display(ws):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.page_margins.header = 0.2
    ws.page_margins.footer = 0.2
    ws.print_options.horizontalCentered = True
    ws.print_title_rows = "1:1"
    widths = {
        "B": 26,
        "C": 12,
        "D": 20,
        "E": 14,
        "F": 14,
        "G": 14,
        "H": 18,
        "I": 12,
        "J": 18,
        "K": 18,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = max(ws.column_dimensions[column].width or 0, width)
    ws.row_dimensions[1].height = max(ws.row_dimensions[1].height or 18, 30)
    for cell in ws[1]:
        alignment = copy(cell.alignment)
        alignment.wrap_text = True
        alignment.horizontal = "center"
        alignment.vertical = "center"
        cell.alignment = alignment


def improve_chart_labels(chart):
    formulas = []
    for series in chart.series:
        category = getattr(series, "cat", None)
        if category is None:
            continue
        for name in ("strRef", "numRef", "multiLvlStrRef"):
            ref = getattr(category, name, None)
            if ref is not None and getattr(ref, "f", None):
                formulas.append(ref.f.replace("'", ""))
    long_axis = any(
        token in formula
        for formula in formulas
        for token in ("$A$4:$A$27", "$A$34:$A$57")
    )
    if long_axis and getattr(chart, "x_axis", None) is not None:
        chart.x_axis.txPr = RichText(
            bodyPr=RichTextProperties(rot=-2700000),
            p=[
                Paragraph(
                    pPr=ParagraphProperties(
                        defRPr=CharacterProperties(sz=800, latin=DrawingFont(typeface="Microsoft YaHei"))
                    )
                )
            ],
        )


def improve_analysis_display(ws):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.page_margins.header = 0.2
    ws.page_margins.footer = 0.2
    ws.print_options.horizontalCentered = True
    ws.print_area = "A1:M61"
    for column, width in {
        "A": 22,
        "B": 18,
        "C": 12,
        "E": 14,
        "F": 16,
        "G": 14,
        "H": 14,
        "I": 14,
        "J": 14,
        "K": 13,
    }.items():
        ws.column_dimensions[column].width = max(ws.column_dimensions[column].width or 0, width)
    for row in (2, 32):
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 18, 30)
    for row in (3, 33):
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 18, 24)
    ws.row_dimensions[61].height = max(ws.row_dimensions[61].height or 18, 40)
    alignment = copy(ws["A61"].alignment)
    alignment.wrap_text = True
    alignment.vertical = "top"
    ws["A61"].alignment = alignment


def standard_freeze_spec(title: str):
    return {
        "检验记录": (2, "C2"),
        "不良明细": (2, "C2"),
        "不良项目": (1, "B2"),
        "统计摘要": (0, "A2"),
        "月度统计": (1, "B2"),
    }.get(title, (0, "A2"))


def reset_workbook_views(wb, active_title, legacy=False):
    """Do not carry a template's saved scroll position into a fresh report.

    Setting freeze_panes alone leaves topLeftCell and old selections intact.
    The source template scrolled to B453 and selected I454, which conflicts
    with a newly frozen C2 pane in Excel/WPS. Rebuild one consistent view.
    """
    active_sheet = wb[active_title]
    # External templates can hide the primary record sheet. openpyxl refuses
    # to activate a hidden worksheet, so restore the exported primary sheet
    # to visible in the output without modifying the source template.
    if active_sheet.sheet_state != "visible":
        active_sheet.sheet_state = "visible"
    wb.active = active_sheet
    wb.views = [BookView(activeTab=wb.index(active_sheet))]
    for ws in wb:
        view = copy(ws.sheet_view)
        view.workbookViewId = 0
        view.view = "normal"
        view.topLeftCell = "A1"
        view.tabSelected = ws.title == active_title
        view.pane = None
        view.selection = [Selection(activeCell="A1", sqref="A1")]
        if legacy and ws.title == active_title:
            view.pane = Pane(
                xSplit=2, ySplit=1, topLeftCell="C2", activePane="bottomRight", state="frozen"
            )
            view.selection = [
                Selection(pane="topRight", activeCell="C1", sqref="C1"),
                Selection(pane="bottomLeft", activeCell="B2", sqref="B2"),
                Selection(pane="bottomRight", activeCell="C2", sqref="C2"),
            ]
        elif not legacy:
            x_split, top_left = standard_freeze_spec(ws.title)
            if x_split:
                split_column = get_column_letter(x_split + 1)
                view.pane = Pane(
                    xSplit=x_split,
                    ySplit=1,
                    topLeftCell=top_left,
                    activePane="bottomRight",
                    state="frozen",
                )
                view.selection = [
                    Selection(
                        pane="topRight",
                        activeCell=f"{split_column}1",
                        sqref=f"{split_column}1",
                    ),
                    Selection(pane="bottomLeft", activeCell="A2", sqref="A2"),
                    Selection(pane="bottomRight", activeCell=top_left, sqref=top_left),
                ]
            else:
                view.pane = Pane(
                    ySplit=1,
                    topLeftCell=top_left,
                    activePane="bottomLeft",
                    state="frozen",
                )
                view.selection = [
                    Selection(pane="bottomLeft", activeCell=top_left, sqref=top_left)
                ]
        ws.views.sheetView = [view]


def add_signature(ws, index: int, path: Path, preserve_height=False):
    if not path.is_file():
        raise ValueError(f"签名图片丢失：{path.name}，请补充后再导出")
    image = Image(BytesIO(path.read_bytes()))
    if not preserve_height:
        ws.row_dimensions[index].height = max(ws.row_dimensions[index].height or 18, 40)
    width = (ws.column_dimensions["J"].width or 13) * 7 + 5
    height = (ws.row_dimensions[index].height or ws.sheet_format.defaultRowHeight) * 4 / 3
    scale = min(max(1, width - 8) / image.width, max(1, height - 8) / image.height, 1)
    image.width, image.height = image.width * scale, image.height * scale
    image.anchor = OneCellAnchor(
        _from=AnchorMarker(
            col=9,
            row=index - 1,
            colOff=pixels_to_EMU((width - image.width) / 2),
            rowOff=pixels_to_EMU((height - image.height) / 2),
        ),
        ext=XDRPositiveSize2D(pixels_to_EMU(image.width), pixels_to_EMU(image.height)),
    )
    ws.add_image(image)


def repair_analysis(
    wb,
    start: date,
    end: date,
    record_count: int,
    teams: list[str] | None = None,
    defect_names: dict[str, str] | None = None,
    other_teams: bool = False,
):
    ws = wb["数据分析表"]
    last = max(record_count + 1, 2)
    dates = f"'成品日检表'!$B$2:$B${last}"
    codes = f"'成品日检表'!$H$2:$H${last}"
    groups = f"'成品日检表'!$C$2:$C${last}"
    judgments = f"'成品日检表'!$I$2:$I${last}"
    ws["L2"], ws["M2"] = start, end
    week_start = end - timedelta(days=end.weekday())
    ws["L32"], ws["M32"] = week_start, week_start + timedelta(days=6)
    ws["L1"], ws["M1"] = "开始日期（含）", "结束日期（含）"
    ws["L31"], ws["M31"] = "周开始（含）", "周结束（含）"
    ws["A2"] = '=IF(AND(YEAR(L2)=YEAR(M2),MONTH(L2)=MONTH(M2)),MONTH(L2),"区间")'
    ws["A32"] = "=WEEKNUM(L32,21)"
    for cell in ("L2", "M2", "L32", "M32"):
        ws[cell].number_format = "yyyy-mm-dd"
    for first, total_row, control in ((4, 28, 2), (34, 58, 32)):
        date_args = f'{dates},">="&$L${control},{dates},"<"&($M${control}+1)'
        ws.cell(control, 2, "月出货抽检不良统计表" if first == 4 else "周出货抽检不良统计表")
        for i in range(24):
            r, code = first + i, chr(97 + i)
            if defect_names is not None:
                write_text(ws.cell(r, 1), defect_names.get(code, ""))
            ws.cell(r, 2, f'=COUNTIFS({date_args},{codes},"*{code}*")')
            ws.cell(r, 3, f"=RANK(B{r},$B${first}:$B${first + 23},0)")
        ws.cell(total_row, 2, f"=SUM(B{first}:B{first + 23})")
        for i in range(8):
            r = first + i
            if teams is not None:
                write_text(ws.cell(r, 5), teams[i] if i < len(teams) else "")
            # COUNTIFS treats * and ? as wildcards and ignores case. SQLite
            # team keys do neither, and the other-team remainder must be disjoint.
            ws.cell(r, 6, (
                f'=SUMPRODUCT(--({dates}>=$L${control}),--({dates}<($M${control}+1)),'
                f'--EXACT({groups},E{r}),--({judgments}="返工"))'
            ))
        if other_teams:
            ws.cell(first + 7, 6, f'=COUNTIFS({date_args},{judgments},"返工")-SUM(F{first}:F{first + 6})')
        ws.cell(first + 8, 6, f"=SUM(F{first}:F{first + 7})")
        for target, source in (("H", "E"), ("I", "F"), ("J", "G")):
            ws[f"{target}{first}"] = (
                f"=SUMIFS('成品日检表'!${source}$2:${source}${last},{date_args})"
            )
        ws[f"K{first}"] = f"=IFERROR(J{first}/I{first},0)"
        ws[f"K{first}"].number_format = "0.00%"
    if defect_names is not None and "工具" in wb.sheetnames:
        tool = wb["工具"]
        for i in range(24):
            code = chr(97 + i)
            write_text(tool.cell(i + 2, 1), code)
            write_text(tool.cell(i + 2, 2), defect_names.get(code, ""))
    # Keep the template's visible layout untouched. L/M are helper cells used
    # only by formulas; writing values must not resize, re-align, or restyle
    # the user's source form.
    ws.data_validations.dataValidation.clear()
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)


def recalculate_com(path: Path) -> bool:
    if os.name != "nt":
        return False
    excel = workbook = None
    initialized = False
    try:
        import pythoncom
        import win32com.client

        pythoncom.CoInitialize()
        initialized = True
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.AutomationSecurity = 3
        workbook = excel.Workbooks.Open(str(path.resolve()), UpdateLinks=0, ReadOnly=False)
        excel.CalculateFullRebuild()
        workbook.Save()
        return True
    except Exception:
        logging.getLogger("qc.excel").info("Excel COM 不可用，保留 openpyxl 导出", exc_info=True)
        return False
    finally:
        # A crashed Excel process can also fail during cleanup. Keep the
        # already-generated workbook and still release the remaining resources.
        if workbook is not None:
            try:
                workbook.Close(False)
            except Exception:
                logging.getLogger("qc.excel").warning("关闭 Excel 工作簿失败", exc_info=True)
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                logging.getLogger("qc.excel").warning("退出 Excel 失败", exc_info=True)
        if initialized:
            try:
                pythoncom.CoUninitialize()
            except Exception:
                logging.getLogger("qc.excel").warning("释放 Excel 组件失败", exc_info=True)


def export_workbook(ctx, path: Path, filters, legacy=False, prefer_com=True, *, expected_count=None) -> Path:
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("导出文件必须使用 .xlsx 扩展名")
    template = Path(ctx.settings.get("template_path") or ctx.paths.template)
    if path.resolve() in {template.resolve(), ctx.paths.template.resolve()}:
        raise ValueError("导出不能覆盖模板，请选择新文件名")
    path.parent.mkdir(parents=True, exist_ok=True)
    if legacy:
        if not template.is_file():
            raise ValueError("Excel 模板不存在，请在设置中选择正确模板")
        wb = load_compatible(template)
        if not {"成品日检表", "成品日检表报表", "数据分析表", "工具"} <= set(wb.sheetnames):
            raise ValueError("模板缺少必要工作表")
        ws = wb["成品日检表"]
        write_text(ws["N1"], "数据来源")
        ws.column_dimensions["N"].hidden = True
        legacy_layout = {}
        for name in ("成品日检表", "成品日检表报表"):
            sheet = wb[name]
            legacy_layout[name] = (
                [copy(c._style) for c in sheet[2]],
                sheet.row_dimensions[2].height or sheet.sheet_format.defaultRowHeight,
            )
            sheet.delete_rows(2, max(sheet.max_row - 1, 1))
            sheet._images.clear()
            for index in list(sheet.row_dimensions):
                if index > 1:
                    del sheet.row_dimensions[index]
        row_style, row_height = legacy_layout["成品日检表"]
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "检验记录"
        ws.append(HEADERS)
        detail = wb.create_sheet("不良明细")
        detail.append(["填写ID", "编码", "名称", "已知件数（空=未知）", "备注"])
        dictionary = wb.create_sheet("不良项目")
        dictionary.append(["编码", "名称", "分类", "启用", "排序", "说明"])
    count, minimum, maximum = 0, None, None
    exported_totals = dict.fromkeys(
        ("batches", "inspection_quantity", "sampling_quantity", "defect_quantity",
         "pass_batches", "rework_batches"), 0,
    )
    monthly = {}
    exported_defect_names = {}
    legacy_period = new_chart_snapshot() if legacy else None
    legacy_weeks = {}
    for count, row in enumerate(ctx.inspections.iter_records(filters), 1):
        index = count + 1
        exported_totals["batches"] += 1
        for key in ("inspection_quantity", "sampling_quantity", "defect_quantity"):
            exported_totals[key] += row[key]
        exported_totals["pass_batches"] += int(row["judgment"] == "合格")
        exported_totals["rework_batches"] += int(row["judgment"] == "返工")
        codes = [d["code"] for d in row["defects"]]
        exported_defect_names.update({d["code"]: d["name"] for d in row["defects"]})
        if legacy and any(
            len(code) != 1 or code not in "abcdefghijklmnopqrstuvwx" for code in codes
        ):
            raise ValueError("该记录含扩展编码，请选择标准报表；旧版仅支持 a–x")
        stamp = datetime.fromisoformat(f"{row['inspection_date']}T{row['inspection_time']}")
        minimum = min(minimum, stamp.date()) if minimum else stamp.date()
        maximum = max(maximum, stamp.date()) if maximum else stamp.date()
        if legacy:
            update_chart_snapshot(legacy_period, row)
            week_start = stamp.date() - timedelta(days=stamp.date().weekday())
            update_chart_snapshot(
                legacy_weeks.setdefault(week_start, new_chart_snapshot()),
                row,
            )
        values = [
            row["inspection_no"],
            stamp,
            row["team"],
            row["work_order"],
            row["inspection_quantity"],
            row["sampling_quantity"],
            row["defect_quantity"],
            ("" if legacy else ";").join(codes),
            row["judgment"],
            "",
        ]
        if not legacy:
            month = stamp.strftime("%Y-%m")
            values += [
                row["inspector"],
                row["remark"],
                source_label(row["source"]),
                stamp.year,
                month,
                stamp.date(),
            ]
            bucket = monthly.setdefault(
                month,
                {
                    "batches": 0,
                    "inspection_quantity": 0,
                    "sampling_quantity": 0,
                    "defect_quantity": 0,
                    "rework_batches": 0,
                },
            )
            bucket["batches"] += 1
            bucket["inspection_quantity"] += row["inspection_quantity"]
            bucket["sampling_quantity"] += row["sampling_quantity"]
            bucket["defect_quantity"] += row["defect_quantity"]
            bucket["rework_batches"] += int(row["judgment"] == "返工")
        for col, value in enumerate(values, 1):
            cell = ws.cell(index, col)
            write_text(cell, value)
            if legacy and col <= len(row_style):
                cell._style = copy(row_style[col - 1])
        if not legacy:
            ws.cell(index, 2).number_format = "yyyy-mm-dd hh:mm:ss"
        if legacy:
            write_text(ws.cell(index, 14), source_label(row["source"]))
            ws.row_dimensions[index].height = row_height
        if row["signature_path"]:
            add_signature(ws, index, ctx.paths.signatures / row["signature_path"], legacy)
        if not legacy:
            for d in row["defects"]:
                detail.append(
                    [row["inspection_no"], d["code"], d["name"], d["quantity"], d["remark"]]
                )
    if expected_count is not None and count != expected_count:
        wb.close()
        raise ValueError(
            f"记录数已变化（预检 {expected_count} 条，当前 {count} 条），请刷新后重新导出；原文件未覆盖"
        )
    if legacy:
        extend_legacy_form(ws, row_style, row_height, count)
        extend_legacy_print_area(ws, count)
        duplicate = wb["成品日检表报表"]
        duplicate_style, duplicate_height = legacy_layout["成品日检表报表"]
        extend_legacy_form(duplicate, duplicate_style, duplicate_height, 0)
        # The hidden duplicate is intentionally empty to avoid two copies being imported.
        analysis_start = filters.start or minimum or date.today()
        analysis_end = filters.end or maximum or date.today()
        teams = [team["name"] for team in ctx.settings.teams()][:8]
        other_teams = any(team not in teams for team in legacy_period["rework_batches"])
        if other_teams:
            other_label = "其他组别"
            while other_label in teams[:7]:
                other_label += "（合计）"
            teams = teams[:7] + [other_label]
        defect_names = {
            item["code"]: item["name"]
            for item in ctx.defects.list()
            if item["code"] in "abcdefghijklmnopqrstuvwx"
        }
        defect_names.update(exported_defect_names)
        repair_analysis(
            wb,
            analysis_start,
            analysis_end,
            count,
            teams=teams,
            defect_names=defect_names,
            other_teams=other_teams,
        )
        for chart in wb["数据分析表"]._charts:
            clear_chart_caches(chart)
            repair_chart_ranges(chart)
        label_legacy_chart_scope(wb["数据分析表"], analysis_start, analysis_end)
        selected_week = analysis_end - timedelta(days=analysis_end.weekday())
        embed_legacy_analysis_chart_data(
            wb,
            legacy_period,
            legacy_weeks.get(selected_week, new_chart_snapshot()),
            teams,
            defect_names,
            other_teams=other_teams,
        )
    else:
        remaining_names = dict(exported_defect_names)
        for item in ctx.defects.list():
            dictionary.append([
                item["code"], remaining_names.pop(item["code"], item["name"]),
                item["category"], item["enabled"], item["sort_order"], item["description"],
            ])
        # A used item could be renamed after its record was changed externally.
        # Retain the exported code/name even if no current dictionary row has it.
        for code, name in remaining_names.items():
            dictionary.append([code, name, "", None, None, ""])
        summary = wb.create_sheet("统计摘要")
        summary.append(["指标", "数值"])
        labels = {
            "batches": "检验批次",
            "inspection_quantity": "检验数量",
            "sampling_quantity": "抽检数量",
            "defect_quantity": "不良件数",
            "defect_rate": "不良率",
            "rework_batches": "返工批次",
            "rework_rate": "返工率",
            "pass_batches": "合格批次",
            "pass_rate": "合格率",
        }
        for key, value in rates(exported_totals).items():
            summary.append([labels.get(key, key), value])
            if key.endswith("rate"):
                summary.cell(summary.max_row, 2).number_format = "0.00%"
        summary.append(["口径", "不良率=不良件数/抽检件数；项目件数未知保留空白，多缺陷可重叠。"])
        summary.append(["数据范围", source_label(filters.source)])
        summary.append(["日期范围", f"{filters.start or minimum or ''} 至 {filters.end or maximum or ''}"])
        add_monthly_analysis(
            wb,
            monthly,
            filters.start or minimum,
            filters.end or maximum,
        )
        for sheet in wb:
            for cells in sheet:
                for cell in cells:
                    if cell.data_type == "f":
                        cell.data_type = "s"
            style_table(sheet)
        apply_standard_report_identity(
            wb,
            ctx.settings.all(),
            filters,
            filters.start or minimum,
            filters.end or maximum,
        )
    wb.calculation = CalcProperties(
        calcId=191029,
        calcMode="auto",
        fullCalcOnLoad=True,
        calcOnSave=True,
        forceFullCalc=True,
    )
    # Standard reports show seconds explicitly; template exports keep the
    # template's original width and number format for visual fidelity.
    if not legacy:
        ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width, 26)
    reset_workbook_views(wb, ws.title, legacy)
    with NamedTemporaryFile(suffix=".xlsx", dir=path.parent, delete=False) as handle:
        temp = Path(handle.name)
    try:
        wb.save(temp)
        wb.close()
        if legacy and prefer_com:
            recalculate_com(temp)
        temp.replace(path)
    finally:
        temp.unlink(missing_ok=True)
    logging.getLogger("qc.excel").info("导出 %s rows=%s legacy=%s", path, count, legacy)
    return path


def clear_chart_caches(chart):
    """Keep native chart formatting while removing saved production values."""
    sources = [chart.title.tx] if chart.title and chart.title.tx else []
    for series in chart.series:
        sources.extend(
            getattr(series, name, None)
            for name in ("cat", "val", "xVal", "yVal", "bubbleSize", "tx")
        )
    for source in sources:
        for reference in ("numRef", "strRef", "multiLvlStrRef"):
            ref = getattr(source, reference, None)
            for cache in ("numCache", "strCache", "multiLvlStrCache"):
                if ref is not None and hasattr(ref, cache):
                    setattr(ref, cache, None)


def repair_chart_ranges(chart):
    """Preserve the source design without restoring its known range bugs."""
    for series in chart.series:
        for data in (series.cat, series.val):
            if data is None:
                continue
            for name in ("numRef", "strRef"):
                ref = getattr(data, name, None)
                if ref is None or not ref.f:
                    continue
                ref.f = ref.f.replace("$A$4:$A$24", "$A$4:$A$27").replace(
                    "$B$4:$B$24", "$B$4:$B$27"
                )
                if getattr(chart.anchor, "_from", None) and chart.anchor._from.row >= 30:
                    ref.f = ref.f.replace("$H$3:$K$3", "$H$33:$K$33").replace(
                        "$H$4:$K$4", "$H$34:$K$34"
                    )
    # The source's weekly chart accidentally included the ranking column as
    # quantities. A zero-defect week must not show a bar of 1 for every defect.
    rank_ranges = {"数据分析表!$C$4:$C$27", "数据分析表!$C$34:$C$57"}
    chart.series = [
        series
        for series in chart.series
        if not (
            series.val and series.val.numRef and series.val.numRef.f.replace("'", "") in rank_ranges
        )
    ]


def create_empty_template(source: Path, target: Path) -> None:
    """Developer utility: preserve the layout but remove ALL production content."""
    wb = load_compatible(source)
    for name in ("成品日检表", "成品日检表报表"):
        ws = wb[name]
        row_style = [copy(c._style) for c in ws[2]]
        row_height = ws.row_dimensions[2].height or ws.sheet_format.defaultRowHeight
        ws.delete_rows(2, max(ws.max_row - 1, 1))
        ws._images.clear()
        for index in list(ws.row_dimensions):
            if index > 1:
                del ws.row_dimensions[index]
        extend_legacy_form(ws, row_style, row_height, 0)
    analysis = wb["数据分析表"]
    analysis.delete_rows(62, max(analysis.max_row - 61, 1))
    analysis["M17"] = None  # An unused scratch calculation from the source data.
    for chart in analysis._charts:
        clear_chart_caches(chart)
        repair_chart_ranges(chart)
    repair_analysis(wb, date(2026, 1, 1), date(2026, 1, 31), 0)
    reset_workbook_views(wb, "成品日检表", legacy=True)
    wb.properties.creator = "Product-QC-Daily"
    wb.properties.lastModifiedBy = "Product-QC-Daily"
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    wb.close()
