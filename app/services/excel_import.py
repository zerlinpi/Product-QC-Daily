import hashlib
import json
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl.utils.datetime import from_excel
from pydantic import ValidationError
from sqlalchemy import select

from app.core.labels import imported_source
from app.core.schemas import InspectionInput
from app.core.validation import validation_message
from app.database.models import InspectionRecord
from app.services.excel_common import load_compatible, wps_images

ALIASES = {
    "填写ID": "inspection_no",
    "检验编号": "inspection_no",
    "时间": "datetime",
    "检验日期": "datetime",
    "组别": "team",
    "加工单号": "work_order",
    "检验数量": "inspection_quantity",
    "抽检数": "sampling_quantity",
    "抽检数量": "sampling_quantity",
    "不良数": "defect_quantity",
    "不良数量": "defect_quantity",
    "不良项目": "codes",
    "判定": "judgment",
    "图片": "signature",
    "签名": "signature",
    "检验员": "inspector",
    "备注": "remark",
    "数据来源": "source",
}


def parse_codes(value, known: set[str], legacy=True) -> list[str]:
    raw = str(value or "").strip().lower()
    if not raw:
        return []
    tokens = [v for v in re.split(r"[\s.,，;；/、]+", raw) if v]
    result = []
    for token in tokens:
        if legacy and token.isascii() and token.isalpha() and all(c in known for c in token):
            result.extend(token)
        elif not legacy and token in known:
            result.append(token)
        else:
            raise ValueError(f"无法识别不良项目：{token}；请核对原表或先维护字典")
    return list(dict.fromkeys(result))


def parse_datetime(value, epoch) -> datetime:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        result = from_excel(value, epoch)
        if isinstance(result, datetime):
            return result
    if isinstance(value, str):
        for pattern in (
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%d %H:%M",
            "%Y-%m-%d",
            "%Y/%m/%d %H:%M:%S",
            "%Y/%m/%d %H:%M",
            "%Y/%m/%d",
        ):
            try:
                return datetime.strptime(value.strip(), pattern)
            except ValueError:
                continue
        try:
            return datetime.fromisoformat(value.strip()).replace(tzinfo=None)
        except ValueError:
            pass
    raise ValueError("日期格式无法识别")


@dataclass
class ImportRow:
    row_number: int
    inspection_no: str
    raw: list
    status: str = "valid"
    message: str = ""
    data: InspectionInput | None = None
    fingerprint: str = ""
    signature: bytes | None = None


@dataclass
class ImportPreview:
    path: Path
    file_hash: str
    sheet: str
    rows: list[ImportRow] = field(default_factory=list)

    @property
    def counts(self):
        counts = Counter(r.status for r in self.rows)
        return {
            status: counts[status]
            for status in ("valid", "duplicate", "conflict", "invalid", "unrecognized")
        }


def preview_workbook(ctx, path: Path) -> ImportPreview:
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("请选择 .xlsx 文件；旧 .xls 请先另存为 .xlsx")
    wb = load_compatible(path)
    candidates = []
    for ws in wb:
        for row_number, row in enumerate(
            ws.iter_rows(min_row=1, max_row=min(ws.max_row, 10), values_only=True), 1
        ):
            mapping = {}
            for index, value in enumerate(row):
                label = str(value or "").replace("：", ":")
                if label.startswith("反馈单"):
                    continue
                key = ALIASES.get(label.split(":")[-1].strip())
                if key:
                    mapping[key] = index
            required = {
                "inspection_no",
                "datetime",
                "team",
                "work_order",
                "inspection_quantity",
                "sampling_quantity",
                "defect_quantity",
                "codes",
                "judgment",
            }
            if required <= mapping.keys():
                # Prefer what the user can actually see. A hidden legacy/cached
                # "成品日检表" must not outrank a visible current record sheet.
                score = (200 if ws.sheet_state == "visible" else 0) + (
                    100 if ws.title == "成品日检表" else 0
                )
                candidates.append((score, ws.title, row_number, mapping))
                break
    if not candidates:
        raise ValueError("未找到包含记录编号、时间、组别、工单和检验数量的记录表")
    _, title, header, mapping = max(candidates, key=lambda x: x[0])
    ws = wb[title]
    if ws.max_row > 100_010 or ws.max_column > 200:
        raise ValueError("导入单表最多 10 万行、200 列，请拆分工作簿")
    dictionaries = {d["code"]: d for d in ctx.defects.list()}
    teams = {t["name"] for t in ctx.settings.teams(enabled_only=True)}
    images = wps_images(path)
    floating = {}
    for img in ws._images:
        if hasattr(img.anchor, "_from") and img.anchor._from.col == mapping.get("signature"):
            floating[img.anchor._from.row + 1] = img._data()
    details = {}
    if "不良明细" in wb:
        for values in wb["不良明细"].iter_rows(min_row=2, values_only=True):
            if len(values) >= 5 and values[0]:
                details.setdefault(str(values[0]), {})[str(values[1])] = {
                    "quantity": values[3],
                    "remark": str(values[4] or ""),
                }
    preview = ImportPreview(path, hashlib.sha256(path.read_bytes()).hexdigest(), title)
    seen = {}
    for number, raw in enumerate(ws.iter_rows(min_row=header + 1, values_only=True), header + 1):
        if not any(
            raw[i] is not None and str(raw[i]).strip()
            for key, i in mapping.items()
            if key not in ("signature", "remark")
        ):
            continue
        fields = {key: raw[index] for key, index in mapping.items()}
        no = str(fields.get("inspection_no") or "").strip()
        item = ImportRow(number, no, list(raw))
        preview.rows.append(item)
        try:
            if not no or len(no) > 200:
                raise ValueError("记录编号为空或过长")
            timestamp = parse_datetime(fields["datetime"], wb.epoch)
            try:
                codes = parse_codes(
                    fields.get("codes"), set(dictionaries), legacy=title != "检验记录"
                )
            except ValueError as exc:
                item.status, item.message = "unrecognized", str(exc)
                continue
            defects = []
            for code in codes:
                if not dictionaries[code]["enabled"]:
                    raise ValueError(f"不良项目 {code} 已停用")
                detail = details.get(no, {}).get(code, {})
                defects.append(
                    {
                        "defect_id": dictionaries[code]["id"],
                        "quantity": detail.get("quantity"),
                        "remark": detail.get("remark", ""),
                    }
                )
            if str(fields["team"] or "").strip() not in teams:
                raise ValueError("组别不存在或已停用，请先在系统设置中添加")
            if fields.get("defect_quantity") in (None, "") and codes:
                raise ValueError("填写了不良项目但缺少不良件数，请核对原表")
            source = imported_source(fields.get("source"))
            item.data = InspectionInput(
                inspection_date=timestamp.date(),
                inspection_time=timestamp.time(),
                team=str(fields["team"] or ""),
                work_order=str(fields["work_order"] or ""),
                inspection_quantity=fields["inspection_quantity"],
                sampling_quantity=fields["sampling_quantity"],
                defect_quantity=fields.get("defect_quantity") or 0,
                judgment=str(fields.get("judgment") or "").strip(),
                inspector=str(fields.get("inspector") or "历史导入（原表未提供）"),
                remark=str(fields.get("remark") or ""),
                source=source,
                defects=defects,
            )
            formula = str(fields.get("signature") or "")
            match = re.search(r'DISPIMG\("([^"]+)"', formula)
            item.signature = images.get(match.group(1)) if match else floating.get(number)
            if match and not item.signature:
                item.message = "签名引用无法解析；记录可导入，请稍后补充签名"
            content = item.data.model_dump(mode="json", exclude={"signature_path"})
            content["signature_hash"] = (
                hashlib.sha256(item.signature).hexdigest() if item.signature else ""
            )
            item.fingerprint = hashlib.sha256(
                json.dumps(content, sort_keys=True, ensure_ascii=False).encode()
            ).hexdigest()
            if no in seen:
                item.status = "duplicate" if seen[no] == item.fingerprint else "conflict"
                item.message = (
                    "工作簿中记录编号重复"
                    if item.status == "duplicate"
                    else "同一记录编号的内容不同，请核对后再导入"
                )
            else:
                seen[no] = item.fingerprint
        except (ValueError, TypeError) as exc:
            item.status = "invalid"
            item.message = validation_message(exc) if isinstance(exc, ValidationError) else str(exc)
    with ctx.db.session() as session:
        identifiers = list(seen)
        existing = {}
        for offset in range(0, len(identifiers), 400):
            existing.update(
                session.execute(
                    select(
                        InspectionRecord.inspection_no, InspectionRecord.import_fingerprint
                    ).where(InspectionRecord.inspection_no.in_(identifiers[offset : offset + 400]))
                ).all()
            )
    # If a workbook repeats an ID with conflicting content, import neither version.
    conflicted = {r.inspection_no for r in preview.rows if r.status == "conflict"}
    for item in preview.rows:
        if item.status == "valid" and item.inspection_no in conflicted:
            item.status, item.message = "conflict", "工作簿内同一记录编号有不同内容"
        if item.status == "valid" and item.inspection_no in existing:
            item.status = (
                "duplicate" if existing[item.inspection_no] == item.fingerprint else "conflict"
            )
            item.message = (
                "数据库已存在此记录（含回收站）"
                if item.status == "duplicate"
                else "已有相同记录编号但内容不同的记录，不会覆盖"
            )
    wb.close()
    return preview
