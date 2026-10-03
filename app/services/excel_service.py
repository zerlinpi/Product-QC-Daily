import hashlib
import logging
from pathlib import Path

from openpyxl import Workbook
from sqlalchemy import select

from app.database.models import InspectionRecord
from app.services.excel_common import write_text
from app.services.excel_export import export_workbook, style_table
from app.services.excel_import import ImportPreview, preview_workbook


class ExcelService:
    def __init__(self, context):
        self.ctx = context

    def preview(self, path: Path) -> ImportPreview:
        return preview_workbook(self.ctx, path)

    def import_preview(self, preview: ImportPreview) -> int:
        if hashlib.sha256(preview.path.read_bytes()).hexdigest() != preview.file_hash:
            raise ValueError("源文件已改变，请重新生成导入预览")
        copied, count = [], 0
        try:
            with self.ctx.db.session() as session:
                for item in preview.rows:
                    if item.status != "valid" or item.data is None:
                        continue
                    existing = session.scalar(
                        select(InspectionRecord).where(
                            InspectionRecord.inspection_no == item.inspection_no
                        )
                    )
                    if existing:
                        if existing.import_fingerprint == item.fingerprint:
                            continue
                        raise ValueError("预览后记录发生变化，请重新预览；此次未导入任何记录")
                    data = item.data
                    if item.signature:
                        name, created = self.ctx.signatures.store(item.signature)
                        if created:
                            copied.append(name)
                        data = data.model_copy(update={"signature_path": name})
                    self.ctx.inspections.save_in_session(
                        session,
                        data,
                        inspection_no=item.inspection_no,
                        fingerprint=item.fingerprint,
                    )
                    count += 1
            logging.getLogger("qc.excel").info(
                "导入 sheet=%s records=%s issues=%s", preview.sheet, count, preview.counts
            )
            return count
        except Exception:
            for name in copied:
                (self.ctx.paths.signatures / name).unlink(missing_ok=True)
            raise

    def export(self, path, filters, legacy=False, prefer_com=True) -> Path:
        return export_workbook(self.ctx, path, filters, legacy, prefer_com)

    def export_issues(self, preview: ImportPreview, path: Path) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = "导入问题"
        ws.append(["工作表", "Excel行号", "填写ID", "状态", "原因", "原始内容"])
        for item in preview.rows:
            if item.status != "valid" or item.message:
                values = [
                    preview.sheet,
                    item.row_number,
                    item.inspection_no,
                    item.status,
                    item.message,
                    str(item.raw),
                ]
                ws.append(values)
                for cell in ws[ws.max_row]:
                    write_text(cell, cell.value)
        style_table(ws)
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        wb.close()
        return Path(path)
