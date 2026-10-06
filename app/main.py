"""Desktop entry point; all user data remains outside the installation."""

import argparse
import json
import logging
import os
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from PySide6.QtCore import QDate, QLockFile, Qt, QTimer
from PySide6.QtGui import QIcon
from PySide6.QtWidgets import (
    QApplication,
    QDialogButtonBox,
    QFileDialog,
    QFrame,
    QGroupBox,
    QMessageBox,
    QScrollArea,
    QStyleFactory,
)

from app import __version__
from app.core.context import AppContext
from app.core.logger import setup_logging
from app.core.paths import AppPaths, resource_path
from app.core.schemas import InspectionInput, RecordFilter
from app.services.excel_common import file_sha256
from app.services.excel_export import header_footer_text
from app.ui.common import (
    BUTTON_MIN_WIDTH,
    CONTROL_MIN_HEIGHT,
    LAYOUT_SPACING,
    PAGE_MARGINS,
    TABLE_ROW_HEIGHT,
    button,
    dialog_button_box,
    friendly_error,
    stack_layout,
)
from app.ui.dialogs.file_dialogs import ExcelSaveDialog
from app.ui.localization import configure_chinese_ui
from app.ui.main_window import MainWindow
from app.ui.styles.theme import configure_platform_style, preferred_style_name


def smoke_test(ctx: AppContext, app: QApplication, report_path: Path | None) -> int:
    """Real packaged-EXE test using an explicitly isolated data directory."""
    if not os.environ.get("QC_DATA_DIR"):
        raise ValueError("自检必须通过 QC_DATA_DIR 指定独立测试目录")
    record = ctx.inspections.save(
        InspectionInput(
            team="U1",
            work_order="PACKAGED-SMOKE",
            inspection_quantity=100,
            sampling_quantity=20,
            inspector="打包自检",
            source="demo",
        )
    )
    ctx.db.dispose()
    reopened = AppContext(ctx.paths.root)
    assert reopened.inspections.get(record["id"])["work_order"] == "PACKAGED-SMOKE"
    window = MainWindow(reopened)
    if sys.platform == "win32":
        selected_style = preferred_style_name(sys.platform, QStyleFactory.keys())
        assert selected_style and selected_style.lower() in (
            "windowsvista",
            "windows",
        ), "Windows 原生 Qt style 不可用"
    stylesheet = app.styleSheet()
    for selector in (
        "QPushButton {",
        "QListWidget {",
        "QLineEdit",
        "QComboBox",
        "QDateEdit",
        "QMenu {",
        "QTableWidget {",
        "QHeaderView::section",
        "QGroupBox {",
        "QMessageBox {",
    ):
        assert selector not in stylesheet
    controls = dialog_button_box(
        QDialogButtonBox.StandardButton.Save | QDialogButtonBox.StandardButton.Cancel,
        default=QDialogButtonBox.StandardButton.Save,
    )
    controls.setParent(window)
    assert "保存" in controls.button(QDialogButtonBox.StandardButton.Save).text()
    assert "取消" in controls.button(QDialogButtonBox.StandardButton.Cancel).text()
    assert controls.button(QDialogButtonBox.StandardButton.Save).minimumHeight() == CONTROL_MIN_HEIGHT
    assert controls.button(QDialogButtonBox.StandardButton.Cancel).minimumHeight() == CONTROL_MIN_HEIGHT
    assert controls.button(QDialogButtonBox.StandardButton.Save).minimumWidth() == BUTTON_MIN_WIDTH
    assert controls.button(QDialogButtonBox.StandardButton.Cancel).minimumWidth() == BUTTON_MIN_WIDTH
    picker_directory = ctx.paths.exports / "自检导出目录"
    picker = ExcelSaveDialog(
        window, "导出报表", picker_directory / "日检报告.xlsx", "电子表格 (*.xlsx)"
    )
    assert "位置" in picker.labelText(QFileDialog.DialogLabel.LookIn)
    assert Path(picker.directory().absolutePath()) == picker_directory
    assert picker.defaultSuffix() == "xlsx"
    controls.deleteLater()
    picker.deleteLater()
    window.show()
    assert window.navigation.count() == 7
    for index in range(7):
        window.navigate(index)
        app.processEvents()
        assert window.navigation.currentRow() == index
    for page in window.pages:
        assert page.layout.getContentsMargins() == PAGE_MARGINS
        assert page.layout.spacing() == LAYOUT_SPACING
    embedded_scrolls = [
        scroll
        for page in window.pages
        for scroll in page.findChildren(QScrollArea)
    ]
    assert embedded_scrolls
    assert all(scroll.frameShape() == QFrame.Shape.NoFrame for scroll in embedded_scrolls)
    assert all(
        scroll.horizontalScrollBarPolicy() == Qt.ScrollBarPolicy.ScrollBarAlwaysOff
        for scroll in embedded_scrolls
    )
    expected_groups = {
        1: {"检验信息", "不良项目"},
        2: {"筛选条件"},
        3: {"分析范围"},
        4: {"项目列表"},
        5: {"导入历史日检表", "导出质量报表"},
        6: {"基础设置", "组别管理", "数据维护"},
    }
    for index, titles in expected_groups.items():
        present = {group.title() for group in window.pages[index].findChildren(QGroupBox)}
        assert titles.issubset(present)

    dashboard, entry, records, analytics, defects, reports, settings = window.pages
    field_controls = [
        dashboard.source,
        entry.inspection_date,
        entry.inspection_time,
        entry.team,
        entry.work_order,
        entry.inspection_quantity,
        entry.sampling_quantity,
        entry.defect_quantity,
        entry.judgment,
        entry.inspector,
        records.start,
        records.end,
        records.team,
        records.judgment,
        records.source,
        records.defect,
        records.has_defects,
        records.search,
        records.work_order,
        records.inspector,
        analytics.preset,
        analytics.start,
        analytics.end,
        analytics.source,
        analytics.metric_choice,
        defects.search,
        reports.preset,
        reports.start,
        reports.end,
        reports.source,
        settings.theme,
        settings.retention,
        *settings.fields.values(),
        entry.defects.search,
        entry.defects.category,
        *[qty for _, _, qty in entry.defects.entries],
    ]
    assert all(control.minimumHeight() == CONTROL_MIN_HEIGHT for control in field_controls)

    settings_field_titles = {
        "公司名称",
        "工厂名称",
        "默认检验员",
        "默认组别",
        "界面主题",
        "原表模板",
        "导出目录",
        "备份目录",
        "自动备份",
        "自动备份保留",
    }
    settings_labels = {
        item.text(): item.objectName()
        for item in settings.findChildren(type(settings.title_label))
        if item.text() in settings_field_titles
    }
    assert settings_labels.keys() == settings_field_titles
    assert all(kind == "fieldLabel" for kind in settings_labels.values())
    company_label = next(
        item
        for item in settings.findChildren(type(settings.title_label))
        if item.text() == "公司名称"
    )
    assert company_label.buddy() is settings.fields["company"]
    assert settings.fields["company"].accessibleName() == "公司名称"

    density_tables = [
        (analytics.ranking, 10),
        (analytics.teams_table, 8),
        (settings.teams, 7),
        (entry.defects.table, 9),
    ]
    for widget, rows in density_tables:
        expected = CONTROL_MIN_HEIGHT + TABLE_ROW_HEIGHT * rows + widget.frameWidth() * 2
        assert widget.minimumHeight() == expected

    primary_probe = button("主要操作", primary=True)
    danger_probe = button("危险操作", danger=True)
    assert primary_probe.minimumHeight() == CONTROL_MIN_HEIGHT
    assert danger_probe.minimumHeight() == CONTROL_MIN_HEIGHT
    assert primary_probe.minimumWidth() == BUTTON_MIN_WIDTH
    assert danger_probe.minimumWidth() == BUTTON_MIN_WIDTH
    assert primary_probe.font().bold()
    assert not danger_probe.icon().isNull()
    primary_probe.deleteLater()
    danger_probe.deleteLater()
    stack_probe = QFrame()
    stack_probe_layout = stack_layout(stack_probe)
    assert stack_probe_layout.getContentsMargins() == (0, 0, 0, 0)
    assert stack_probe_layout.spacing() == LAYOUT_SPACING
    stack_probe.deleteLater()

    def grid_position(layout, widget):
        index = layout.indexOf(widget)
        assert index >= 0
        return layout.getItemPosition(index)

    window.resize(1080, 720)
    app.processEvents()
    dashboard = window.pages[0]
    window.navigate(0)
    app.processEvents()
    assert dashboard._layout_mode == "narrow"
    assert grid_position(dashboard.grid, dashboard.cards[2][0]) == (1, 0, 1, 5)
    entry = window.pages[1]
    window.navigate(1)
    app.processEvents()
    assert entry._layout_mode == "narrow"
    assert grid_position(entry.content_grid, entry.right_panel) == (1, 0, 1, 1)
    records = window.pages[2]
    window.navigate(2)
    app.processEvents()
    assert records._filter_layout_mode == "narrow"
    assert grid_position(records.filters_grid, records.defect_field) == (3, 0, 1, 2)
    analytics = window.pages[3]
    window.navigate(3)
    app.processEvents()
    assert analytics._layout_mode == "narrow"
    assert grid_position(analytics.grid, analytics.trend_frame) == (4, 0, 1, 10)
    reports = window.pages[5]
    window.navigate(5)
    app.processEvents()
    assert reports._layout_mode == "narrow"
    assert grid_position(reports.filters_grid, reports.source) == (3, 1, 1, 1)

    window.resize(1440, 920)
    app.processEvents()
    window.navigate(0)
    app.processEvents()
    assert dashboard._layout_mode == "wide"
    assert grid_position(dashboard.grid, dashboard.cards[2][0]) == (0, 4, 1, 2)
    window.navigate(1)
    app.processEvents()
    assert entry._layout_mode == "wide"
    assert grid_position(entry.content_grid, entry.right_panel) == (0, 1, 1, 1)
    window.navigate(2)
    app.processEvents()
    assert records._filter_layout_mode == "wide"
    window.navigate(3)
    app.processEvents()
    assert analytics._layout_mode == "wide"
    assert grid_position(analytics.grid, analytics.trend_frame) == (2, 0, 1, 5)
    window.navigate(5)
    app.processEvents()
    assert reports._layout_mode == "wide"
    assert grid_position(reports.filters_grid, reports.source) == (1, 3, 1, 1)
    window.navigate(4)
    app.processEvents()
    assert window.pages[4].search.isClearButtonEnabled()
    window.navigate(6)
    app.processEvents()
    assert window.pages[6].team_count.text().startswith("共 ")
    assert window.pages[0].refreshed.objectName() == "summary"
    assert window.pages[2].selection_count.objectName() == "summary"
    assert window.pages[4].count.objectName() == "summary"
    assert window.pages[5].import_status.objectName() == "summary"
    assert window.pages[6].team_count.objectName() == "summary"
    assert window.pages[1].defects.total.objectName() == "summary"
    assert window.pages[2].count.objectName() == "summary"
    assert window.pages[0].charts[0].empty.objectName() == "empty"
    assert (
        window.pages[2].table.horizontalHeaderItem(4).textAlignment()
        & Qt.AlignmentFlag.AlignRight
    )
    assert (
        window.pages[4].table.horizontalHeaderItem(3).textAlignment()
        & Qt.AlignmentFlag.AlignHCenter
    )
    window.refresh_theme()

    reports = window.pages[5]
    reports.preset.setCurrentText("自定义")
    reports.start.setDate(QDate(2026, 10, 2))
    reports.end.setDate(QDate(2026, 10, 1))
    assert not reports.original_export.isEnabled()
    assert not reports.detailed_export.isEnabled()
    assert reports.export_scope.text() == "日期范围无效：开始日期不能晚于结束日期"
    reports.end.setDate(QDate(2026, 10, 2))
    assert reports.original_export.isEnabled()
    assert reports.detailed_export.isEnabled()

    analytics = window.pages[3]
    analytics.preset.setCurrentText("自定义")
    analytics.start.setDate(QDate(2026, 10, 2))
    analytics.end.setDate(QDate(2026, 10, 1))
    assert not analytics.analyze_button.isEnabled()
    assert "日期范围无效" in analytics.scope.text()

    records = window.pages[2]
    records.range_enabled.setChecked(True)
    records.start.setDate(QDate(2026, 10, 2))
    records.end.setDate(QDate(2026, 10, 1))
    assert not records.query_button.isEnabled()
    records.refresh()
    assert "上一次查询结果" in records.count.text()
    records.range_enabled.setChecked(False)
    records.source.setCurrentText("演示数据")
    assert records.filters_dirty
    assert "筛选条件尚未应用" in records.action_buttons["export"].toolTip()
    records.refresh()
    assert not records.filters_dirty
    assert records.applied_filters.source == "demo"

    assert header_footer_text("A" * 63 + "&TRAILING", 64).endswith("&&")

    output = reopened.excel.export(
        ctx.paths.exports / "smoke.xlsx", RecordFilter(source="demo"), legacy=True, prefer_com=False
    )
    assert output.exists()
    workbook = load_workbook(output)
    sheet = workbook["成品日检表"]
    assert workbook.active == sheet
    assert sheet.sheet_view.topLeftCell == "A1"
    assert sheet.freeze_panes == "C2"
    assert sheet.sheet_view.selection[-1].activeCell == "C2"
    workbook.close()
    preview = reopened.excel.preview(output)
    assert preview.file_hash == file_sha256(output)
    year = date.today().year
    reopened.settings.update({"company": "自检公司", "factory": "一厂"})
    standard = reopened.excel.export(
        ctx.paths.exports / "smoke-standard.xlsx",
        RecordFilter(
            start=date(year, 1, 1),
            end=date(year, 12, 31),
            source="demo",
        ),
        legacy=False,
        prefer_com=False,
    )
    workbook = load_workbook(standard)
    assert workbook["检验记录"]["N1"].value == "月份"
    assert workbook["检验记录"].auto_filter.ref.endswith(
        f"N{workbook['检验记录'].max_row}"
    )
    monthly = workbook["月度统计"]
    assert monthly.max_row == 14
    assert monthly["A14"].value == "合计"
    assert len(monthly._charts) == 2
    assert workbook["检验记录"].page_setup.orientation == "landscape"
    assert workbook.properties.creator == "自检公司 · 一厂"
    assert "自检公司" in workbook["检验记录"].oddHeader.left.text
    assert "&P" in workbook["检验记录"].oddFooter.center.text
    assert workbook["检验记录"].freeze_panes == "C2"
    assert workbook["检验记录"].sheet_view.zoomScale == 85
    assert monthly.auto_filter.ref == f"A1:H{monthly.max_row - 1}"
    assert len(monthly.conditional_formatting) == 7
    assert not workbook["检验记录"].sheet_view.showGridLines
    workbook.close()
    backup = reopened.backup.backup()
    reopened.backup.restore(backup)
    window.close()
    reopened.db.dispose()
    result = {
        "ok": True,
        "version": __version__,
        "frozen": bool(getattr(sys, "frozen", False)),
        "pages": 7,
        "database_persistence": True,
        "template_export": True,
        "backup_restore": True,
        "chinese_controls": True,
        "export_view_reset": True,
        "export_path_dialog": True,
        "annual_standard_export": True,
        "native_windows_ui": True,
        "adaptive_native_layout": True,
        "native_utility_pages": True,
        "unified_native_ui": True,
        "component_ui_consistency": True,
        "full_ui_style_consistency": True,
    }
    if report_path:
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(result, indent=2), encoding="utf-8")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--self-test", action="store_true")
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    if args.self_test:
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    app = QApplication(sys.argv[:1])
    configure_chinese_ui(app)
    app.setApplicationName("Product-QC-Daily")
    app.setApplicationVersion(__version__)
    app.setOrganizationName("Product-QC-Daily")
    configure_platform_style(app)
    app.setWindowIcon(QIcon(str(resource_path("assets/icons/app.svg"))))
    ctx = None
    try:
        paths = AppPaths()
        setup_logging(paths.logs)
        lock = QLockFile(str(paths.root / "application.lock"))
        lock.setStaleLockTime(0)
        if not lock.tryLock(100):
            if args.self_test:
                raise ValueError("同一数据目录已有正在运行的软件实例")
            QMessageBox.information(
                None, "软件已运行", "此数据目录已有日检软件在运行，请返回已打开的窗口。"
            )
            return 1
        ctx = AppContext(paths.root)
        logging.getLogger("qc").info("启动 v%s", __version__)
        if args.self_test:
            return smoke_test(ctx, app, args.report)
        window = MainWindow(ctx)
        sys.excepthook = lambda kind, error, traceback: friendly_error(window, error)
        window.show()
        QTimer.singleShot(
            200,
            lambda: window.run_job(
                "检查每日备份",
                ctx.backup.daily_backup,
                lambda path: window.notify("每日备份已完成" if path else "本地数据库已就绪"),
            ),
        )
        return app.exec()
    except Exception as error:
        if args.self_test:
            logging.getLogger("qc").exception("自检失败")
            if args.report:
                args.report.parent.mkdir(parents=True, exist_ok=True)
                args.report.write_text(
                    json.dumps({"ok": False, "error": str(error)}), encoding="utf-8"
                )
        else:
            friendly_error(None, error)
        return 1
    finally:
        if ctx:
            ctx.db.dispose()


if __name__ == "__main__":
    raise SystemExit(main())
