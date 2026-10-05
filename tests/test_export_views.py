from copy import copy

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.views import Pane, Selection

from app.core.schemas import RecordFilter
from app.services.excel_export import LEGACY_FORM_ROWS, create_empty_template


def assert_clean_views(wb, active_title, legacy):
    assert wb.active.title == active_title
    assert len(wb.views) == 1
    assert wb.views[0].showVerticalScroll and wb.views[0].showHorizontalScroll
    assert wb.views[0].showSheetTabs
    for ws in wb:
        assert len(ws.views.sheetView) == 1
        view = ws.sheet_view
        assert view.topLeftCell == "A1"
        assert view.view == "normal"
        assert view.tabSelected == (ws.title == active_title)
        assert view.workbookViewId == 0
        if legacy and ws.title == active_title:
            assert ws.freeze_panes == "C2"
            assert (view.pane.xSplit, view.pane.ySplit) == (2, 1)
            assert view.pane.activePane == "bottomRight"
            assert [(s.pane, s.activeCell, s.sqref) for s in view.selection] == [
                ("topRight", "C1", "C1"),
                ("bottomLeft", "B2", "B2"),
                ("bottomRight", "C2", "C2"),
            ]
        elif legacy:
            assert view.pane is None
            assert [(s.pane, s.activeCell, s.sqref) for s in view.selection] == [
                (None, "A1", "A1")
            ]
        elif ws.title == active_title:
            assert ws.freeze_panes == "C2"
            assert (view.pane.xSplit, view.pane.ySplit) == (2, 1)
            assert view.pane.activePane == "bottomRight"
            assert [(s.pane, s.activeCell, s.sqref) for s in view.selection] == [
                ("topRight", "C1", "C1"),
                ("bottomLeft", "A2", "A2"),
                ("bottomRight", "C2", "C2"),
            ]
        else:
            assert ws.freeze_panes == "A2"
            assert [(s.pane, s.activeCell, s.sqref) for s in view.selection] == [
                ("bottomLeft", "A2", "A2")
            ]


@pytest.mark.parametrize("legacy", [True, False])
@pytest.mark.parametrize("count", [0, 1, 505])
def test_exports_start_at_top_with_consistent_frozen_panes(ctx, payload, tmp_path, legacy, count):
    for _ in range(count):
        ctx.inspections.save(payload)
    output = ctx.excel.export(
        tmp_path / "report.xlsx", RecordFilter(), legacy=legacy, prefer_com=False
    )
    wb = load_workbook(output)
    title = "成品日检表" if legacy else "检验记录"
    expected_rows = max(count + 1, LEGACY_FORM_ROWS + 1) if legacy else count + 1
    assert wb[title].max_row == expected_rows
    assert_clean_views(wb, title, legacy)
    wb.close()


@pytest.mark.parametrize("operation", ["prepare_template", "export"])
def test_external_template_old_scroll_split_and_window_states_are_not_reused(
    ctx, payload, tmp_path, operation
):
    source, target = tmp_path / "external.xlsx", tmp_path / "result.xlsx"
    wb = load_workbook(ctx.paths.template)
    wb.views[0].showVerticalScroll = False
    wb.views[0].showHorizontalScroll = False
    wb.views[0].showSheetTabs = False
    wb.views.append(copy(wb.views[0]))
    for ws in wb:
        view = ws.sheet_view
        view.topLeftCell = "B453"
        view.view = "pageBreakPreview"
        view.tabSelected = True
        view.pane = Pane(ySplit=453, topLeftCell="B454", state="frozenSplit")
        view.selection = [Selection(activeCell="I454", sqref="I454:I455")]
        second = copy(view)
        second.workbookViewId = 1
        ws.views.sheetView.append(second)
    wb.save(source)
    wb.close()
    original = source.read_bytes()
    if operation == "prepare_template":
        create_empty_template(source, target)
    else:
        ctx.inspections.save(payload)
        ctx.settings.update({"template_path": str(source)})
        ctx.excel.export(target, RecordFilter(), legacy=True, prefer_com=False)
    result = load_workbook(target)
    assert_clean_views(result, "成品日检表", True)
    assert len(result["数据分析表"]._charts) == 6
    assert source.read_bytes() == original
    result.close()


def test_bundled_template_has_no_saved_production_scroll_state(ctx):
    wb = load_workbook(ctx.paths.template)
    assert_clean_views(wb, "成品日检表", True)
    wb.close()



@pytest.mark.parametrize("operation", ["prepare_template", "export"])
def test_external_template_hidden_primary_sheet_is_restored_and_source_is_unchanged(
    ctx, payload, tmp_path, operation
):
    source, target = tmp_path / "hidden-primary.xlsx", tmp_path / "result.xlsx"
    wb = load_workbook(ctx.paths.template)
    wb["成品日检表"].sheet_state = "hidden"
    wb.active = wb["数据分析表"]
    wb.save(source)
    wb.close()
    original = source.read_bytes()

    if operation == "prepare_template":
        create_empty_template(source, target)
    else:
        ctx.inspections.save(payload)
        ctx.settings.update({"template_path": str(source)})
        ctx.excel.export(target, RecordFilter(), legacy=True, prefer_com=False)

    result = load_workbook(target)
    assert result["成品日检表"].sheet_state == "visible"
    assert_clean_views(result, "成品日检表", True)
    assert source.read_bytes() == original
    result.close()
