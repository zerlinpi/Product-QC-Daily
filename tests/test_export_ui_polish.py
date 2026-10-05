from threading import Event

from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication

from app.core.schemas import RecordFilter
from app.services.excel_export import LEGACY_FORM_ROWS
from app.ui.main_window import MainWindow


def test_legacy_export_keeps_full_blank_form_after_last_record(ctx, payload, tmp_path):
    ctx.inspections.save(payload)
    output = ctx.excel.export(
        tmp_path / "continuous-form.xlsx", RecordFilter(), legacy=True, prefer_com=False
    )
    wb = load_workbook(output)
    ws = wb["成品日检表"]
    assert ws.max_row == LEGACY_FORM_ROWS + 1
    assert ws.row_dimensions[50].height == ws.row_dimensions[2].height
    assert ws["B50"].value is None and ws["J50"].value is None
    assert ws["B50"].font == ws["B2"].font
    assert ws["B50"].fill == ws["B2"].fill
    assert ws["B50"].border == ws["B2"].border
    assert ws["B50"].alignment == ws["B2"].alignment
    assert ws["B2"].number_format == "yyyy-mm-dd hh:mm:ss"
    assert ws["J50"].style_id == ws["J2"].style_id
    assert ws.column_dimensions["B"].width >= 26
    assert ws.column_dimensions["D"].width >= 20
    assert ws.row_dimensions[1].height >= 30
    assert all(cell.alignment.wrap_text for cell in ws[1])
    wb.close()


def test_analysis_sheet_uses_readable_widths_and_wrapped_note(ctx, payload, tmp_path):
    ctx.inspections.save(payload)
    output = ctx.excel.export(
        tmp_path / "analysis-layout.xlsx", RecordFilter(), legacy=True, prefer_com=False
    )
    wb = load_workbook(output)
    ws = wb["数据分析表"]
    assert ws.column_dimensions["A"].width >= 22
    assert ws.column_dimensions["B"].width >= 18
    assert ws.row_dimensions[2].height >= 30
    assert ws.row_dimensions[32].height >= 30
    assert ws["A61"].alignment.wrap_text
    assert ws.row_dimensions[61].height >= 40
    wb.close()


def test_export_job_uses_modal_animation_instead_of_statusbar_progress(ctx, qtbot):
    window = MainWindow(ctx)
    qtbot.addWidget(window)
    window.show()
    started, release = Event(), Event()

    def work():
        started.set()
        assert release.wait(5)
        return "done.xlsx"

    try:
        window.run_job("导出测试报表", work)
        qtbot.waitUntil(started.is_set, timeout=5000)
        qtbot.waitUntil(lambda: window._job_dialog is not None, timeout=2000)
        assert window._job_dialog.objectName() == "taskProgressDialog"
        assert window._job_dialog.isVisible()
        assert not window.progress.isVisible()
        assert QApplication.activeModalWidget() is window._job_dialog
        release.set()
        qtbot.waitUntil(lambda: window._job is None, timeout=10000)
        assert window._job_dialog is None
    finally:
        release.set()
        qtbot.waitUntil(lambda: window._job is None, timeout=10000)


def test_export_completion_and_dropdowns_are_visibly_customized(ctx, qtbot, tmp_path):
    window = MainWindow(ctx)
    qtbot.addWidget(window)
    window.show()
    result = tmp_path / "完成.xlsx"
    window.export_completed(result)
    qtbot.waitUntil(lambda: window._export_message is not None, timeout=2000)
    assert window._export_message.windowTitle() == "导出完成"
    assert window._export_message.text() == "导出完成"
    assert str(result) in window._export_message.informativeText()
    stylesheet = QApplication.instance().styleSheet()
    assert "QComboBox::drop-down" in stylesheet
    assert "QComboBox QAbstractItemView::item" in stylesheet
    assert "QDateEdit::drop-down" in stylesheet
    window._export_message.close()
