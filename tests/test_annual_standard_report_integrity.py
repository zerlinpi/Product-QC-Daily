"""针对用户年度报表只有少量正式记录的回归：全年演示须完整落在明细及图表。"""

from collections import Counter
from datetime import date, datetime

from openpyxl import load_workbook

from app.core.schemas import RecordFilter


def test_1825_demo_standard_export_charts_match_every_exported_record(
    ctx, payload, tmp_path
):
    start, end = date(2026, 1, 1), date(2026, 12, 31)
    ctx.inspections.save(payload.model_copy(update={"inspection_date": date(2026, 10, 3)}))
    assert ctx.demo.generate(1825, start, end, ["U1", "U2", "U3"], seed=2026) == 1825
    filters = RecordFilter(start=start, end=end, source="demo", descending=False)
    target = ctx.excel.export(tmp_path / "demo-year.xlsx", filters, legacy=False)
    wb = load_workbook(target)
    try:
        rows = list(wb["检验记录"].values)
        header, data = rows[0], rows[1:]
        assert len(data) == 1825
        assert header[-3:] == ("年份", "月份", "日期")
        assert all(row[12] == "演示数据" and row[13] == 2026 for row in data)
        assert len({row[15] for row in data}) == 365
        first, last = data[0][15], data[-1][15]
        assert (first.date() if isinstance(first, datetime) else first) == start
        assert (last.date() if isinstance(last, datetime) else last) == end
        month_counts = Counter(row[14] for row in data)
        months = wb["月度统计"]
        assert [months.cell(i, 1).value for i in range(2, 14)] == [
            f"2026-{m:02d}" for m in range(1, 13)
        ]
        assert [
            months.cell(i, 2).value for i in range(2, 14)
        ] == [month_counts[f"2026-{m:02d}"] for m in range(1, 13)]
        assert sum(month_counts.values()) == 1825
        for column, index in [(3, 4), (4, 5), (5, 6)]:
            assert months.cell(14, column).value == sum(row[index] for row in data)
        assert dict(wb["统计摘要"].values)["检验批次"] == 1825
        assert dict(wb["统计摘要"].values)["数据范围"] == "演示数据"
        assert len(months._charts) == 2
        chart_batches, chart_rates = months._charts
        assert [point.v for point in chart_batches.series[0].val.numRef.numCache.pt] == [
            months.cell(i, 2).value for i in range(2, 14)
        ]
        for series, column in zip(chart_rates.series, [6, 8], strict=True):
            assert [point.v for point in series.val.numRef.numCache.pt] == [
                months.cell(i, column).value for i in range(2, 14)
            ]
    finally:
        wb.close()
