from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

HEADERS = [
    "填写ID",
    "时间",
    "检验:组别",
    "检验:加工单号",
    "检验:检验数量",
    "检验:抽检数",
    "检验:不良数",
    "检验:不良项目",
    "检验:判定",
    "签名:图片",
]


def workbook(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "成品日检表"
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def row(identifier="old-1", code="eg", defect=2):
    return [
        identifier,
        datetime(2023, 12, 31, 15, 30),
        "U1",
        "MO-001",
        100,
        20,
        defect,
        code,
        "合格",
    ]


def test_split_codes_without_inventing_quantities():
    from app.services.excel_import import parse_codes

    assert parse_codes("D.e.K.", set("abcdefghijklmnopqrstuvwx")) == ["d", "e", "k"]
    for value in ["α", "y", "1", "e.槽变形"]:
        with pytest.raises(ValueError):
            parse_codes(value, set("abcdefghijklmnopqrstuvwx"))


def test_import_preview_duplicates_conflicts_invalid_and_unknown(ctx, tmp_path):
    from app.core.schemas import RecordFilter

    path = workbook(
        tmp_path / "input.xlsx",
        [row(), row("old-2", "", None), row("bad", "y"), row("bad2", "e", 25), row()],
    )
    preview = ctx.excel.preview(path)
    assert preview.counts == {
        "valid": 2,
        "duplicate": 1,
        "conflict": 0,
        "invalid": 1,
        "unrecognized": 1,
    }
    assert ctx.inspections.query(RecordFilter())[1] == 0
    assert ctx.excel.import_preview(preview) == 2
    assert ctx.inspections.query(RecordFilter())[1] == 2
    imported = ctx.inspections.query(RecordFilter())[0]
    assert any(d["code"] == "g" and d["quantity"] is None for r in imported for d in r["defects"])
    assert ctx.excel.import_preview(ctx.excel.preview(path)) == 0
    changed = workbook(tmp_path / "changed.xlsx", [row(defect=3)])
    assert ctx.excel.preview(changed).counts["conflict"] == 1
    report = ctx.excel.export_issues(preview, tmp_path / "issues.xlsx")
    issue_wb = load_workbook(report)
    issue_ws = issue_wb.active
    assert issue_ws.max_row == 4
    statuses = {issue_ws.cell(row, 4).value for row in range(2, issue_ws.max_row + 1)}
    assert statuses == {"重复", "异常", "无法识别"}
    assert not statuses & {"duplicate", "invalid", "unrecognized"}
    issue_wb.close()


def test_header_detection_prefers_visible_main(ctx, tmp_path):
    path = workbook(tmp_path / "input.xlsx", [row()])
    wb = load_workbook(path)
    other = wb.copy_worksheet(wb.active)
    other.title = "成品日检表报表"
    other.sheet_state = "hidden"
    wb.save(path)
    assert ctx.excel.preview(path).sheet == "成品日检表"
    assert ctx.excel.preview(path).counts["valid"] == 1


def test_standard_export_roundtrip_details_and_formula_text(ctx, payload, tmp_path):
    from app.core.schemas import InspectionInput, RecordFilter

    ctx.inspections.save(
        InspectionInput(
            **(
                payload.model_dump()
                | {
                    "work_order": "=1+1",
                    "defect_quantity": 2,
                    "defects": [
                        {"defect_id": 1, "quantity": 2},
                        {"defect_id": 24, "quantity": None},
                    ],
                }
            )
        )
    )
    path = ctx.excel.export(tmp_path / "standard.xlsx", RecordFilter(), legacy=False)
    wb = load_workbook(path)
    assert wb["检验记录"]["D2"].value == "=1+1"
    assert wb["检验记录"]["D2"].data_type == "s"
    assert wb["不良明细"].max_row == 3
    from app.core.context import AppContext

    other = AppContext(tmp_path / "other")
    preview = other.excel.preview(path)
    assert preview.counts["valid"] == 1
    other.excel.import_preview(preview)
    saved = other.inspections.query(RecordFilter())[0][0]
    assert [d["quantity"] for d in saved["defects"]] == [2, None]
    other.db.dispose()


def test_legacy_template_preserves_sheets_charts_and_correct_formulas(ctx, payload, tmp_path):
    from app.core.schemas import RecordFilter

    ctx.inspections.save(payload)
    path = ctx.excel.export(tmp_path / "legacy.xlsx", RecordFilter(), legacy=True, prefer_com=False)
    wb = load_workbook(path)
    assert wb.sheetnames[:4] == ["成品日检表报表", "成品日检表", "数据分析表", "工具"]
    analysis = wb["数据分析表"]
    assert len(analysis._charts) == 6
    monthly_totals = analysis._charts[1].series[0]
    weekly_totals = analysis._charts[3].series[0]
    assert [point.v for point in monthly_totals.cat.strRef.strCache.pt] == [
        "检验数量",
        "抽检数",
        "不良数",
        "不良率",
    ]
    assert [point.v for point in monthly_totals.val.numRef.numCache.pt][:3] == [
        payload.inspection_quantity,
        payload.sampling_quantity,
        payload.defect_quantity,
    ]
    assert [point.v for point in weekly_totals.val.numRef.numCache.pt][:3] == [
        payload.inspection_quantity,
        payload.sampling_quantity,
        payload.defect_quantity,
    ]
    assert "SUMIFS" in analysis["J4"].value
    assert '"*x*"' in wb["数据分析表"]["B27"].value
    assert "$B$4:$B$27" in wb["数据分析表"]["C4"].value
    assert "B34:B57" in wb["数据分析表"]["B58"].value
    assert wb["成品日检表"]["A2"].value.startswith("QC-")


def test_import_preview_changed_file_is_not_silently_used(ctx, tmp_path):
    path = workbook(tmp_path / "input.xlsx", [row()])
    preview = ctx.excel.preview(path)
    workbook(path, [row("new")])
    with pytest.raises(ValueError, match="改变"):
        ctx.excel.import_preview(preview)


def test_legacy_codes_remain_concatenated_after_adding_ambiguous_code(ctx, tmp_path):
    ctx.defects.save({"code": "eg", "name": "扩展项目"})
    path = workbook(tmp_path / "legacy-input.xlsx", [row()])
    preview = ctx.excel.preview(path)
    assert preview.counts["valid"] == 1
    assert [d.defect_id for d in preview.rows[0].data.defects] == [5, 7]


def test_legacy_demo_export_cannot_turn_into_production(ctx, payload, tmp_path):
    from app.core.context import AppContext
    from app.core.schemas import InspectionInput, RecordFilter

    ctx.inspections.save(InspectionInput(**(payload.model_dump() | {"source": "demo"})))
    path = ctx.excel.export(
        tmp_path / "demo-legacy.xlsx", RecordFilter(source="demo"), legacy=True, prefer_com=False
    )
    other = AppContext(tmp_path / "isolated")
    try:
        other.excel.import_preview(other.excel.preview(path))
        assert other.inspections.query(RecordFilter())[1] == 0
        assert other.inspections.query(RecordFilter(source="demo"))[1] == 1
    finally:
        other.db.dispose()


@pytest.mark.parametrize("legacy", [False, True])
def test_export_full_timestamp_fits_and_stays_a_real_date(ctx, payload, tmp_path, legacy):
    from app.core.schemas import InspectionInput, RecordFilter

    stamp = datetime(2026, 12, 31, 23, 59, 59)
    ctx.inspections.save(
        InspectionInput(
            **(
                payload.model_dump()
                | {
                    "inspection_date": stamp.date(),
                    "inspection_time": stamp.time(),
                }
            )
        )
    )
    path = ctx.excel.export(tmp_path / "date.xlsx", RecordFilter(), legacy=legacy, prefer_com=False)
    wb = load_workbook(path)
    ws = wb["成品日检表" if legacy else "检验记录"]
    assert ws["B2"].value == stamp
    assert ws["B2"].is_date
    if legacy:
        template = load_workbook(ctx.paths.template)
        template_ws = template["成品日检表"]
        assert ws["B2"].number_format == template_ws["B2"].number_format
        assert ws.column_dimensions["B"].width == template_ws.column_dimensions["B"].width
        template.close()
    else:
        assert ws["B2"].number_format == "yyyy-mm-dd hh:mm:ss"
        assert ws.column_dimensions["B"].width >= 24
    wb.close()


@pytest.mark.parametrize("legacy", [False, True])
def test_chinese_source_exports_and_demo_roundtrip(ctx, payload, tmp_path, legacy):
    from app.core.context import AppContext
    from app.core.schemas import InspectionInput, RecordFilter

    for source in ("manual", "excel", "demo"):
        ctx.inspections.save(InspectionInput(**(payload.model_dump() | {"source": source})))
    path = ctx.excel.export(
        tmp_path / "sources.xlsx", RecordFilter(source="all"), legacy=legacy, prefer_com=False
    )
    wb = load_workbook(path)
    ws = wb["成品日检表" if legacy else "检验记录"]
    column = 14 if legacy else 13
    assert {ws.cell(r, column).value for r in range(2, 5)} == {"手动录入", "表格导入", "演示数据"}
    if not legacy:
        assert dict(wb["统计摘要"].values)["数据范围"] == "全部数据"
    wb.close()
    other = AppContext(tmp_path / "other-source")
    try:
        assert other.excel.import_preview(other.excel.preview(path)) == 3
        assert other.inspections.query(RecordFilter(source="demo"))[1] == 1
        assert other.inspections.query(RecordFilter())[1] == 2
    finally:
        other.db.dispose()


def test_legacy_layout_keeps_original_chart_positions_and_signature_cell(ctx, payload, tmp_path):
    from PIL import Image

    from app.core.schemas import InspectionInput, RecordFilter

    signature = tmp_path / "signature.png"
    Image.new("RGB", (400, 100), "white").save(signature)
    ctx.inspections.save(
        InspectionInput(**(payload.model_dump() | {"signature_path": str(signature)}))
    )
    path = ctx.excel.export(tmp_path / "layout.xlsx", RecordFilter(), legacy=True, prefer_com=False)
    wb = load_workbook(path)
    ws, analysis = wb["成品日检表"], wb["数据分析表"]
    assert [(c.anchor._from.col, c.anchor._from.row) for c in analysis._charts] == [
        (4, 12),
        (7, 4),
        (4, 17),
        (7, 34),
        (3, 49),
        (3, 42),
    ]
    template = load_workbook(ctx.paths.template)
    template_analysis = template["数据分析表"]
    assert analysis.sheet_view.showGridLines == template_analysis.sheet_view.showGridLines
    assert analysis.sheet_view.zoomScale == template_analysis.sheet_view.zoomScale
    assert analysis.page_setup.orientation == template_analysis.page_setup.orientation
    assert analysis.page_setup.fitToWidth == template_analysis.page_setup.fitToWidth
    assert analysis.print_options.horizontalCentered == template_analysis.print_options.horizontalCentered
    assert analysis.print_area == template_analysis.print_area
    assert analysis._charts[2].series[0].val.numRef.f.endswith("$B$4:$B$27")
    assert analysis._charts[3].series[0].val.numRef.f.endswith("$H$34:$K$34")
    assert len(analysis._charts[4].series) == 1
    assert analysis._charts[4].series[0].val.numRef.f.endswith("$B$34:$B$57")
    assert ws.column_dimensions["A"].hidden
    assert ws.row_dimensions[2].height == pytest.approx(34.45)
    assert ws["C2"].font.name == "微软雅黑"
    template_ws = template["成品日检表"]
    assert ws.sheet_view.showGridLines == template_ws.sheet_view.showGridLines
    assert ws.sheet_view.zoomScale == template_ws.sheet_view.zoomScale
    assert ws.page_setup.orientation == template_ws.page_setup.orientation
    assert ws.page_setup.fitToWidth == template_ws.page_setup.fitToWidth
    assert ws.print_title_rows == "$1:$1"
    anchor = ws._images[0].anchor
    assert (anchor._from.col, anchor._from.row) == (9, 1)
    assert (anchor._from.colOff + anchor.ext.cx) / 9525 <= ws.column_dimensions["J"].width * 7 + 5
    assert (anchor._from.rowOff + anchor.ext.cy) / 9525 <= ws.row_dimensions[2].height * 4 / 3
    template.close()
    wb.close()


@pytest.mark.parametrize(
    "value,expected",
    [("demo", "demo"), ("演示数据", "demo"), ("manual", "excel"), ("表格导入", "excel")],
)
def test_source_import_accepts_old_and_chinese_labels(ctx, tmp_path, value, expected):
    path = workbook(tmp_path / "source-input.xlsx", [row()])
    wb = load_workbook(path)
    wb.active["N1"], wb.active["N2"] = "数据来源", value
    wb.save(path)
    preview = ctx.excel.preview(path)
    assert preview.counts["valid"] == 1
    assert preview.rows[0].data.source == expected


def test_template_keeps_chart_design_and_removes_production_caches(ctx, payload, tmp_path):
    from copy import copy
    from zipfile import ZipFile

    from openpyxl.chart.data_source import NumData, NumVal

    from app.core.schemas import RecordFilter
    from app.services.excel_export import create_empty_template

    source, target = tmp_path / "private-template.xlsx", tmp_path / "blank-template.xlsx"
    wb = load_workbook(ctx.paths.template)
    wb["成品日检表"]["A2"] = "PRIVATE-PRODUCTION-RECORD"
    chart = wb["数据分析表"]._charts[0]
    chart.series[0].graphicalProperties.solidFill = "00AA44"
    chart.series[0].val.numRef.numCache = NumData(ptCount=1, pt=[NumVal(idx=0, v=123456789)])
    original_anchor = (copy(chart.anchor._from), copy(chart.anchor.to))
    wb.save(source)
    wb.close()
    before = source.read_bytes()
    create_empty_template(source, target)
    cleaned = load_workbook(target)
    assert (
        cleaned["数据分析表"]._charts[0].series[0].graphicalProperties.solidFill.srgbClr == "00AA44"
    )
    saved_anchor = cleaned["数据分析表"]._charts[0].anchor
    assert (saved_anchor._from, saved_anchor.to) == original_anchor
    assert cleaned["成品日检表"]["A2"].value is None
    assert source.read_bytes() == before
    with ZipFile(target) as archive:
        assert not any(name.startswith("xl/media/") for name in archive.namelist())
        for name in archive.namelist():
            if name.startswith("xl/charts/"):
                content = archive.read(name)
                assert b"123456789" not in content
                assert b"numCache" not in content
    cleaned.close()
    # Selecting an external populated template must not reuse its old chart values either.
    ctx.inspections.save(payload)
    ctx.settings.update({"template_path": str(source)})
    output = ctx.excel.export(
        tmp_path / "external.xlsx", RecordFilter(), legacy=True, prefer_com=False
    )
    with ZipFile(output) as archive:
        assert b"123456789" not in archive.read("xl/charts/chart1.xml")


@pytest.mark.parametrize("operation", ["prepare_template", "export"])
def test_external_template_quantity_chart_excludes_rank(ctx, payload, tmp_path, operation):
    from copy import deepcopy

    from app.core.schemas import RecordFilter
    from app.services.excel_export import create_empty_template

    source, target = tmp_path / "external-template.xlsx", tmp_path / "result.xlsx"
    wb = load_workbook(ctx.paths.template)
    chart = wb["数据分析表"]._charts[4]
    chart.series = chart.series[:1]
    rank_series = deepcopy(chart.series[0])
    rank_series.val.numRef.f = "'数据分析表'!$C$34:$C$57"
    chart.series.append(rank_series)
    anchor = deepcopy(chart.anchor)
    wb.save(source)
    wb.close()
    source_bytes = source.read_bytes()
    if operation == "prepare_template":
        create_empty_template(source, target)
    else:
        ctx.inspections.save(payload)
        ctx.settings.update({"template_path": str(source)})
        ctx.excel.export(target, RecordFilter(), legacy=True, prefer_com=False)
    result = load_workbook(target)
    chart = result["数据分析表"]._charts[4]
    assert len(chart.series) == 1
    assert chart.series[0].val.numRef.f.endswith("$B$34:$B$57")
    assert (chart.anchor._from, chart.anchor.to) == (anchor._from, anchor.to)
    assert source.read_bytes() == source_bytes
    result.close()



def test_legacy_export_syncs_current_team_and_defect_names(ctx, tmp_path):
    from app.core.schemas import RecordFilter

    team = next(item for item in ctx.settings.teams() if item["name"] == "U1")
    ctx.settings.save_team("一组", team["id"], enabled=True, sort_order=team["sort_order"])
    defect = next(item for item in ctx.defects.list() if item["code"] == "a")
    ctx.defects.save(defect | {"name": "端子包角（新名称）"}, defect["id"])

    path = ctx.excel.export(
        tmp_path / "renamed-labels.xlsx",
        RecordFilter(),
        legacy=True,
        prefer_com=False,
    )
    wb = load_workbook(path)
    analysis, tool = wb["数据分析表"], wb["工具"]
    assert analysis["A4"].value == "端子包角（新名称）"
    assert analysis["A34"].value == "端子包角（新名称）"
    assert analysis["E4"].value == "一组"
    assert analysis["E34"].value == "一组"
    assert tool["A2"].value == "a"
    assert tool["B2"].value == "端子包角（新名称）"
    wb.close()



def test_annual_demo_standard_export_supports_month_filter_and_charts(ctx, tmp_path):
    from datetime import date
    from zipfile import ZipFile

    from app.core.schemas import RecordFilter

    ctx.demo.generate(
        240,
        date(2026, 1, 1),
        date(2026, 12, 31),
        ["U1"],
        rework_rate=0.08,
        defect_rate=0.02,
        seed=42,
    )
    path = ctx.excel.export(
        tmp_path / "annual-demo.xlsx",
        RecordFilter(
            start=date(2026, 1, 1),
            end=date(2026, 12, 31),
            source="demo",
        ),
        legacy=False,
    )
    wb = load_workbook(path)
    records = wb["检验记录"]
    headers = [cell.value for cell in records[1]]
    assert headers[13:16] == ["年份", "月份", "日期"]
    assert headers[16] == "不良项目编码"
    assert records.column_dimensions["Q"].hidden
    assert records.auto_filter.ref == f"A1:Q{records.max_row}"
    years = {records.cell(row, 14).value for row in range(2, records.max_row + 1)}
    months = {records.cell(row, 15).value for row in range(2, records.max_row + 1)}
    dates = {records.cell(row, 16).value for row in range(2, records.max_row + 1)}
    assert years == {2026}
    assert months == {f"2026-{month:02d}" for month in range(1, 13)}
    assert all(value.year == 2026 for value in dates)
    assert records["P2"].number_format == "yyyy-mm-dd"

    monthly = wb["月度统计"]
    assert [monthly.cell(row, 1).value for row in range(2, 14)] == [
        f"2026-{month:02d}" for month in range(1, 13)
    ]
    assert sum(monthly.cell(row, 2).value for row in range(2, 14)) == 240
    assert monthly["F2"].number_format == "0.00%"
    assert monthly["H2"].number_format == "0.00%"
    assert len(monthly._charts) == 2

    volume, rates = monthly._charts
    assert volume.x_axis.tickLblSkip == 1
    assert rates.x_axis.tickLblSkip == 1
    assert volume.series[0].cat.strRef is not None
    assert [point.v for point in volume.series[0].cat.strRef.strCache.pt] == [
        f"2026-{month:02d}" for month in range(1, 13)
    ]
    assert [point.v for point in volume.series[0].val.numRef.numCache.pt] == [
        monthly.cell(row, 2).value for row in range(2, 14)
    ]
    assert [point.v for point in rates.series[0].val.numRef.numCache.pt] == [
        monthly.cell(row, 6).value for row in range(2, 14)
    ]
    assert [point.v for point in rates.series[1].val.numRef.numCache.pt] == [
        monthly.cell(row, 8).value for row in range(2, 14)
    ]
    assert wb.calculation.calcMode == "auto"
    assert wb.calculation.fullCalcOnLoad
    assert wb.calculation.forceFullCalc
    assert dict(wb["统计摘要"].values)["数据范围"] == "演示数据"
    wb.close()

    with ZipFile(path) as archive:
        chart_xml = b"".join(
            archive.read(name)
            for name in archive.namelist()
            if name.startswith("xl/charts/chart") and name.endswith(".xml")
        )
        assert b"strCache" in chart_xml
        assert b"numCache" in chart_xml
        assert b"2026-01" in chart_xml
        assert b"2026-12" in chart_xml
        workbook_xml = archive.read("xl/workbook.xml")
        assert b'calcMode="auto"' in workbook_xml
        assert b'fullCalcOnLoad="1"' in workbook_xml
        assert b'forceFullCalc="1"' in workbook_xml



def test_multi_year_monthly_charts_reduce_axis_label_density():
    from datetime import date

    from openpyxl import Workbook

    from app.services.excel_export import add_monthly_analysis

    wb = Workbook()
    wb.remove(wb.active)
    monthly = add_monthly_analysis(
        wb,
        {},
        date(2024, 1, 1),
        date(2026, 12, 31),
    )
    volume, rates = monthly._charts
    assert monthly.max_row == 38
    assert volume.x_axis.tickLblSkip == 2
    assert rates.x_axis.tickLblSkip == 2
    wb.close()


def test_wrapped_row_height_respects_explicit_line_breaks():
    from app.services.excel_export import wrapped_row_height

    single = wrapped_row_height(("第一行", 40))
    multiline = wrapped_row_height(("第一行\n第二行\n第三行", 40))
    assert single == 22
    assert multiline >= 51


def test_standard_export_is_print_ready_and_visually_grouped(ctx, payload, tmp_path):
    from app.core.schemas import InspectionInput, RecordFilter

    ctx.inspections.save(
        InspectionInput(
            **(
                payload.model_dump()
                | {
                    "work_order": "WO-LONG-QUALITY-REPORT-001",
                    "remark": "返工原因：尺寸偏差，已复检并记录处理结果；该备注用于验证导出后长文本能够完整换行显示而不是被固定行高截断。",
                    "judgment": "返工",
                    "defect_quantity": 2,
                    "defects": [
                        {
                            "defect_id": 1,
                            "quantity": 2,
                            "remark": "不良位置较长，需要在导出明细中自动增加行高并保持完整可见。",
                        }
                    ],
                }
            )
        )
    )
    path = ctx.excel.export(
        tmp_path / "styled-standard.xlsx",
        RecordFilter(start=payload.inspection_date, end=payload.inspection_date, source="production"),
        legacy=False,
        prefer_com=False,
    )
    wb = load_workbook(path)
    records = wb["检验记录"]
    assert records.page_setup.orientation == "landscape"
    assert not records.sheet_view.showGridLines
    assert records.auto_filter.ref == f"A1:Q{records.max_row}"
    assert records.column_dimensions["D"].width >= 22
    assert records.column_dimensions["L"].width >= 34
    assert records["L2"].alignment.wrap_text
    assert records["I2"].fill.fgColor.rgb.endswith("FCE4D6")
    assert records["I2"].font.bold
    assert records.freeze_panes == "C2"
    assert records.sheet_view.zoomScale == 85
    assert records.print_options.horizontalCentered
    assert records.page_margins.header == 0.2
    assert records["A2"].alignment.wrap_text
    assert records["D2"].alignment.wrap_text
    assert records["E2"].number_format == "#,##0"
    assert records["E2"].alignment.horizontal == "right"
    assert records["G2"].alignment.horizontal == "right"
    assert records["G2"].fill.fgColor.rgb.endswith("FCE4D6")
    assert records["G2"].font.bold
    assert records["A1"].border.bottom.style == "medium"
    assert records.row_dimensions[2].height > 22
    assert records.print_area

    detail = wb["不良明细"]
    assert detail.row_dimensions[2].height > 22
    assert detail.print_area
    assert detail.freeze_panes == "C2"
    assert detail["A2"].alignment.wrap_text
    assert detail["D2"].alignment.horizontal == "right"
    assert detail["D2"].number_format == "#,##0"

    dictionary = wb["不良项目"]
    assert dictionary.print_area
    assert dictionary.freeze_panes == "B2"
    assert dictionary["E2"].alignment.horizontal == "right"
    assert dictionary["E2"].number_format == "#,##0"

    summary = wb["统计摘要"]
    assert summary.page_setup.orientation == "portrait"
    assert not summary.sheet_view.showGridLines
    assert summary.column_dimensions["B"].width >= 48
    assert summary.sheet_view.zoomScale == 100
    assert summary["B2"].number_format == "#,##0"
    assert summary.print_area
    assert summary.freeze_panes == "A2"
    note_rows = {
        summary.cell(row, 1).value: row for row in range(2, summary.max_row + 1)
    }
    assert summary.row_dimensions[note_rows["口径"]].height > 22
    assert summary.cell(note_rows["口径"], 2).alignment.horizontal == "left"
    metric_rows = {
        summary.cell(row, 1).value: row for row in range(2, summary.max_row + 1)
    }
    assert summary.cell(metric_rows["不良率"], 2).fill.fgColor.rgb.endswith("FCE4D6")
    assert summary.cell(metric_rows["不良率"], 2).font.color.rgb.endswith("C65911")
    assert summary.cell(metric_rows["合格率"], 2).fill.fgColor.rgb.endswith("E2F0D9")
    assert summary.cell(metric_rows["合格率"], 2).font.color.rgb.endswith("375623")
    assert summary.cell(metric_rows["检验数量"], 2).alignment.horizontal == "right"

    monthly = wb["月度统计"]
    assert monthly["A2"].value
    assert monthly.cell(monthly.max_row, 1).value == "合计"
    assert monthly.cell(monthly.max_row, 6).number_format == "0.00%"
    assert monthly.cell(monthly.max_row, 8).number_format == "0.00%"
    assert len(monthly._charts) == 2
    assert len(monthly.conditional_formatting) == 7
    assert monthly.auto_filter.ref == f"A1:H{monthly.max_row - 1}"
    assert monthly.sheet_view.zoomScale == 85
    assert monthly.freeze_panes == "B2"
    assert monthly["B2"].number_format == "#,##0"
    assert monthly["G2"].number_format == "#,##0"
    assert monthly["B2"].alignment.horizontal == "right"
    assert monthly["F2"].alignment.horizontal == "right"
    assert monthly.cell(monthly.max_row, 1).border.top.style == "medium"
    assert monthly.print_area
    assert monthly.row_dimensions[monthly.max_row].height >= 24
    wb.close()



def test_standard_export_print_identity_and_summary_hierarchy(ctx, payload, tmp_path):
    from app.core.schemas import RecordFilter

    ctx.settings.update({"company": "A&B Quality", "factory": "东莞一厂"})
    ctx.inspections.save(payload)
    path = ctx.excel.export(
        tmp_path / "identity-report.xlsx",
        RecordFilter(
            start=payload.inspection_date,
            end=payload.inspection_date,
            source="production",
        ),
        legacy=False,
        prefer_com=False,
    )
    wb = load_workbook(path)
    assert wb.properties.creator == "A&B Quality · 东莞一厂"
    assert wb.properties.lastModifiedBy == "Product-QC-Daily"
    assert "正式数据" in wb.properties.subject
    assert str(payload.inspection_date) in wb.properties.title

    for ws in wb:
        assert "A&&B Quality" in ws.oddHeader.left.text
        assert ws.title in ws.oddHeader.center.text
        assert str(payload.inspection_date) in ws.oddHeader.right.text
        assert ws.oddFooter.left.text == "正式数据"
        assert "&P" in ws.oddFooter.center.text
        assert "&N" in ws.oddFooter.center.text
        assert ws.oddFooter.right.text == "&F"

    summary = wb["统计摘要"]
    assert summary["A2"].fill.fgColor.rgb.endswith("D9EAF7")
    note_rows = {
        summary.cell(row, 1).value: row for row in range(2, summary.max_row + 1)
    }
    for label in ("口径", "数据范围", "日期范围"):
        row = note_rows[label]
        assert summary.cell(row, 1).fill.fgColor.rgb.endswith("F2F2F2")
        assert summary.cell(row, 2).fill.fgColor.rgb.endswith("F2F2F2")
        assert not summary.cell(row, 2).font.bold
    wb.close()



def test_header_footer_truncation_never_splits_ampersand_escape():
    from app.services.excel_export import header_footer_text

    value = "A" * 63 + "&" + "TRAILING"
    rendered = header_footer_text(value, 64)
    assert rendered == "A" * 63 + "&&"
    assert rendered.replace("&&", "&") == value[:64]



def test_import_preview_rejects_source_changed_during_preview(ctx, tmp_path, monkeypatch):
    from app.core.schemas import RecordFilter
    from app.services import excel_import

    path = workbook(tmp_path / "changing.xlsx", [row()])

    def mutate_source(source):
        changed = load_workbook(source)
        changed.active["D2"] = "CHANGED-DURING-PREVIEW"
        changed.save(source)
        changed.close()
        return {}

    monkeypatch.setattr(excel_import, "wps_images", mutate_source)
    with pytest.raises(ValueError, match="预览过程中发生变化"):
        ctx.excel.preview(path)
    assert ctx.inspections.query(RecordFilter())[1] == 0


def test_import_preview_closes_workbook_on_early_validation_error(ctx, tmp_path, monkeypatch):
    from app.services import excel_import

    path = tmp_path / "no-header.xlsx"
    path.write_bytes(b"snapshot")
    closed = []

    class EmptyWorkbook:
        def __iter__(self):
            return iter(())

        def close(self):
            closed.append(True)

    monkeypatch.setattr(excel_import, "load_compatible", lambda _: EmptyWorkbook())
    with pytest.raises(ValueError, match="未找到包含记录编号"):
        ctx.excel.preview(path)
    assert closed == [True]



def test_import_hashing_is_streamed_without_path_read_bytes(ctx, tmp_path, monkeypatch):
    from pathlib import Path

    from app.core.schemas import RecordFilter

    path = workbook(tmp_path / "stream-hash.xlsx", [row("STREAM-1")])
    real_read_bytes = Path.read_bytes

    def reject_whole_file_read(self):
        if self == path:
            raise AssertionError("导入源文件不应使用 Path.read_bytes 整体读入内存")
        return real_read_bytes(self)

    monkeypatch.setattr(Path, "read_bytes", reject_whole_file_read)
    preview = ctx.excel.preview(path)
    assert preview.counts["valid"] == 1
    assert ctx.excel.import_preview(preview) == 1
    assert ctx.inspections.query(RecordFilter())[1] == 1
