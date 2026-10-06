import pytest
from PySide6.QtCore import Qt


def test_ui_entry_save_and_reopen_without_duplicate(ctx, qtbot):
    from app.ui.main_window import MainWindow

    window = MainWindow(ctx)
    qtbot.addWidget(window)
    window.show()
    window.navigate(1)
    entry = window.pages[1]
    entry.work_order.setText("UI-100")
    entry.inspector.setText("王工")
    entry.inspection_quantity.setValue(100)
    entry.sampling_quantity.setValue(20)
    entry.save_record()
    saved_id = entry.record_id
    assert saved_id
    entry.work_order.setText("UI-101")
    entry.save_record()
    from app.core.schemas import RecordFilter

    assert ctx.inspections.query(RecordFilter())[1] == 1
    assert ctx.inspections.get(saved_id)["work_order"] == "UI-101"
    entry.save_record(new=True)
    assert entry.record_id is None
    assert entry.team.currentText() == "U1"
    assert entry.inspector.text() == "王工"
    assert ctx.inspections.query(RecordFilter())[1] == 1
    # Open again against the same persistent context.
    window.close()
    again = MainWindow(ctx)
    qtbot.addWidget(again)
    again.navigate(2)
    assert again.pages[2].table.rowCount() == 1


def test_navigation_all_pages_empty_database(ctx, qtbot):
    from app.ui.main_window import MainWindow

    window = MainWindow(ctx)
    qtbot.addWidget(window)
    window.show()
    for i in range(7):
        window.navigate(i)
        assert window.stack.currentIndex() == i
        if i == 1:
            assert not window.pages[1].dirty


def test_copy_selection_creates_new_record_and_preserves_original(ctx, payload, qtbot):
    from app.core.schemas import RecordFilter
    from app.ui.main_window import MainWindow

    saved = ctx.inspections.save(payload)
    window = MainWindow(ctx)
    qtbot.addWidget(window)
    window.open_record(saved["id"], copy_record=True)
    entry = window.pages[1]
    assert entry.record_id is None
    entry.work_order.setText("COPY-2")
    entry.save_record()
    assert ctx.inspections.query(RecordFilter())[1] == 2
    assert ctx.inspections.get(saved["id"])["work_order"] == "MO-001"


def test_defect_selector_counts_and_unknown(ctx, qtbot):
    from app.ui.widgets.defect_selector import DefectSelector

    widget = DefectSelector(ctx)
    qtbot.addWidget(widget)
    widget.set_values([{"defect_id": 5, "quantity": 2}, {"defect_id": 7, "quantity": None}])
    values = widget.values()
    assert {(v["defect_id"], v["quantity"]) for v in values} == {(5, 2), (7, None)}


def test_keyboard_save_and_new(ctx, qtbot):
    from app.ui.main_window import MainWindow

    window = MainWindow(ctx)
    qtbot.addWidget(window)
    window.show()
    window.navigate(1)
    entry = window.pages[1]
    entry.work_order.setText("KEYBOARD")
    entry.inspector.setText("工人")
    entry.inspection_quantity.setValue(100)
    entry.sampling_quantity.setValue(20)
    entry.work_order.setFocus()
    window.activateWindow()
    qtbot.waitUntil(window.isActiveWindow)
    qtbot.wait(50)
    qtbot.keyClick(entry.work_order, Qt.Key.Key_S, Qt.KeyboardModifier.ControlModifier)
    qtbot.waitUntil(lambda: entry.record_id is not None)
    qtbot.keyClick(entry.work_order, Qt.Key.Key_N, Qt.KeyboardModifier.ControlModifier)
    assert entry.record_id is None


def test_confirmed_discard_restores_saved_record(ctx, payload, qtbot, monkeypatch):
    from app.ui.main_window import MainWindow

    saved = ctx.inspections.save(payload)
    window = MainWindow(ctx)
    qtbot.addWidget(window)
    window.open_record(saved["id"])
    entry = window.pages[1]
    entry.work_order.setText("ABANDONED")
    monkeypatch.setattr("app.ui.pages.inspection_page.confirm", lambda *a, **kw: True)
    window.navigate(2)
    window.navigate(1)
    assert entry.work_order.text() == "MO-001"
    entry.save_record()
    assert ctx.inspections.get(saved["id"])["work_order"] == "MO-001"


def test_editor_preserves_per_defect_remarks(ctx, payload, qtbot):
    from app.core.schemas import InspectionInput
    from app.ui.main_window import MainWindow

    saved = ctx.inspections.save(
        InspectionInput(
            **(
                payload.model_dump()
                | {
                    "defect_quantity": 1,
                    "defects": [{"defect_id": 5, "quantity": 1, "remark": "已返修确认"}],
                }
            )
        )
    )
    window = MainWindow(ctx)
    qtbot.addWidget(window)
    window.open_record(saved["id"])
    window.pages[1].defects.reload()
    window.pages[1].save_record()
    assert ctx.inspections.get(saved["id"])["defects"][0]["remark"] == "已返修确认"


def test_required_popup_names_every_missing_field_in_chinese(ctx, qtbot, monkeypatch):
    import app.ui.common as common
    from app.core.schemas import RecordFilter
    from app.ui.main_window import MainWindow

    window = MainWindow(ctx)
    qtbot.addWidget(window, before_close_func=lambda w: setattr(w.pages[1], "dirty", False))
    window.navigate(1)
    messages = []

    class FakeMessageBox:
        def exec(self):
            return None

    monkeypatch.setattr(
        common,
        "message_box",
        lambda parent, title, text, **kwargs: (
            messages.append(text) or FakeMessageBox()
        ),
    )
    entry = window.pages[1]
    entry.team.setCurrentIndex(-1)
    entry.work_order.setText("  ")
    entry.inspector.clear()
    entry.inspection_quantity.setValue(0)
    entry.sampling_quantity.setValue(0)
    entry.save_record()
    assert len(messages) == 1
    for name in ("组别", "加工单号", "检验员", "检验数量", "抽检数量"):
        assert name in messages[0]
    assert "大于 0" in messages[0]
    assert "String should" not in messages[0]
    assert "Input should" not in messages[0]
    assert ctx.inspections.query(RecordFilter())[1] == 0


def test_record_source_cells_use_chinese(ctx, payload, qtbot):
    from app.core.schemas import InspectionInput
    from app.ui.main_window import MainWindow

    for source in ("manual", "excel", "demo"):
        ctx.inspections.save(InspectionInput(**(payload.model_dump() | {"source": source})))
    window = MainWindow(ctx)
    qtbot.addWidget(window)
    page = window.pages[2]
    page.source.setCurrentIndex(2)
    page.refresh()
    assert {page.table.item(r, 9).text() for r in range(3)} == {"手动录入", "表格导入", "演示数据"}


def test_record_export_defaults_to_original_layout(ctx, payload, qtbot, monkeypatch, tmp_path):
    from openpyxl import load_workbook

    from app.ui.dialogs import file_dialogs
    from app.ui.main_window import MainWindow

    ctx.inspections.save(payload)
    window = MainWindow(ctx)
    qtbot.addWidget(window)
    page = window.pages[2]
    page.refresh()
    output = tmp_path / "ui-export.xlsx"
    monkeypatch.setattr(file_dialogs, "save_excel", lambda *args: (str(output), ""))
    monkeypatch.setattr(window, "run_job", lambda title, work, done: done(work()))
    page.export()
    wb = load_workbook(output)
    assert "成品日检表" in wb.sheetnames
    assert "数据分析表" in wb.sheetnames
    wb.close()



def test_report_import_summary_uses_chinese_status_labels(ctx, qtbot):
    from types import SimpleNamespace

    from app.ui.main_window import MainWindow

    window = MainWindow(ctx)
    qtbot.addWidget(window)
    page = window.pages[5]
    preview = SimpleNamespace(
        counts={
            "valid": 2,
            "duplicate": 1,
            "conflict": 1,
            "invalid": 1,
            "unrecognized": 1,
        }
    )
    page.import_done(2, preview)
    summary = page.import_status.text()
    for status in ("正常", "重复", "编号冲突", "异常", "无法识别"):
        assert status in summary
    for internal in ("valid", "duplicate", "conflict", "invalid", "unrecognized"):
        assert internal not in summary



def test_records_filters_reset_and_selection_feedback(ctx, payload, qtbot):
    from app.ui.main_window import MainWindow

    ctx.inspections.save(payload)
    window = MainWindow(ctx)
    qtbot.addWidget(window)
    window.navigate(2)
    page = window.pages[2]

    assert not page.start.isEnabled()
    assert not page.end.isEnabled()
    page.range_enabled.setChecked(True)
    assert page.start.isEnabled()
    assert page.end.isEnabled()

    page.search.setText("MO-001")
    page.judgment.setCurrentIndex(1)
    page.source.setCurrentIndex(2)
    page.table.selectRow(0)
    assert page.selection_count.text() == "已选择 1 条"
    assert page.action_buttons["edit"].isEnabled()
    assert page.action_buttons["delete"].isEnabled()
    assert not page.action_buttons["restore"].isEnabled()

    page.reset_filters()
    assert not page.range_enabled.isChecked()
    assert not page.start.isEnabled()
    assert page.search.text() == ""
    assert page.judgment.currentIndex() == 0
    assert page.source.currentIndex() == 0
    assert page.selection_count.text() == "未选择记录"


def test_analysis_reports_and_dashboard_show_scope_feedback(ctx, qtbot):
    from app.ui.main_window import MainWindow

    window = MainWindow(ctx)
    qtbot.addWidget(window)

    dashboard = window.pages[0]
    dashboard.refresh()
    assert dashboard.refreshed.text().startswith("更新于 ")

    window.navigate(3)
    analytics = window.pages[3]
    assert analytics.scope.text().startswith("当前范围：")
    assert analytics.source.currentText() in analytics.scope.text()

    window.navigate(5)
    reports = window.pages[5]
    assert reports.export_scope.text().startswith("将导出：")
    assert reports.source.currentText() in reports.export_scope.text()



def test_reports_invalid_custom_range_disables_export_before_save_dialog(
    ctx, qtbot, monkeypatch
):
    from PySide6.QtCore import QDate

    from app.ui.dialogs import file_dialogs
    from app.ui.main_window import MainWindow

    window = MainWindow(ctx)
    qtbot.addWidget(window)
    page = window.pages[5]
    page.preset.setCurrentText("自定义")
    page.start.setDate(QDate(2026, 10, 2))
    page.end.setDate(QDate(2026, 10, 1))

    assert page.export_scope.text() == "日期范围无效：开始日期不能晚于结束日期"
    assert not page.original_export.isEnabled()
    assert not page.detailed_export.isEnabled()

    monkeypatch.setattr(
        file_dialogs,
        "save_excel",
        lambda *args, **kwargs: pytest.fail("无效日期范围不应打开保存窗口"),
    )
    page.export(False)

    page.end.setDate(QDate(2026, 10, 2))
    assert page.export_scope.text().startswith("将导出：")
    assert page.original_export.isEnabled()
    assert page.detailed_export.isEnabled()
