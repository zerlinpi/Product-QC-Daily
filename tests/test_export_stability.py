"""Reproductions for export snapshot, period labels and print truncation."""

from datetime import date

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.core.schemas import DefectInput, RecordFilter


def seed_records(ctx, payload, count):
    with ctx.db.session() as session:
        for index in range(count):
            ctx.inspections.save_in_session(
                session,
                payload,
                inspection_no=f"SNAP-{index:05d}",
                validate_references=False,
                flush=False,
            )


def test_stream_keeps_defects_in_same_snapshot_across_500_row_boundary(ctx, payload):
    data = payload.model_copy(
        update={
            "defect_quantity": 1,
            "defects": [DefectInput(defect_id=1, quantity=1)],
        }
    )
    seed_records(ctx, data, 501)
    iterator = ctx.inspections.iter_records(RecordFilter(descending=False))
    first = next(iterator)
    with ctx.db.session() as session:
        session.execute(text("UPDATE inspection_defects SET quantity=NULL"))
    exported = [first, *iterator]
    assert len(exported) == 501
    assert all(row["defects"][0]["quantity"] == 1 for row in exported)
    assert ctx.inspections.get(exported[-1]["id"])["defects"][0]["quantity"] is None


def test_summary_uses_exported_rows_when_database_changes(ctx, payload, monkeypatch, tmp_path):
    ctx.inspections.save(payload)
    original = ctx.inspections.iter_records

    def changing_records(filters):
        yield from original(filters)
        ctx.inspections.save(payload)

    monkeypatch.setattr(ctx.inspections, "iter_records", changing_records)
    wb = load_workbook(ctx.excel.export(tmp_path / "snapshot.xlsx", RecordFilter()))
    try:
        assert wb["检验记录"].max_row == 2
        assert dict(wb["统计摘要"].values)["检验批次"] == 1
        assert dict(wb["统计摘要"].values)["检验数量"] == payload.inspection_quantity
        assert wb["月度统计"]["B2"].value == 1
    finally:
        wb.close()


@pytest.mark.parametrize("fixed_area", [False, True])
def test_annual_print_area_covers_last_record_and_repeats_header(
    ctx, payload, tmp_path, fixed_area
):
    seed_records(ctx, payload, 501)
    template = tmp_path / "template.xlsx"
    wb = load_workbook(ctx.paths.template)
    if fixed_area:
        wb["成品日检表"].print_area = "B1:K501"
    wb.save(template)
    wb.close()
    original = template.read_bytes()
    ctx.settings.update({"template_path": str(template)})
    result = load_workbook(
        ctx.excel.export(
            tmp_path / "annual.xlsx",
            RecordFilter(),
            legacy=True,
            prefer_com=False,
        )
    )
    try:
        assert result["成品日检表"].print_area == "'成品日检表'!$B$1:$K$502"
        assert result["成品日检表"].print_title_rows == "$1:$1"
        assert template.read_bytes() == original
    finally:
        result.close()


@pytest.mark.parametrize(
    "start,end,iso_week",
    [
        (date(2026, 1, 1), date(2026, 12, 31), "2026-W53"),
        (date(2020, 12, 30), date(2021, 1, 1), "2020-W53"),
        (date(2024, 2, 1), date(2024, 2, 29), "2024-W09"),
        (date(2026, 1, 5), date(2026, 1, 6), "2026-W02"),
    ],
)
def test_six_chart_titles_describe_period_and_only_final_iso_week(
    ctx, tmp_path, start, end, iso_week
):
    wb = load_workbook(
        ctx.excel.export(
            tmp_path / "scope.xlsx",
            RecordFilter(start=start, end=end),
            legacy=True,
            prefer_com=False,
        )
    )
    try:
        analysis = wb["数据分析表"]
        titles = ["".join(chart.title.to_tree().itertext()) for chart in analysis._charts]
        assert len(titles) == 6
        assert all("所选期间" in title for title in titles[:3])
        assert all("截止周" in title for title in titles[3:])
        # Use the ISO Thursday even if a user edits L32 to a non-Monday date.
        assert analysis["A32"].value == (
            '=YEAR(L32-WEEKDAY(L32,2)+4)&"-W"&TEXT(WEEKNUM(L32,21),"00")'
        )
        assert analysis["L32"].value == "=$M$2-WEEKDAY($M$2,2)+1"
        assert analysis["M32"].value == "=$L$32+6"
        period_end = wb["工具"]["F2"].value.date()
        assert f"{period_end.isocalendar().year}-W{period_end.isocalendar().week:02d}" == iso_week
        assert "仅统计本次导出明细" in analysis["A61"].value
    finally:
        wb.close()


@pytest.mark.parametrize("legacy", [False, True])
def test_changed_preflight_count_preserves_existing_file(ctx, payload, tmp_path, legacy):
    ctx.inspections.save(payload)
    target = tmp_path / "existing.xlsx"
    target.write_bytes(b"existing report")
    with pytest.raises(ValueError, match="记录数已变化"):
        ctx.excel.export(target, RecordFilter(), legacy=legacy, prefer_com=False, expected_count=0)
    assert target.read_bytes() == b"existing report"
    assert not list(tmp_path.glob("tmp*.xlsx"))


def test_legacy_charts_include_rework_outside_first_eight_configured_teams(ctx, payload, tmp_path):
    ctx.settings.save_team("U9", sort_order=99)
    ctx.inspections.save(payload.model_copy(update={"team": "U9", "judgment": "返工"}))
    wb = load_workbook(
        ctx.excel.export(
            tmp_path / "ninth-team.xlsx",
            RecordFilter(start=payload.inspection_date, end=payload.inspection_date),
            legacy=True,
            prefer_com=False,
        )
    )
    try:
        chart = wb["数据分析表"]._charts[0]
        assert sum(p.v for p in chart.series[0].val.numRef.numCache.pt) == 1
        assert "其他组别" in [p.v for p in chart.series[0].cat.strRef.strCache.pt]
        assert "SUM(F4:F10)" in wb["数据分析表"]["F11"].value
    finally:
        wb.close()


def test_legacy_chart_labels_keep_names_from_export_snapshot(ctx, payload, tmp_path, monkeypatch):
    ctx.inspections.save(
        payload.model_copy(
            update={
                "defect_quantity": 1,
                "defects": [DefectInput(defect_id=1, quantity=1)],
            }
        )
    )
    old_name = ctx.defects.list()[0]["name"]
    original = ctx.inspections.iter_records

    def changing_names(filters):
        yield from original(filters)
        ctx.defects.save({"code": "a", "name": "导出期间修改的名称"}, 1)

    monkeypatch.setattr(ctx.inspections, "iter_records", changing_names)
    wb = load_workbook(
        ctx.excel.export(tmp_path / "names.xlsx", RecordFilter(), legacy=True, prefer_com=False)
    )
    try:
        assert wb["数据分析表"]["A4"].value == old_name
        assert wb["工具"]["B2"].value == old_name
    finally:
        wb.close()


def test_legacy_blank_feedback_column_keeps_template_borders_on_filled_rows(ctx, payload, tmp_path):
    seed_records(ctx, payload, 501)
    template = load_workbook(ctx.paths.template)
    wb = load_workbook(
        ctx.excel.export(tmp_path / "borders.xlsx", RecordFilter(), legacy=True, prefer_com=False)
    )
    try:
        for row in (2, 500, 501, 502):
            for col in range(1, 12):
                assert (
                    wb["成品日检表"].cell(row, col)._style
                    == template["成品日检表"].cell(2, col)._style
                )
    finally:
        template.close()
        wb.close()


@pytest.mark.parametrize("legacy", [False, True])
def test_save_failure_does_not_replace_existing_workbook(
    ctx, payload, tmp_path, monkeypatch, legacy
):
    from pathlib import Path

    from openpyxl.workbook.workbook import Workbook

    ctx.inspections.save(payload)
    target = tmp_path / "good.xlsx"
    target.write_bytes(b"previous good workbook")

    def fail_save(self, path):
        Path(path).write_bytes(b"partial write")
        raise OSError("磁盘写入失败")

    monkeypatch.setattr(Workbook, "save", fail_save)
    with pytest.raises(OSError, match="磁盘写入失败"):
        ctx.excel.export(target, RecordFilter(), legacy=legacy, prefer_com=False)
    assert target.read_bytes() == b"previous good workbook"
    assert not list(tmp_path.glob("tmp*.xlsx"))


def test_literal_case_sensitive_team_formulas_and_other_total(ctx, payload, tmp_path):
    ctx.settings.save_team("U*", team_id=1, sort_order=1)
    ctx.settings.save_team("u2", sort_order=99)
    ctx.settings.save_team("U9", sort_order=100)
    for team in ("U2", "u2", "U9"):
        ctx.inspections.save(payload.model_copy(update={"team": team, "judgment": "返工"}))
    wb = load_workbook(
        ctx.excel.export(tmp_path / "literal.xlsx", RecordFilter(), legacy=True, prefer_com=False)
    )
    try:
        analysis = wb["数据分析表"]
        categories = [p.v for p in analysis._charts[0].series[0].cat.strRef.strCache.pt]
        values = [p.v for p in analysis._charts[0].series[0].val.numRef.numCache.pt]
        assert values[categories.index("U*")] == 0
        assert values[categories.index("U2")] == 1
        assert values[-1] == 2 and sum(values) == 3
        for first in (4, 34):
            for r in range(first, first + 7):
                assert analysis.cell(r, 6).value.startswith("=SUMPRODUCT(")
                assert f"EXACT('成品日检表'!$C$2:$C$4,E{r})" in analysis.cell(r, 6).value
    finally:
        wb.close()


def test_standard_dictionary_name_matches_stream_snapshot(ctx, payload, tmp_path, monkeypatch):
    ctx.inspections.save(
        payload.model_copy(
            update={
                "defect_quantity": 1,
                "defects": [DefectInput(defect_id=1, quantity=1)],
            }
        )
    )
    original = ctx.inspections.iter_records

    def rename_before_stream(filters):
        ctx.defects.save({"code": "a", "name": "=快照名称<&>"}, 1)
        yield from original(filters)

    monkeypatch.setattr(ctx.inspections, "iter_records", rename_before_stream)
    wb = load_workbook(ctx.excel.export(tmp_path / "dictionary.xlsx", RecordFilter()))
    try:
        assert wb["不良明细"]["C2"].value == "=快照名称<&>"
        assert wb["不良项目"]["B2"].value == "=快照名称<&>"
        assert wb["不良项目"]["B2"].data_type == "s"
    finally:
        wb.close()


@pytest.mark.parametrize("legacy", [True, False])
@pytest.mark.parametrize("configured", [True, False])
def test_both_formats_protect_original_and_bundled_templates(
    ctx, tmp_path, monkeypatch, legacy, configured
):
    original = tmp_path / "protected-template.xlsx"
    original.write_bytes(ctx.paths.template.read_bytes())
    before = original.read_bytes()
    if configured:
        ctx.settings.update({"template_path": str(original)})
    else:
        monkeypatch.setattr(ctx.paths, "template", original)
    with pytest.raises(ValueError, match="不能覆盖模板"):
        ctx.excel.export(original, RecordFilter(), legacy=legacy, prefer_com=False)
    assert original.read_bytes() == before


@pytest.mark.parametrize("change_time", ["before", "after"])
def test_standard_dictionary_resolves_all_codes_used_by_snapshot(
    ctx, payload, tmp_path, monkeypatch, change_time
):
    row = ctx.inspections.save(
        payload.model_copy(
            update={
                "defect_quantity": 1,
                "defects": [DefectInput(defect_id=1, quantity=1)],
            }
        )
    )
    original = ctx.inspections.iter_records

    def change_codes(filters):
        if change_time == "before":
            ctx.defects.save({"code": "new-code", "name": "新项目"})
            item = next(d for d in ctx.defects.list() if d["code"] == "new-code")
            ctx.inspections.save(
                payload.model_copy(
                    update={
                        "defect_quantity": 1,
                        "defects": [DefectInput(defect_id=item["id"], quantity=1)],
                    }
                ),
                row["id"],
            )
        yield from original(filters)
        if change_time == "after":
            ctx.inspections.save(payload, row["id"])
            ctx.defects.save({"code": "renamed-unused", "name": "改名的闲置项目"}, 1)

    monkeypatch.setattr(ctx.inspections, "iter_records", change_codes)
    wb = load_workbook(ctx.excel.export(tmp_path / "codes.xlsx", RecordFilter()))
    try:
        names = {r[0]: r[1] for r in list(wb["不良项目"].values)[1:]}
        details = list(wb["不良明细"].values)[1:]
        assert len(details) == 1
        for detail in details:
            assert names[detail[1]] == detail[2]
        assert details[0][1] == ("new-code" if change_time == "before" else "a")
    finally:
        wb.close()
