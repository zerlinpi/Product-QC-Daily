import pytest
from PySide6.QtCore import QDate, QLocale, Qt
from PySide6.QtWidgets import QDialogButtonBox, QFileDialog, QPushButton

from app.ui.dialogs import file_dialogs
from app.ui.main_window import MainWindow


def test_standard_controls_are_chinese_even_on_english_system(ctx, qtbot):
    QLocale.setDefault(QLocale("en_US"))
    window = MainWindow(ctx)
    qtbot.addWidget(window)
    buttons = QDialogButtonBox(
        QDialogButtonBox.StandardButton.Save | QDialogButtonBox.StandardButton.Cancel, window
    )
    assert "保存" in buttons.button(QDialogButtonBox.StandardButton.Save).text()
    assert "取消" in buttons.button(QDialogButtonBox.StandardButton.Cancel).text()
    picker = QFileDialog(window)
    assert "文件" in picker.labelText(QFileDialog.DialogLabel.FileName)
    assert "位置" in picker.labelText(QFileDialog.DialogLabel.LookIn)
    assert "类型" in picker.labelText(QFileDialog.DialogLabel.FileType)
    assert window.pages[1].inspection_date.locale().language() == QLocale.Language.Chinese


def test_custom_report_dates_update_visible_export_scope(ctx, qtbot):
    window = MainWindow(ctx)
    qtbot.addWidget(window)
    page = window.pages[5]
    page.preset.setCurrentText("自定义")
    page.start.setDate(QDate(2025, 12, 30))
    page.end.setDate(QDate(2026, 1, 3))
    assert "2025-12-30 至 2026-01-03" in page.export_scope.text()


@pytest.mark.parametrize("detailed", [False, True])
def test_export_selection_uses_displayed_rows_not_unapplied_filters(
    ctx, payload, qtbot, monkeypatch, tmp_path, detailed
):
    from openpyxl import load_workbook

    ctx.inspections.save(payload)
    window = MainWindow(ctx)
    qtbot.addWidget(window)
    page = window.pages[2]
    page.refresh()
    page.table.selectRow(0)
    page.source.setCurrentText("演示数据")
    output = tmp_path / "selected.xlsx"
    monkeypatch.setattr(
        file_dialogs,
        "save_excel",
        lambda *args: (str(output), "明细报表 (*.xlsx)" if detailed else "原表格式 (*.xlsx)"),
    )
    monkeypatch.setattr(window, "run_job", lambda title, work, done: done(work()))
    page.export()
    wb = load_workbook(output)
    sheet = wb["检验记录" if detailed else "成品日检表"]
    assert "MO-001" in [cell.value for cell in sheet[2]]
    assert ("不良明细" in wb.sheetnames) is detailed
    wb.close()


def test_recycle_bin_switch_refreshes_rows_and_available_actions(ctx, payload, qtbot):
    from app.core.schemas import InspectionInput

    ctx.inspections.save(payload)
    deleted = ctx.inspections.save(
        InspectionInput(**(payload.model_dump() | {"work_order": "DELETED"}))
    )
    ctx.inspections.delete([deleted["id"]])
    window = MainWindow(ctx)
    qtbot.addWidget(window)
    page = window.pages[2]
    page.refresh()
    page.table.selectRow(0)
    page.trash.setChecked(True)
    assert [r["work_order"] for r in page.rows] == ["DELETED"]
    assert not page.selected_ids()
    page.table.selectRow(0)
    assert page.action_buttons["restore"].isEnabled()
    assert not page.action_buttons["edit"].isEnabled()
    assert not page.action_buttons["copy"].isEnabled()


def test_dashboard_new_inspection_does_not_reopen_previous_saved_record(ctx, payload, qtbot):
    saved = ctx.inspections.save(payload)
    window = MainWindow(ctx)
    qtbot.addWidget(window, before_close_func=lambda w: setattr(w.pages[1], "dirty", False))
    window.open_record(saved["id"])
    window.navigate(0)
    control = next(b for b in window.pages[0].findChildren(QPushButton) if "新建检验" in b.text())
    qtbot.mouseClick(control, Qt.MouseButton.LeftButton)
    assert window.stack.currentIndex() == 1
    assert window.pages[1].record_id is None
    assert ctx.inspections.get(saved["id"])["work_order"] == "MO-001"


@pytest.mark.parametrize("accept", [False, True])
def test_confirmation_defaults_to_cancel_and_requires_explicit_acceptance(qtbot, accept):
    from PySide6.QtCore import QTimer
    from PySide6.QtWidgets import QApplication, QMessageBox

    from app.ui.common import confirm

    seen = {}

    def choose():
        dialog = QApplication.activeModalWidget()
        seen["default"] = dialog.defaultButton().text()
        chosen = next(
            b
            for b in dialog.buttons()
            if dialog.buttonRole(b)
            == (QMessageBox.ButtonRole.AcceptRole if accept else QMessageBox.ButtonRole.RejectRole)
        )
        chosen.click()

    QTimer.singleShot(0, choose)
    assert confirm(None, "移入回收站", "将所选记录移入回收站？", action="移入回收站") is accept
    assert seen["default"] == "取消"


def test_import_preview_pagination_matches_available_rows(ctx, qtbot, tmp_path):
    from datetime import datetime

    from openpyxl import Workbook

    from app.ui.dialogs.import_dialog import ImportDialog

    path = tmp_path / "preview.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "成品日检表"
    ws.append(
        ["填写ID", "时间", "组别", "加工单号", "检验数量", "抽检数", "不良数", "不良项目", "判定"]
    )
    for i in range(201):
        ws.append([f"PREVIEW-{i}", datetime(2026, 1, 5), "U1", "MO-001", 100, 20, 0, "", "合格"])
    wb.save(path)
    wb.close()
    window = MainWindow(ctx)
    qtbot.addWidget(window)
    dialog = ImportDialog(ctx, window, ctx.excel.preview(path))
    qtbot.addWidget(dialog)
    assert not dialog.previous_button.isEnabled()
    assert dialog.next_button.isEnabled()
    dialog.next_button.click()
    assert dialog.table.rowCount() == 1
    assert dialog.previous_button.isEnabled()
    assert not dialog.next_button.isEnabled()
    dialog.filter.setCurrentText("异常")
    assert dialog.table.rowCount() == 0
    assert dialog.visible_status.text() == "当前筛选无记录"
    assert dialog.visible_status.objectName() == "empty"
    assert not dialog.previous_button.isEnabled()
    assert not dialog.next_button.isEnabled()


def test_entry_footer_tracks_saved_modified_copied_and_new_records(ctx, payload, qtbot):
    saved = ctx.inspections.save(payload)
    window = MainWindow(ctx)
    qtbot.addWidget(window, before_close_func=lambda w: setattr(w.pages[1], "dirty", False))
    page = window.pages[1]
    page.load_record(saved)
    assert "已保存" in page.saved_note.text()
    assert page.saved_note.objectName() == "success"
    page.remark.setPlainText("修改后的备注")
    assert "修改尚未保存" in page.saved_note.text()
    assert page.saved_note.objectName() == "warning"
    page.refresh()
    assert "修改尚未保存" in page.saved_note.text()
    page.discard_changes()
    assert "已保存" in page.saved_note.text()
    page.load_record(saved, copy_record=True)
    assert "尚未保存" in page.saved_note.text()
    page.save_record()
    assert "保存成功" in page.saved_note.text()
    assert page.saved_note.objectName() == "success"
    page.refresh()
    assert "保存成功" in page.saved_note.text()
    page.reset()
    assert "本条尚未保存" in page.saved_note.text()



def test_record_context_menu_targets_clicked_row(ctx, payload, qtbot):
    from app.core.schemas import InspectionInput

    ctx.inspections.save(InspectionInput(**(payload.model_dump() | {"work_order": "RIGHT-A"})))
    ctx.inspections.save(InspectionInput(**(payload.model_dump() | {"work_order": "RIGHT-B"})))
    window = MainWindow(ctx)
    qtbot.addWidget(window)
    window.show()
    window.navigate(2)
    page = window.pages[2]
    page.table.selectRow(0)
    target_id = page.rows[1]["id"]
    point = page.table.visualItemRect(page.table.item(1, 0)).center()
    page.select_context_row(point)

    assert page.selected_ids() == [target_id]
    assert page.selection_count.text() == "已选择 1 条"


def test_defect_dialog_uses_desktop_window_flags_and_default_save(ctx, qtbot):
    from PySide6.QtWidgets import QDialogButtonBox

    from app.ui.dialogs.defect_dialog import DefectDialog

    dialog = DefectDialog(ctx, None)
    qtbot.addWidget(dialog)
    assert not bool(dialog.windowFlags() & Qt.WindowType.WindowContextHelpButtonHint)
    buttons = dialog.findChild(QDialogButtonBox)
    assert buttons.button(QDialogButtonBox.StandardButton.Save).isDefault()
    assert dialog.fields["code"].placeholderText() == "例如 A01"
    assert dialog.fields["name"].placeholderText() == "请输入不良项目名称"
    assert dialog.enabled.text() == "启用此项目"


def test_team_dialog_uses_desktop_window_flags_and_default_save(ctx, qtbot, monkeypatch):
    from PySide6.QtWidgets import QDialog, QDialogButtonBox

    window = MainWindow(ctx)
    qtbot.addWidget(window)
    page = window.pages[6]
    observed = {}

    def inspect(dialog):
        observed["help"] = bool(
            dialog.windowFlags() & Qt.WindowType.WindowContextHelpButtonHint
        )
        buttons = dialog.findChild(QDialogButtonBox)
        observed["default_save"] = buttons.button(QDialogButtonBox.StandardButton.Save).isDefault()
        observed["minimum_width"] = dialog.minimumWidth()
        return QDialog.DialogCode.Rejected

    monkeypatch.setattr(QDialog, "exec", inspect)
    page.edit_team(False)

    from app.ui.common import FORM_DIALOG_MIN_WIDTH

    assert observed == {
        "help": False,
        "default_save": True,
        "minimum_width": FORM_DIALOG_MIN_WIDTH,
    }



def test_import_preview_uses_native_dialog_button_box(ctx, qtbot, tmp_path):
    from datetime import datetime

    from openpyxl import Workbook

    from app.ui.dialogs.import_dialog import ImportDialog

    path = tmp_path / "native-preview.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "成品日检表"
    ws.append(
        ["填写ID", "时间", "组别", "加工单号", "检验数量", "抽检数", "不良数", "不良项目", "判定"]
    )
    ws.append(["NATIVE-1", datetime(2026, 1, 5), "U1", "MO-1", 100, 20, 0, "", "合格"])
    wb.save(path)
    wb.close()

    window = MainWindow(ctx)
    qtbot.addWidget(window)
    dialog = ImportDialog(ctx, window, ctx.excel.preview(path))
    qtbot.addWidget(dialog)
    buttons = dialog.findChild(QDialogButtonBox)
    assert buttons is not None
    accept = next(
        button
        for button in buttons.buttons()
        if buttons.buttonRole(button) == QDialogButtonBox.ButtonRole.AcceptRole
    )
    assert accept.isDefault()
    assert buttons.button(QDialogButtonBox.StandardButton.Cancel) is not None
    assert dialog.visible_status.text() == "当前显示 1 条"
    assert dialog.table.item(0, 2).foreground().color().name() == "#16a34a"
    assert dialog.filter.accessibleName() == "导入状态筛选"



def test_records_pending_filters_are_explicit_and_do_not_mix_pagination(
    ctx, payload, qtbot
):
    from app.core.schemas import InspectionInput
    from app.ui.main_window import MainWindow

    with ctx.db.session() as session:
        for index in range(51):
            ctx.inspections.save_in_session(
                session,
                InspectionInput(
                    **(
                        payload.model_dump()
                        | {
                            "work_order": f"PENDING-{index:02d}",
                            "source": "manual",
                        }
                    )
                ),
            )

    window = MainWindow(ctx)
    qtbot.addWidget(window)
    window.navigate(2)
    page = window.pages[2]
    assert page.total == 51
    assert page.next_button.isEnabled()
    original_summary = page.count.text()
    assert not page.filters_dirty

    page.source.setCurrentText("演示数据")
    assert page.filters_dirty
    assert "筛选条件已更改" in page.count.text()
    assert "上一次查询结果" in page.count.text()
    assert not page.previous_button.isEnabled()
    assert not page.next_button.isEnabled()
    assert "筛选条件尚未应用" in page.action_buttons["export"].toolTip()
    assert page.applied_filters.source == "production"

    page.source.setCurrentText("正式数据")
    assert not page.filters_dirty
    assert page.count.text() == original_summary
    assert page.next_button.isEnabled()

    page.search.setText("不存在的工单")
    assert page.filters_dirty
    page.refresh()
    assert not page.filters_dirty
    assert page.total == 0
    assert page.applied_filters.search == "不存在的工单"
    assert page.count.text().startswith("共 0 条")
