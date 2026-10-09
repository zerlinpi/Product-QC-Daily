"""Native month controls and readable defect names must survive export/import."""

from datetime import date, datetime

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from app.core.schemas import DefectInput, RecordFilter


@pytest.mark.parametrize("legacy", [False, True])
@pytest.mark.parametrize("source", ["demo", "manual"])
def test_export_shows_defect_names_and_roundtrips_identity(
    ctx, payload, tmp_path, legacy, source
):
    record = ctx.inspections.save(payload.model_copy(update={
        "source": source,
        "defect_quantity": 2,
        "defects": [DefectInput(defect_id=1, quantity=None), DefectInput(defect_id=2, quantity=2)],
    }))
    path = ctx.excel.export(
        tmp_path / "names.xlsx", RecordFilter(source="demo" if source == "demo" else "production"),
        legacy=legacy, prefer_com=False,
    )
    wb = load_workbook(path)
    try:
        ws = wb["成品日检表" if legacy else "检验记录"]
        assert ws["H2"].value == "端子包角不良；走线槽漏扎带"
        code_column = "O" if legacy else "Q"
        assert ws[f"{code_column}1"].value == "不良项目编码"
        assert ws[f"{code_column}2"].value == "a;b"
        assert ws.column_dimensions[code_column].hidden
        assert ws["H2"].alignment.wrap_text
        if legacy:
            # Two six-character Chinese names need three lines at the original
            # 13-character width; keep the template width, font and all text.
            assert ws.column_dimensions["H"].width == 13
            assert ws["H2"].font.sz == 10
            assert ws.row_dimensions[2].height == 48
        assert ws.auto_filter.ref == ("B1:K2" if legacy else "A1:P2")
    finally:
        wb.close()
    preview = ctx.excel.preview(path)
    assert len(preview.rows) == 1
    row = preview.rows[0]
    assert row.data is not None, row.message
    assert row.inspection_no == record["inspection_no"]
    assert row.data.source == ("demo" if source == "demo" else "excel")
    assert [d.defect_id for d in row.data.defects] == [1, 2]
    assert [d.quantity for d in row.data.defects] == ([None, None] if legacy else [None, 2])


@pytest.mark.parametrize("legacy", [False, True])
def test_edited_display_name_does_not_silently_import_stale_hidden_code(
    ctx, payload, tmp_path, legacy
):
    ctx.inspections.save(payload.model_copy(update={
        "defect_quantity": 1,
        "defects": [DefectInput(defect_id=1)],
    }))
    path = ctx.excel.export(tmp_path / "edited.xlsx", RecordFilter(), legacy=legacy, prefer_com=False)
    wb = load_workbook(path)
    wb["成品日检表" if legacy else "检验记录"]["H2"] = "漏胶"
    wb.save(path)
    wb.close()
    row = ctx.excel.preview(path).rows[0]
    assert row.status == "unrecognized"
    assert row.data is None
    assert "名称与编码不一致" in row.message


@pytest.mark.parametrize("legacy", [False, True])
def test_exported_names_remain_importable_after_dictionary_rename(ctx, payload, tmp_path, legacy):
    # Names may contain separators or formula-like text; the explicit code
    # identifies the item without trying to split or execute its display name.
    ctx.defects.save({"code": "a", "name": "=外观；划伤/压痕"}, defect_id=1)
    ctx.inspections.save(payload.model_copy(update={
        "defect_quantity": 1, "defects": [DefectInput(defect_id=1)],
    }))
    path = ctx.excel.export(tmp_path / "renamed.xlsx", RecordFilter(), legacy=legacy, prefer_com=False)
    wb = load_workbook(path)
    ws = wb["成品日检表" if legacy else "检验记录"]
    assert ws["H2"].data_type == "s"
    assert ws["H2"].value == "=外观；划伤/压痕"
    wb.close()
    ctx.defects.save({"code": "a", "name": "外观损伤（新名称）"}, defect_id=1)
    row = ctx.excel.preview(path).rows[0]
    assert row.data is not None, row.message
    assert [d.defect_id for d in row.data.defects] == [1]
    assert row.data.defects[0].quantity is None


def test_month_control_preserves_unrelated_validation_and_removes_stale_week_input(ctx, tmp_path):
    template = tmp_path / "template.xlsx"
    wb = load_workbook(ctx.paths.template)
    ws = wb["数据分析表"]
    validation = DataValidation(type="list", formula1='"1,2,3"')
    validation.sqref = "A2 E2 H2 A32 E32 H32 N20:N21"
    ws.add_data_validation(validation)
    wb.save(template)
    wb.close()
    ctx.settings.update({"template_path": str(template)})
    path = ctx.excel.export(tmp_path / "validation.xlsx", RecordFilter(), legacy=True, prefer_com=False)
    wb = load_workbook(path)
    rules = wb["数据分析表"].data_validations.dataValidation
    assert any("N20" in rule.sqref and "N21" in rule.sqref for rule in rules)
    assert not any(cell in rule.sqref for cell in ("A32", "E32", "H32") for rule in rules)
    wb.close()


@pytest.mark.parametrize(
    "start,end,labels,bounds",
    [
        (date(2026, 1, 1), date(2026, 12, 31),
         ["全部", *[f"2026-{m:02d}" for m in range(1, 13)]],
         [(date(2026, 1, 1), date(2026, 12, 31)), (date(2026, 1, 1), date(2026, 1, 31))]),
        (date(2023, 12, 20), date(2024, 2, 29),
         ["全部", "2023-12", "2024-01", "2024-02"],
         [(date(2023, 12, 20), date(2024, 2, 29)), (date(2023, 12, 20), date(2023, 12, 31)),
          (date(2024, 1, 1), date(2024, 1, 31)), (date(2024, 2, 1), date(2024, 2, 29))]),
    ],
)
def test_original_form_keeps_working_month_selector(ctx, tmp_path, start, end, labels, bounds):
    path = ctx.excel.export(
        tmp_path / "months.xlsx", RecordFilter(start=start, end=end), legacy=True, prefer_com=False
    )
    wb = load_workbook(path)
    try:
        ws = wb["数据分析表"]
        assert ws["A2"].value == "全部"
        validations = [v for v in ws.data_validations.dataValidation if "A2" in v.sqref]
        assert len(validations) == 1
        validation = validations[0]
        assert validation.type == "list" and not validation.showDropDown
        assert validation.showErrorMessage and validation.errorStyle == "stop"
        name = wb.defined_names[validation.formula1.lstrip("=")]
        sheet_name, address = next(name.destinations)
        values = [row[0].value for row in wb[sheet_name][address]]
        assert values == labels
        for row_number, (low, high) in enumerate(bounds, 2):
            assert wb["工具"].cell(row_number, 5).value == datetime.combine(low, datetime.min.time())
            assert wb["工具"].cell(row_number, 6).value == datetime.combine(high, datetime.min.time())
        assert all(ws[c].data_type == "f" for c in ("L2", "M2", "L32", "M32"))
        assert ws["E2"].value == ws["H2"].value == "=$A$2"
        assert len(ws._charts) == 6
        assert all("所选期间" in "".join(c.title.to_tree().itertext()) for c in ws._charts[:3])
        assert wb["成品日检表"].auto_filter.ref == "B1:K2"
    finally:
        wb.close()
