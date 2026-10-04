from itertools import groupby
from pathlib import Path
from types import SimpleNamespace

import pytest
from PySide6.QtWidgets import QFileDialog, QPushButton

from app.core.schemas import InspectionInput, RecordFilter
from app.ui.main_window import MainWindow


def test_renaming_default_team_keeps_default_and_history_usable(ctx, payload):
    row = ctx.inspections.save(payload)
    team = next(t for t in ctx.settings.teams() if t["name"] == "U1")
    ctx.settings.save_team("一组", team["id"])
    assert ctx.inspections.get(row["id"])["team"] == "一组"
    assert ctx.settings.get("default_team") == "一组"
    new = InspectionInput(**(payload.model_dump() | {"team": ctx.settings.get("default_team")}))
    assert ctx.inspections.save(new)["team"] == "一组"


@pytest.mark.parametrize("descending", [True, False])
def test_date_column_sorts_time_within_same_day(ctx, payload, descending):
    for stamp in ["09:00:00", "17:00:00", "12:00:00"]:
        ctx.inspections.save(InspectionInput(**(payload.model_dump() | {"inspection_time": stamp})))
    rows, _ = ctx.inspections.query(RecordFilter(sort="inspection_date", descending=descending))
    assert [r["inspection_time"] for r in rows] == sorted(
        ["09:00:00", "17:00:00", "12:00:00"], reverse=descending
    )


def test_source_column_sorts_by_source_instead_of_date(ctx, payload):
    for source in ["manual", "excel", "demo", "excel"]:
        ctx.inspections.save(InspectionInput(**(payload.model_dump() | {"source": source})))
    rows, _ = ctx.inspections.query(RecordFilter(source="all", sort="source", descending=False))
    sources = [r["source"] for r in rows]
    assert all(sources.count(v) == len(list(g)) for v, g in groupby(sources))
    reversed_rows, _ = ctx.inspections.query(
        RecordFilter(source="all", sort="source", descending=True)
    )
    assert [r["source"] for r in reversed_rows] == list(reversed(sources))


@pytest.mark.parametrize("page_index", [2, 5])
def test_saved_export_directory_is_used_by_export_pages(
    ctx, payload, qtbot, monkeypatch, tmp_path, page_index
):
    directory = tmp_path / "自定义导出"
    directory.mkdir()
    ctx.settings.update({"export_directory": str(directory)})
    ctx.inspections.save(payload)
    window = MainWindow(ctx)
    qtbot.addWidget(window)
    page = window.pages[page_index]
    page.refresh()
    chosen = []
    monkeypatch.setattr(
        QFileDialog, "getSaveFileName", lambda *args: (chosen.append(Path(args[2])) or "", "")
    )
    if page_index == 2:
        page.export()
    else:
        page.export(False)
    assert chosen[0].parent == directory


def test_defect_search_stays_correct_after_dictionary_reload(ctx, qtbot):
    from app.ui.widgets.defect_selector import DefectSelector

    selector = DefectSelector(ctx)
    qtbot.addWidget(selector)
    selector.search.setText("端子包角")
    assert not selector.table.isRowHidden(0)
    ctx.defects.save({"code": "a", "name": "修改后的名称"}, 1)
    selector.reload()
    assert selector.search.text() == "端子包角"
    assert selector.table.isRowHidden(0)


def test_signature_warning_can_be_exported_and_uses_export_directory(
    ctx, qtbot, monkeypatch, tmp_path
):
    from datetime import datetime

    from openpyxl import Workbook

    from app.ui.dialogs.import_dialog import ImportDialog

    path = tmp_path / "missing-image.xlsx"
    wb = Workbook()
    wb.active.title = "成品日检表"
    wb.active.append(
        [
            "填写ID",
            "时间",
            "组别",
            "加工单号",
            "检验数量",
            "抽检数",
            "不良数",
            "不良项目",
            "判定",
            "签名",
        ]
    )
    wb.active.append(
        [
            "WARN-1",
            datetime(2026, 1, 5),
            "U1",
            "MO-1",
            100,
            20,
            0,
            "",
            "合格",
            '=DISPIMG("missing",1)',
        ]
    )
    wb.save(path)
    wb.close()
    preview = ctx.excel.preview(path)
    assert preview.counts["valid"] == 1 and preview.rows[0].message
    window = MainWindow(ctx)
    qtbot.addWidget(window)
    dialog = ImportDialog(ctx, window, preview)
    qtbot.addWidget(dialog)
    control = next(b for b in dialog.findChildren(QPushButton) if b.text() == "导出异常报告")
    assert control.isEnabled()
    directory = tmp_path / "报告目录"
    ctx.settings.update({"export_directory": str(directory)})
    chosen = []
    monkeypatch.setattr(
        QFileDialog, "getSaveFileName", lambda *args: (chosen.append(Path(args[2])) or "", "")
    )
    dialog.report()
    assert chosen[0].parent == directory


def test_com_cleanup_failure_does_not_abort_fallback_export(ctx, payload, tmp_path, monkeypatch):
    import sys

    from openpyxl import load_workbook

    from app.services import excel_export

    cleaned = []

    def fail_close(*args):
        cleaned.append("close")
        raise RuntimeError("Excel closed unexpectedly")

    workbook = SimpleNamespace(Close=fail_close)

    def fail_calculate():
        raise RuntimeError("Excel calculation failed")

    excel = SimpleNamespace(
        Workbooks=SimpleNamespace(Open=lambda *args, **kwargs: workbook),
        CalculateFullRebuild=fail_calculate,
        Quit=lambda: cleaned.append("quit"),
    )
    monkeypatch.setitem(
        sys.modules,
        "pythoncom",
        SimpleNamespace(
            CoInitialize=lambda: cleaned.append("init"),
            CoUninitialize=lambda: cleaned.append("uninit"),
        ),
    )
    client = SimpleNamespace(DispatchEx=lambda name: excel)
    monkeypatch.setitem(sys.modules, "win32com", SimpleNamespace(client=client))
    monkeypatch.setitem(sys.modules, "win32com.client", client)
    monkeypatch.setattr(excel_export, "os", SimpleNamespace(name="nt"))
    ctx.inspections.save(payload)
    path = ctx.excel.export(tmp_path / "fallback.xlsx", RecordFilter(), legacy=True)
    with_path = load_workbook(path)
    assert with_path["成品日检表"]["D2"].value == "MO-001"
    with_path.close()
    assert cleaned == ["init", "close", "quit", "uninit"]


def test_defect_category_filter_is_preserved_on_reload(ctx, qtbot):
    from app.ui.widgets.defect_selector import DefectSelector

    ctx.defects.save({"code": "appearance", "name": "外观检查", "category": "外观"})
    selector = DefectSelector(ctx)
    qtbot.addWidget(selector)
    selector.category.setCurrentText("外观")
    selector.reload()
    assert selector.category.currentText() == "外观"
    visible = [
        item["category"]
        for row, (item, _, _) in enumerate(selector.entries)
        if not selector.table.isRowHidden(row)
    ]
    assert visible == ["外观"]


def test_background_jobs_recover_after_failure_and_can_run_again(ctx, payload, qtbot, monkeypatch):
    from threading import Event, get_ident

    window = MainWindow(ctx)
    qtbot.addWidget(window)
    window.show()
    main_thread = get_ident()
    completed, errors = [], []
    started, release = Event(), Event()

    def work():
        started.set()
        assert release.wait(5)
        return get_ident(), ctx.inspections.save(payload)["id"]

    def saved(result):
        completed.append((result, get_ident()))

    monkeypatch.setattr(
        "app.ui.main_window.friendly_error", lambda parent, error: errors.append(str(error))
    )
    try:
        window.run_job("保存测试记录", work, saved)
        qtbot.waitUntil(started.is_set, timeout=5000)
        assert not window.centralWidget().isEnabled()
        window.close()
        assert window.isVisible()
        release.set()
        qtbot.waitUntil(lambda: window._job is None, timeout=10000)
        assert window.centralWidget().isEnabled()
        assert completed[0][0][0] != main_thread and completed[0][1] == main_thread
        assert ctx.inspections.get(completed[0][0][1])["work_order"] == "MO-001"

        def fail():
            raise ValueError("模拟任务失败")

        window.run_job("失败任务", fail)
        qtbot.waitUntil(lambda: window._job is None, timeout=5000)
        assert errors == ["模拟任务失败"]
        assert window.centralWidget().isEnabled()
        backups = []
        window.run_job("随后备份", ctx.backup.backup, backups.append)
        qtbot.waitUntil(lambda: window._job is None, timeout=5000)
        assert backups[0].is_file()
    finally:
        release.set()
        qtbot.waitUntil(lambda: window._job is None, timeout=10000)


def test_export_contains_every_page_and_selected_scope(ctx, payload, tmp_path):
    from openpyxl import load_workbook

    ids = []
    with ctx.db.session() as session:
        for i in range(505):
            row = ctx.inspections.save_in_session(
                session, InspectionInput(**(payload.model_dump() | {"work_order": f"PAGE-{i}"}))
            )
            ids.append(row.id)
    exported = ctx.excel.export(tmp_path / "all-pages.xlsx", RecordFilter(page=4))
    wb = load_workbook(exported)
    assert wb["检验记录"].max_row == 506
    assert {c.value for c in wb["检验记录"]["D"][1:]} == {f"PAGE-{i}" for i in range(505)}
    assert wb["统计摘要"]["B2"].value == 505
    wb.close()
    selected = ctx.excel.export(
        tmp_path / "selected-pages.xlsx", RecordFilter(ids=[ids[0], ids[-1]])
    )
    wb = load_workbook(selected)
    assert wb["检验记录"].max_row == 3
    assert {c.value for c in wb["检验记录"]["D"][1:]} == {"PAGE-0", "PAGE-504"}
    assert wb["统计摘要"]["B2"].value == 2
    wb.close()


def test_import_conflict_after_preview_rolls_back_earlier_rows(ctx, payload, tmp_path):
    from datetime import datetime

    from openpyxl import Workbook

    path = tmp_path / "preview-conflict.xlsx"
    wb = Workbook()
    wb.active.title = "成品日检表"
    wb.active.append(
        ["填写ID", "时间", "组别", "加工单号", "检验数量", "抽检数", "不良数", "不良项目", "判定"]
    )
    for number in ["NEW-1", "NEW-2"]:
        wb.active.append([number, datetime(2026, 1, 5), "U1", "IMPORT", 100, 20, 0, "", "合格"])
    wb.save(path)
    wb.close()
    preview = ctx.excel.preview(path)
    assert preview.counts["valid"] == 2
    with ctx.db.session() as session:
        ctx.inspections.save_in_session(session, payload, inspection_no="NEW-2")
    with pytest.raises(ValueError, match="未导入任何记录"):
        ctx.excel.import_preview(preview)
    rows, count = ctx.inspections.query()
    assert count == 1 and rows[0]["inspection_no"] == "NEW-2"
    assert rows[0]["work_order"] == "MO-001"


def test_backup_tamper_is_rejected_without_changing_current_data(ctx, payload, tmp_path):
    from zipfile import ZIP_DEFLATED, ZipFile

    ctx.inspections.save(payload)
    source = ctx.backup.backup()
    tampered = tmp_path / "tampered.zip"
    with ZipFile(source) as original, ZipFile(tampered, "w", ZIP_DEFLATED) as modified:
        for name in original.namelist():
            value = original.read(name)
            modified.writestr(name, value + b"changed" if name == "product_qc.db" else value)
    later = ctx.inspections.save(payload)
    with pytest.raises(ValueError, match="校验失败"):
        ctx.backup.restore(tampered)
    assert ctx.inspections.query()[1] == 2
    assert ctx.inspections.get(later["id"])["work_order"] == "MO-001"
    assert ctx.db.health_check() == "ok"


def test_renamed_default_team_is_used_without_restarting_window(ctx, qtbot):
    window = MainWindow(ctx)
    qtbot.addWidget(window)
    ctx.settings.save_team("一组", ctx.settings.teams()[0]["id"])
    window.navigate(1)
    assert window.pages[1].team.currentText() == "一组"


def test_returning_to_saved_editor_loads_bulk_changes(ctx, payload, qtbot):
    saved = ctx.inspections.save(payload)
    window = MainWindow(ctx)
    qtbot.addWidget(window)
    window.open_record(saved["id"])
    window.navigate(2)
    ctx.inspections.bulk_update([saved["id"]], team="U2", inspector="李工")
    window.navigate(1)
    entry = window.pages[1]
    assert entry.team.currentText() == "U2"
    assert entry.inspector.text() == "李工"
    assert not entry.dirty
    entry.save_record()
    assert ctx.inspections.get(saved["id"])["inspector"] == "李工"


@pytest.mark.parametrize("remove", ["trash", "clear_demo"])
def test_returning_to_removed_editor_starts_new_record(ctx, payload, qtbot, remove):
    saved = ctx.inspections.save(InspectionInput(**(payload.model_dump() | {"source": "demo"})))
    window = MainWindow(ctx)
    qtbot.addWidget(window)
    window.open_record(saved["id"])
    window.navigate(2)
    if remove == "trash":
        ctx.inspections.delete([saved["id"]])
    else:
        ctx.demo.clear()
    window.navigate(1)
    assert window.pages[1].record_id is None
    assert not window.pages[1].dirty


def test_report_import_preview_and_record_export_use_real_background_jobs(
    ctx, qtbot, monkeypatch, tmp_path
):
    from datetime import datetime

    from openpyxl import Workbook, load_workbook
    from PySide6.QtCore import QTimer
    from PySide6.QtWidgets import QApplication

    from app.ui.dialogs.import_dialog import ImportDialog

    source = tmp_path / "workflow.xlsx"
    wb = Workbook()
    wb.active.title = "成品日检表"
    wb.active.append(
        ["填写ID", "时间", "组别", "加工单号", "检验数量", "抽检数", "不良数", "不良项目", "判定"]
    )
    wb.active.append(
        ["FLOW-1", datetime(2026, 1, 5, 8, 30), "U1", "FLOW-ORDER", 100, 20, 2, "ag", "返工"]
    )
    wb.active.append(["FLOW-BAD", datetime(2026, 1, 5), "U1", "", 100, 20, 0, "", "合格"])
    wb.save(source)
    wb.close()
    window = MainWindow(ctx)
    qtbot.addWidget(window)
    window.show()
    window.navigate(5)
    report = window.pages[5]
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *args: (str(source), ""))
    errors, preview_counts = [], []
    monkeypatch.setattr(
        "app.ui.main_window.friendly_error", lambda parent, error: errors.append(str(error))
    )
    timer = QTimer(window)

    def accept_preview():
        dialog = QApplication.activeModalWidget()
        if isinstance(dialog, ImportDialog):
            preview_counts.append(dialog.preview.counts)
            next(
                b for b in dialog.findChildren(QPushButton) if b.text().startswith("导入 1 条")
            ).click()

    timer.timeout.connect(accept_preview)
    timer.start(20)
    try:
        report.import_file()
        qtbot.waitUntil(
            lambda: report.import_status.text().startswith("已导入") or bool(errors), timeout=10000
        )
        assert not errors
        assert preview_counts[0]["valid"] == 1 and preview_counts[0]["invalid"] == 1
        assert ctx.inspections.query()[1] == 1
        window.navigate(2)
        records = window.pages[2]
        records.table.selectRow(0)
        exported = tmp_path / "workflow-out.xlsx"
        monkeypatch.setattr(
            QFileDialog, "getSaveFileName", lambda *args: (str(exported), "明细报表 (*.xlsx)")
        )
        records.export()
        qtbot.waitUntil(lambda: window._job is None, timeout=10000)
        assert not errors and exported.is_file()
        result = load_workbook(exported)
        assert result["检验记录"]["D2"].value == "FLOW-ORDER"
        assert result["不良明细"].max_row == 3
        result.close()
    finally:
        timer.stop()
        dialog = QApplication.activeModalWidget()
        if isinstance(dialog, ImportDialog):
            dialog.reject()
        qtbot.waitUntil(lambda: window._job is None, timeout=10000)


def test_record_filters_combine_correctly_and_treat_wildcards_literally(ctx, payload):
    matching = ctx.inspections.save(
        InspectionInput(
            **(
                payload.model_dump()
                | {
                    "team": "U2",
                    "work_order": "ORDER_50%",
                    "inspector": "李工",
                    "judgment": "返工",
                    "defect_quantity": 2,
                    "defects": [{"defect_id": 1, "quantity": None}],
                    "remark": "待复检",
                }
            )
        )
    )
    ctx.inspections.save(InspectionInput(**(payload.model_dump() | {"work_order": "ORDERX500"})))
    ctx.inspections.save(
        InspectionInput(**(payload.model_dump() | {"work_order": "ORDER_50%", "source": "demo"}))
    )
    filters = RecordFilter(
        start="2026-01-05",
        end="2026-01-05",
        team="U2",
        work_order="_50%",
        inspector="李",
        judgment="返工",
        defect_id=1,
        has_defects=True,
        search="待复检",
    )
    rows, count = ctx.inspections.query(filters)
    assert count == 1 and rows[0]["id"] == matching["id"]
    assert ctx.inspections.query(RecordFilter(work_order="_50%"))[1] == 1
    ctx.inspections.delete([matching["id"]])
    assert ctx.inspections.query(filters)[1] == 0
    assert ctx.inspections.query(filters.model_copy(update={"deleted": True}))[1] == 1
    ctx.inspections.restore([matching["id"]])
    assert ctx.inspections.query(filters)[1] == 1
