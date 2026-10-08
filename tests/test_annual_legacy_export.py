"""年度原表导出回归：真实 SQLite、真实演示数据和实际生成的 XLSX。"""

from collections import Counter
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from app.core.schemas import RecordFilter
from app.services.demo_service import suggested_demo_count


@pytest.mark.parametrize("year,expected_count,expected_days", [(2026, 1825, 365), (2024, 1830, 366)])
def test_full_year_legacy_export_preserves_all_rows_and_template(
    ctx, payload, tmp_path, year, expected_count, expected_days
):
    start, end = date(year, 1, 1), date(year, 12, 31)
    teams = [item["name"] for item in ctx.settings.teams(True)]
    assert suggested_demo_count(start, end) == expected_count

    # Same date and same source range: production must never leak into demo exports.
    production = payload.model_copy(update={"inspection_date": f"{year}-01-05"})
    saved_production = ctx.inspections.save(production)
    assert ctx.demo.generate(expected_count, start, end, teams, seed=20261008) == expected_count
    filters = RecordFilter(start=start, end=end, source="demo", descending=False)
    expected_rows = list(ctx.inspections.iter_records(filters))
    assert len(expected_rows) == expected_count

    export_path = ctx.excel.export(
        tmp_path / f"legacy-{year}.xlsx", filters, legacy=True, prefer_com=False
    )
    wb = load_workbook(export_path)
    original = load_workbook(ctx.paths.template)
    try:
        assert wb.sheetnames == original.sheetnames
        ws, source_ws = wb["成品日检表"], original["成品日检表"]
        analysis = wb["数据分析表"]
        assert len(analysis._charts) == 6
        assert wb["成品日检表报表"].sheet_state == "hidden"
        assert ws.max_row >= expected_count + 1
        ids = [ws.cell(i, 1).value for i in range(2, expected_count + 2)]
        assert ids == [r["inspection_no"] for r in expected_rows]
        assert saved_production["inspection_no"] not in ids
        assert all(ws.cell(i, 14).value == "演示数据" for i in range(2, expected_count + 2))
        exported_dates = [ws.cell(i, 2).value.date() for i in range(2, expected_count + 2)]
        assert exported_dates == sorted(exported_dates)
        assert exported_dates[0] == start and exported_dates[-1] == end
        day_counts = Counter(exported_dates)
        assert len(day_counts) == expected_days
        assert min(day_counts.values()) >= 1
        if year == 2024:
            assert date(2024, 2, 29) in day_counts

        # Regression boundary: 501st record and last record must be styled and readable.
        for row_no in (2, 501, 502, expected_count + 1):
            assert isinstance(ws.cell(row_no, 2).value, datetime)
            assert ws.cell(row_no, 2).number_format == source_ws["B2"].number_format
            assert ws.cell(row_no, 2).style_id == ws["B2"].style_id
            assert ws.cell(row_no, 4).style_id == ws["D2"].style_id
            assert ws.row_dimensions[row_no].height == pytest.approx(
                source_ws.row_dimensions[2].height
            )
        for column in ("A", "B", "C", "D", "E", "F", "G", "H", "I", "J"):
            assert ws.column_dimensions[column].width == source_ws.column_dimensions[column].width
        assert ws.print_area == source_ws.print_area
        assert ws.print_title_rows == source_ws.print_title_rows
        assert ws.page_setup.orientation == source_ws.page_setup.orientation
        assert ws.sheet_view.topLeftCell == "A1"
        assert ws.freeze_panes == "C2"
        assert wb.active.title == "成品日检表"

        # Monthly formulas must refer to ALL 1825/1830 rows, not the legacy 500.
        last = expected_count + 1
        assert f"$B${last}" in analysis["B4"].value
        assert f"$B${last}" in analysis["B34"].value
        assert f"$E${last}" in analysis["H4"].value
        assert analysis["L2"].value.date() == start
        assert analysis["M2"].value.date() == end
        assert analysis["C4"].value.endswith("$B$4:$B$27,0)")
        assert analysis["B58"].value == "=SUM(B34:B57)"

        # Six charts retain template locations and have non-stale caches.
        for result_chart, original_chart in zip(analysis._charts, original["数据分析表"]._charts, strict=True):
            assert result_chart.anchor._from == original_chart.anchor._from
        totals = [
            sum(r["inspection_quantity"] for r in expected_rows),
            sum(r["sampling_quantity"] for r in expected_rows),
            sum(r["defect_quantity"] for r in expected_rows),
        ]
        cached = [
            point.v
            for point in analysis._charts[1].series[0].val.numRef.numCache.pt
        ]
        assert cached[:3] == totals
        assert cached[3] == pytest.approx(totals[2] / totals[1])
    finally:
        wb.close()
        original.close()
