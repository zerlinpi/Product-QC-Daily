"""Independent SQLite/OOXML oracle: no production statistics/export helpers."""

import posixpath
import sqlite3
from collections import Counter, defaultdict
from copy import copy
from datetime import date, datetime, time, timedelta
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from app.core.schemas import DefectInput, InspectionInput, RecordFilter

NS = {
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
}


def raw_records(ctx, start, end, source="production", deleted=False):
    with sqlite3.connect(ctx.db.path) as db:
        db.row_factory = sqlite3.Row
        source_sql = "source = 'demo'" if source == "demo" else "source != 'demo'"
        rows = [
            dict(row)
            for row in db.execute(
                f"SELECT * FROM inspection_records WHERE deleted_at IS {'NOT ' if deleted else ''}NULL "
                f"AND {source_sql} AND inspection_date BETWEEN ? AND ? "
                "ORDER BY inspection_date, inspection_time, id",
                (str(start), str(end)),
            )
        ]
        for row in rows:
            row["defects"] = [
                dict(d)
                for d in db.execute(
                    "SELECT i.code,i.name,d.quantity,d.remark FROM inspection_defects d "
                    "JOIN defect_items i ON i.id=d.defect_id "
                    "WHERE d.inspection_id=? ORDER BY d.defect_id",
                    (row["id"],),
                )
            ]
        return rows


def expected_summary(rows):
    samples = sum(r["sampling_quantity"] for r in rows)
    defects = sum(r["defect_quantity"] for r in rows)
    passes = sum(r["judgment"] == "合格" for r in rows)
    rework = sum(r["judgment"] == "返工" for r in rows)
    return {
        "检验批次": len(rows),
        "检验数量": sum(r["inspection_quantity"] for r in rows),
        "抽检数量": samples,
        "不良件数": defects,
        "合格批次": passes,
        "返工批次": rework,
        "不良率": defects / samples if samples else 0,
        "合格率": passes / len(rows) if rows else 0,
        "返工率": rework / len(rows) if rows else 0,
    }


def assert_drawing_relationships(archive, expected_count):
    """Resolve the actual sheet -> drawing -> chart package graph."""
    charts = []
    for name in archive.namelist():
        if not (name.startswith("xl/worksheets/sheet") and name.endswith(".xml")):
            continue
        sheet = ET.fromstring(archive.read(name))
        drawing = sheet.find("x:drawing", NS)
        if drawing is None:
            continue
        rel_name = posixpath.join(
            posixpath.dirname(name), "_rels", posixpath.basename(name) + ".rels"
        )
        relationships = {
            r.get("Id"): r.get("Target") for r in ET.fromstring(archive.read(rel_name))
        }
        target = relationships[next(iter(drawing.attrib.values()))]
        drawing_name = (
            target.lstrip("/")
            if target.startswith("/")
            else posixpath.normpath(posixpath.join(posixpath.dirname(name), target))
        )
        drawing_rels = posixpath.join(
            posixpath.dirname(drawing_name), "_rels", posixpath.basename(drawing_name) + ".rels"
        )
        rels = {r.get("Id"): r.get("Target") for r in ET.fromstring(archive.read(drawing_rels))}
        for anchor in ET.fromstring(archive.read(drawing_name)):
            chart = anchor.find(".//c:chart", NS)
            if chart is None:
                continue
            assert anchor.find("a:from/a:row", NS) is not None
            assert anchor.find("a:to", NS) is not None or anchor.find("a:ext", NS) is not None
            target = rels[next(iter(chart.attrib.values()))]
            chart_name = (
                target.lstrip("/")
                if target.startswith("/")
                else posixpath.normpath(posixpath.join(posixpath.dirname(drawing_name), target))
            )
            assert chart_name in archive.namelist()
            charts.append(chart_name)
    assert len(charts) == len(set(charts)) == expected_count
    return charts


def assert_charts_xml(path, wb, rows, legacy, end):
    with ZipFile(path) as archive:
        chart_names = assert_drawing_relationships(archive, 6 if legacy else 2)
        for name in chart_names:
            xml = ET.fromstring(archive.read(name))
            series = xml.findall(".//c:ser", NS)
            assert len(series) == (2 if not legacy and "chart2" in name else 1)
            for series_node in series:
                value_ref = series_node.find("c:val/c:numRef", NS)
                category_ref = series_node.find("c:cat/c:strRef", NS)
                assert value_ref is not None and category_ref is not None
                formula = value_ref.findtext("c:f", namespaces=NS)
                cat_formula = category_ref.findtext("c:f", namespaces=NS)
                sheet_name, address = formula.rsplit("!", 1)
                ws = wb[sheet_name.strip("'")]
                _, cat_address = cat_formula.rsplit("!", 1)
                x1, y1, x2, y2 = range_boundaries(address)
                cat_x1, cat_y1, cat_x2, cat_y2 = range_boundaries(cat_address)
                categories = [
                    str(ws.cell(y, x).value or "")
                    for y in range(cat_y1, cat_y2 + 1)
                    for x in range(cat_x1, cat_x2 + 1)
                ]
                cat_cache = category_ref.find("c:strCache", NS)
                cache = value_ref.find("c:numCache", NS)
                assert cat_cache is not None and cache is not None
                values = [
                    float(p.findtext("c:v", namespaces=NS)) for p in cache.findall("c:pt", NS)
                ]
                cached_categories = [
                    p.findtext("c:v", default="", namespaces=NS)
                    for p in cat_cache.findall("c:pt", NS)
                ]
                assert (
                    int(cache.find("c:ptCount", NS).get("val"))
                    == len(values)
                    == (y2 - y1 + 1) * (x2 - x1 + 1)
                )
                assert (
                    int(cat_cache.find("c:ptCount", NS).get("val"))
                    == len(categories)
                    == len(values)
                )
                assert [int(p.get("idx")) for p in cache.findall("c:pt", NS)] == list(
                    range(len(values))
                )
                assert cached_categories == categories
                if legacy:
                    assert x1 != 3  # Ranking is never a value series.
                    week_start = end - timedelta(days=end.weekday())
                    selected = (
                        [
                            r
                            for r in rows
                            if week_start <= date.fromisoformat(r["inspection_date"]) <= end
                        ]
                        if y1 >= 34
                        else rows
                    )
                    summary = expected_summary(selected)
                    if x1 == 8:
                        expected = [
                            summary[k] for k in ("检验数量", "抽检数量", "不良件数", "不良率")
                        ]
                    elif x1 == 2:
                        counts = Counter(d["code"] for r in selected for d in r["defects"])
                        expected = [counts[chr(97 + i)] for i in range(24)]
                    else:
                        counts = Counter(r["team"] for r in selected if r["judgment"] == "返工")
                        expected = [counts[category] for category in categories]
                    assert values == pytest.approx(expected)
                    for y in range(y1, y2 + 1):
                        for x in range(x1, x2 + 1):
                            assert ws.cell(y, x).data_type == "f"
                else:
                    assert y1 == 2 and ws.cell(y2, 1).value != "合计"
                    assert values == pytest.approx(
                        [ws.cell(y, x).value for y in range(y1, y2 + 1) for x in range(x1, x2 + 1)]
                    )
                    title = series_node.find("c:tx/c:strRef", NS)
                    assert title.find("c:strCache/c:ptCount", NS).get("val") == "1"
                    if x1 in (6, 8):
                        assert cache.findtext("c:formatCode", namespaces=NS) == "0.00%"
                        assert xml.find(".//c:valAx/c:numFmt", NS).get("formatCode") == "0.0%"


def assert_export(ctx, tmp_path, start, end, legacy, source="production", deleted=False):
    rows = raw_records(ctx, start, end, source, deleted)
    filters = RecordFilter(start=start, end=end, source=source, deleted=deleted, descending=False)
    path = ctx.excel.export(tmp_path / "matrix.xlsx", filters, legacy=legacy, prefer_com=False)
    wb = load_workbook(path)
    try:
        ws = wb["成品日检表" if legacy else "检验记录"]
        exported = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[0] is not None]
        assert len(exported) == len(rows)
        for values, row in zip(exported, rows, strict=True):
            assert values[:9] == (
                row["inspection_no"],
                datetime.combine(
                    date.fromisoformat(row["inspection_date"]),
                    time.fromisoformat(row["inspection_time"]),
                ),
                row["team"],
                row["work_order"],
                row["inspection_quantity"],
                row["sampling_quantity"],
                row["defect_quantity"],
                "；".join(d["name"] for d in row["defects"]) or None,
                row["judgment"],
            )
            assert values[9] is None  # Signatures are image anchors, not arbitrary strings.
            assert values[14 if legacy else 16] == (
                ";".join(d["code"] for d in row["defects"]) or None
            )
            if legacy:
                assert values[13] == (
                    "演示数据"
                    if row["source"] == "demo"
                    else "手动录入"
                    if row["source"] == "manual"
                    else "表格导入"
                )
            else:
                assert values[10:13] == (
                    row["inspector"],
                    row["remark"] or None,
                    "演示数据"
                    if row["source"] == "demo"
                    else "手动录入"
                    if row["source"] == "manual"
                    else "表格导入",
                )
                assert values[13:16] == (
                    int(row["inspection_date"][:4]),
                    row["inspection_date"][:7],
                    datetime.fromisoformat(row["inspection_date"]),
                )
        assert all(ws.cell(i, 4).data_type == "s" for i in range(2, len(rows) + 2))
        assert ws.sheet_view.topLeftCell == "A1" and ws.freeze_panes == "C2"
        if legacy:
            assert ws.print_area == f"'成品日检表'!$B$1:$K${max(2, len(rows) + 1)}"
            assert ws.print_title_rows == "$1:$1"
            assert ws.column_dimensions["A"].hidden and ws.column_dimensions["N"].hidden
            assert wb["成品日检表报表"].sheet_state == "hidden"
            assert wb["数据分析表"].merged_cells.ranges
            original = load_workbook(ctx.paths.template)
            try:
                template = original["成品日检表"]
                for n in {2, 500, 501, 502, len(rows) + 1} & set(range(2, len(rows) + 2)):
                    if not ws.cell(n, 8).value:
                        assert ws.row_dimensions[n].height == template.row_dimensions[2].height
                    else:
                        assert template.row_dimensions[2].height <= ws.row_dimensions[n].height <= 409.5
                    for col in range(1, 11):
                        actual, expected = ws.cell(n, col), template.cell(2, col)
                        for attr in ("font", "border", "fill", "number_format", "protection"):
                            assert copy(getattr(actual, attr)) == copy(getattr(expected, attr))
                        alignment = copy(expected.alignment)
                        if col == 8 and actual.value:
                            alignment.wrap_text = True
                        assert actual.alignment == alignment
            finally:
                original.close()
        else:
            assert list(wb["不良明细"].values)[1:] == [
                (r["inspection_no"], d["code"], d["name"], d["quantity"], d["remark"] or None)
                for r in rows
                for d in r["defects"]
            ]
            summary = dict(wb["统计摘要"].values)
            for key, value in expected_summary(rows).items():
                assert summary[key] == pytest.approx(value)
            months = wb["月度统计"]
            grouped = defaultdict(list)
            for row in rows:
                grouped[row["inspection_date"][:7]].append(row)
            for values in list(months.values)[1:]:
                selected = rows if values[0] == "合计" else grouped[values[0]]
                expected = expected_summary(selected)
                assert values[1:8] == pytest.approx(
                    [
                        expected[k]
                        for k in (
                            "检验批次",
                            "检验数量",
                            "抽检数量",
                            "不良件数",
                            "不良率",
                            "返工批次",
                            "返工率",
                        )
                    ]
                )
            assert months.cell(months.max_row, 1).value == "合计"
        assert_charts_xml(path, wb, rows, legacy, end)
        return rows
    finally:
        wb.close()


@pytest.mark.parametrize("legacy", [True, False])
@pytest.mark.parametrize("count", [0, 1, 11, 499, 500, 501, 5000, 10000])
def test_boundary_sizes_field_by_field(ctx, payload, tmp_path, legacy, count):
    with ctx.db.session() as session:
        for i in range(count):
            data = payload.model_dump() | {
                "inspection_date": date(2026, 1, 1) + timedelta(days=i % 365),
                "inspection_time": time(i % 24, i % 60, i % 60),
                "source": "manual" if i % 2 else "excel",
                "team": f"U{i % 3 + 1}",
                "work_order": "=工单<&>" + "长" * 100 + str(i),
                "inspector": "张 & 李",
                "remark": "" if i % 2 else "=备注\n含中文、<> & 特殊字符",
                "inspection_quantity": 100 + i % 97,
                "sampling_quantity": 10 + i % 71,
                "defect_quantity": 0 if i % 3 == 0 else 2,
                "judgment": "返工" if i % 4 == 0 else "合格",
                "defects": []
                if i % 3 == 0
                else [
                    DefectInput(
                        defect_id=i % 24 + 1,
                        quantity=None if i % 2 else 1,
                        remark="=不良备注<&>" if i % 2 else "",
                    )
                ],
            }
            ctx.inspections.save_in_session(
                session,
                InspectionInput(**data),
                inspection_no=f"边界-{i:06d}",
                validate_references=False,
                flush=False,
            )
    rows = assert_export(ctx, tmp_path, date(2026, 1, 1), date(2026, 12, 31), legacy)
    assert len(rows) == count


@pytest.mark.parametrize("legacy", [True, False])
@pytest.mark.parametrize("year,count,days", [(2026, 1825, 365), (2024, 1830, 366)])
def test_full_annual_daily_coverage_from_raw_sqlite(ctx, tmp_path, legacy, year, count, days):
    start, end = date(year, 1, 1), date(year, 12, 31)
    ctx.demo.generate(count, start, end, ["U1", "U2", "U3"], seed=20261008)
    rows = assert_export(ctx, tmp_path, start, end, legacy, "demo")
    assert len(rows) == count
    assert len({r["inspection_date"] for r in rows}) == days
    assert {r["inspection_date"][:7] for r in rows} == {f"{year}-{m:02d}" for m in range(1, 13)}
    if year == 2024:
        assert any(r["inspection_date"] == "2024-02-29" for r in rows)


@pytest.mark.parametrize("legacy", [True, False])
@pytest.mark.parametrize(
    "start,end",
    [
        ("2026-01-01", "2026-12-31"),
        ("2026-01-01", "2026-01-31"),
        ("2026-01-05", "2026-01-05"),
        ("2026-01-05", "2026-02-05"),
        ("2025-12-30", "2026-01-05"),
        ("2024-02-29", "2024-02-29"),
        ("2023-01-01", "2023-12-31"),
    ],
)
def test_range_and_source_matrix(ctx, payload, tmp_path, legacy, start, end):
    for day in (
        "2025-12-29",
        "2025-12-30",
        "2026-01-05",
        "2026-01-31",
        "2026-02-05",
        "2026-12-31",
        "2024-02-29",
    ):
        for source in ("manual", "excel", "demo"):
            ctx.inspections.save(
                InspectionInput(
                    **(payload.model_dump() | {"inspection_date": day, "source": source})
                )
            )
    deleted = ctx.inspections.save(payload)
    ctx.inspections.delete([deleted["id"]])
    for source in ("production", "demo"):
        assert_export(
            ctx, tmp_path, date.fromisoformat(start), date.fromisoformat(end), legacy, source
        )
